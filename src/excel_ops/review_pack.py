from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation


REVIEW_SHEET = "Review"
HISTORY_SHEET = "Review History"
MANIFEST_SHEET = "_Manifest"
DECISIONS = ("accept", "correct", "reject", "cannot determine")
HEADERS = (
    "Stable record ID",
    "Source value",
    "Source location",
    "Proposed value",
    "Review reason",
    "Confidence",
    "Candidates",
    "Decision",
    "Corrected value",
    "Reviewer note",
)
HISTORY_HEADERS = (
    "Recorded at (UTC)",
    "Stable record ID",
    "Decision",
    "Corrected value",
    "Reviewer note",
)


@dataclass(frozen=True)
class ReviewPackRow:
    record_id: str
    source_value: Any
    source_location: str
    proposed_value: Any
    review_reason: str
    confidence: float
    candidates: tuple[str, ...] = ()


@dataclass(frozen=True)
class ReviewDecision:
    record_id: str
    decision: str
    value: Any | None
    reviewer_note: str


@dataclass(frozen=True)
class ReviewHistoryEntry:
    recorded_at: str
    record_id: str
    decision: str
    corrected_value: Any | None
    reviewer_note: str


@dataclass(frozen=True)
class ReviewImportResult:
    decisions: dict[str, ReviewDecision]
    pending_record_ids: tuple[str, ...]
    accepted_values: dict[str, Any]
    history: tuple[ReviewHistoryEntry, ...]


def write_review_pack(
    rows: Iterable[ReviewPackRow],
    output_path: str | Path,
    *,
    history: Iterable[ReviewHistoryEntry] = (),
) -> Path:
    """Write an offline pack whose editable fields are safe to round-trip."""

    review_rows = tuple(rows)
    _validate_rows(review_rows)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    review = workbook.active
    review.title = REVIEW_SHEET
    review.append(HEADERS)
    _style_header(review)

    for item in review_rows:
        review.append(
            [
                item.record_id,
                item.source_value,
                item.source_location,
                item.proposed_value,
                item.review_reason,
                item.confidence,
                json.dumps(item.candidates, ensure_ascii=False),
                "",
                "",
                "",
            ]
        )

    validation = DataValidation(
        type="list", formula1='"accept,correct,reject,cannot determine"', allow_blank=True
    )
    validation.error = "Choose accept, correct, reject, or cannot determine."
    validation.errorTitle = "Invalid review decision"
    validation.showErrorMessage = True
    review.add_data_validation(validation)
    if review.max_row > 1:
        validation.add(f"H2:H{review.max_row}")

    for row in review.iter_rows(min_row=2):
        for cell in row:
            cell.protection = Protection(locked=cell.column not in (8, 9, 10))
    review.protection.sheet = True
    review.freeze_panes = "A2"
    review.auto_filter.ref = f"A1:J{review.max_row}"

    history_sheet = workbook.create_sheet(HISTORY_SHEET)
    history_sheet.append(HISTORY_HEADERS)
    _style_header(history_sheet)
    for entry in history:
        history_sheet.append(
            [
                entry.recorded_at,
                entry.record_id,
                entry.decision,
                entry.corrected_value,
                entry.reviewer_note,
            ]
        )
    history_sheet.freeze_panes = "A2"

    manifest = workbook.create_sheet(MANIFEST_SHEET)
    manifest.append(["Format", "excel-ops-review-pack-v1"])
    manifest.append(["Headers", json.dumps(HEADERS)])
    manifest.append(["Stable record ID", "Original row digest"])
    for item in review_rows:
        manifest.append([item.record_id, _row_digest(item)])
    manifest.sheet_state = "veryHidden"

    for sheet in (review, history_sheet):
        for column in sheet.columns:
            width = min(56, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width

    workbook.save(output)
    return output


def import_review_pack(
    path: str | Path, *, record_history: bool = True
) -> ReviewImportResult:
    """Validate an edited pack and return decisions keyed by stable record ID.

    Blank decisions remain pending. ``cannot determine`` is deliberately absent
    from ``accepted_values``, so it cannot become an automatic acceptance.
    """

    workbook_path = Path(path)
    workbook = load_workbook(workbook_path)
    _validate_workbook_structure(workbook)
    review = workbook[REVIEW_SHEET]
    manifest = workbook[MANIFEST_SHEET]

    expected = {
        str(record_id): str(digest)
        for record_id, digest in manifest.iter_rows(min_row=4, max_col=2, values_only=True)
        if record_id is not None
    }
    actual_ids: list[str] = []
    decisions: dict[str, ReviewDecision] = {}
    pending: list[str] = []
    new_history: list[ReviewHistoryEntry] = []
    accepted_values: dict[str, Any] = {}
    existing_history = _read_history(workbook[HISTORY_SHEET])
    existing_decisions = {
        (entry.record_id, entry.decision, entry.corrected_value, entry.reviewer_note)
        for entry in existing_history
    }

    for values in review.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
        record_id = str(values[0] or "").strip()
        if not record_id:
            workbook.close()
            raise ValueError("Review Pack contains a row without a stable record ID")
        actual_ids.append(record_id)
        item = ReviewPackRow(
            record_id=record_id,
            source_value=values[1],
            source_location=str(values[2] or ""),
            proposed_value=values[3],
            review_reason=str(values[4] or ""),
            confidence=float(values[5] or 0),
            candidates=_parse_candidates(values[6]),
        )
        if expected.get(record_id) != _row_digest(item):
            workbook.close()
            raise ValueError(f"Original review data changed for stable record ID: {record_id}")

        decision = str(values[7] or "").strip().casefold()
        corrected = values[8]
        note = str(values[9] or "")
        if not decision:
            pending.append(record_id)
            continue
        if decision not in DECISIONS:
            workbook.close()
            raise ValueError(f"Invalid decision for {record_id}: {decision}")
        if decision == "correct" and (corrected is None or str(corrected).strip() == ""):
            workbook.close()
            raise ValueError(f"Decision 'correct' requires a corrected value for {record_id}")

        value = values[3] if decision == "accept" else corrected if decision == "correct" else None
        decisions[record_id] = ReviewDecision(record_id, decision, value, note)
        if decision in ("accept", "correct"):
            accepted_values[record_id] = value
        signature = (record_id, decision, corrected, note)
        if signature not in existing_decisions:
            new_history.append(
                ReviewHistoryEntry(
                    datetime.now(timezone.utc).isoformat(), record_id, decision, corrected, note
                )
            )

    if len(actual_ids) != len(set(actual_ids)):
        workbook.close()
        raise ValueError("Review Pack contains duplicate stable record IDs")
    if set(actual_ids) != set(expected):
        workbook.close()
        raise ValueError("Review Pack stable record IDs do not match the original manifest")

    combined_history = (*existing_history, *new_history)
    if record_history and new_history:
        history_sheet = workbook[HISTORY_SHEET]
        for entry in new_history:
            history_sheet.append(
                [entry.recorded_at, entry.record_id, entry.decision, entry.corrected_value, entry.reviewer_note]
            )
        workbook.save(workbook_path)
    workbook.close()
    return ReviewImportResult(decisions, tuple(pending), accepted_values, combined_history)


def review_rows_from_match_results(results: Iterable[Any]) -> list[ReviewPackRow]:
    """Adapt #2 MatchResult objects without coupling this module to matching.py."""

    rows: list[ReviewPackRow] = []
    for result in results:
        if getattr(result, "status", None) == "matched":
            continue
        record = result.record
        candidates = tuple(
            str(getattr(candidate, "destination_key", candidate))
            for candidate in getattr(result, "candidates", ())
        )
        rows.append(
            ReviewPackRow(
                record_id=str(result.record_id),
                source_value=record.location,
                source_location=record.source,
                proposed_value=getattr(result, "destination_key", None)
                or (candidates[0] if candidates else None),
                review_reason="; ".join(getattr(result, "reasons", ())),
                confidence=float(getattr(result, "confidence", 0)),
                candidates=candidates,
            )
        )
    return rows


def _validate_rows(rows: tuple[ReviewPackRow, ...]) -> None:
    ids = [item.record_id.strip() for item in rows]
    if any(not record_id for record_id in ids):
        raise ValueError("Stable record IDs must not be blank")
    if len(ids) != len(set(ids)):
        raise ValueError("Stable record IDs must be unique")


def _validate_workbook_structure(workbook) -> None:
    required = {REVIEW_SHEET, HISTORY_SHEET, MANIFEST_SHEET}
    if not required.issubset(workbook.sheetnames):
        raise ValueError("Review Pack is missing a required sheet")
    headers = tuple(workbook[REVIEW_SHEET].cell(1, index).value for index in range(1, len(HEADERS) + 1))
    if headers != HEADERS:
        raise ValueError("Review Pack headers were changed")
    if workbook[MANIFEST_SHEET]["B1"].value != "excel-ops-review-pack-v1":
        raise ValueError("Unsupported or damaged Review Pack manifest")


def _row_digest(item: ReviewPackRow) -> str:
    payload = json.dumps(
        [
            item.record_id,
            item.source_value,
            item.source_location,
            item.proposed_value,
            item.review_reason,
            item.confidence,
            item.candidates,
        ],
        ensure_ascii=False,
        sort_keys=True,
        default=str,
        separators=(",", ":"),
    )
    return sha256(payload.encode("utf-8")).hexdigest()


def _parse_candidates(value: Any) -> tuple[str, ...]:
    try:
        parsed = json.loads(str(value or "[]"))
    except json.JSONDecodeError as exc:
        raise ValueError("Review Pack candidates were changed or damaged") from exc
    if not isinstance(parsed, list):
        raise ValueError("Review Pack candidates must be a JSON list")
    return tuple(str(item) for item in parsed)


def _read_history(sheet) -> tuple[ReviewHistoryEntry, ...]:
    entries: list[ReviewHistoryEntry] = []
    for values in sheet.iter_rows(min_row=2, max_col=len(HISTORY_HEADERS), values_only=True):
        if values[1] is None:
            continue
        entries.append(
            ReviewHistoryEntry(
                str(values[0] or ""), str(values[1]), str(values[2] or ""), values[3], str(values[4] or "")
            )
        )
    return tuple(entries)


def _style_header(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
