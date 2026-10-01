"""Apply an explicit health-repair allowlist to a new workbook copy."""

from __future__ import annotations

from decimal import Decimal
from hashlib import sha256
import os
from pathlib import Path
import secrets
import stat
import tempfile
from typing import Any, Iterable
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, range_to_tuple

from .cells import is_merged_non_anchor, is_protected_formula_value
from .health_models import (
    HealthRepairAction,
    HealthRepairApplyError,
    HealthRepairChange,
    HealthRepairPlan,
    HealthRepairResult,
)
from .health_repair_plan import plan_workbook_health_repairs

_APPLICABLE_OPERATIONS = {
    "normalize_whitespace",
    "remove_invisible_characters",
    "convert_numeric_text",
    "rename_duplicate_header",
    "align_filter_range",
    "align_freeze_panes",
    "align_print_area",
    "declare_date_format",
    "declare_blank_zero_policy",
}


class _DirectoryGuard:
    """Hold or address one output directory without following a replaced path."""

    def __init__(self, path: Path, expected: tuple[int, int]) -> None:
        self.path = path
        self.expected = expected
        self.fd: int | None = None
        self.handle: int | None = None

    def __enter__(self) -> "_DirectoryGuard":
        if os.name == "nt":
            self._open_windows_lock()
        elif os.open in os.supports_dir_fd and os.link in os.supports_dir_fd:
            self.fd = os.open(self.path, os.O_RDONLY | getattr(os, "O_DIRECTORY", 0))
            if self._directory_identity() != self.expected:
                self.close()
                raise HealthRepairApplyError("output directory changed before locking")
        else:
            raise HealthRepairApplyError("platform cannot safely lock or address output directory")
        return self

    def _open_windows_lock(self) -> None:
        import ctypes
        from ctypes import wintypes

        create_file = ctypes.windll.kernel32.CreateFileW
        create_file.argtypes = [
            wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
            wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE,
        ]
        create_file.restype = wintypes.HANDLE
        handle = create_file(
            str(self.path), 0x80000000, 0x00000001 | 0x00000002,
            None, 3, 0x02000000, None,
        )
        invalid = ctypes.c_void_p(-1).value
        if handle == invalid:
            raise HealthRepairApplyError("could not lock output directory against replacement")
        self.handle = int(handle)
        if _identity(self.path) != self.expected:
            self.close()
            raise HealthRepairApplyError("output directory changed before locking")

    def _directory_identity(self) -> tuple[int, int]:
        if self.fd is not None:
            details = os.fstat(self.fd)
            return details.st_dev, details.st_ino
        return _identity(self.path)

    def create_temp(self, prefix: str, suffix: str) -> tuple[Path, int, tuple[int, int]]:
        if self.fd is None:
            descriptor, name = tempfile.mkstemp(prefix=prefix, suffix=suffix, dir=self.path)
            path = Path(name)
        else:
            for _ in range(100):
                leaf = f"{prefix}{secrets.token_hex(12)}{suffix}"
                try:
                    descriptor = os.open(
                        leaf, os.O_RDWR | os.O_CREAT | os.O_EXCL,
                        0o600, dir_fd=self.fd,
                    )
                except FileExistsError:
                    continue
                path = self.path / leaf
                break
            else:
                raise HealthRepairApplyError("could not allocate a secure staged filename")
        details = os.fstat(descriptor)
        return path, descriptor, (details.st_dev, details.st_ino)

    def link(self, staged: Path, output: Path) -> None:
        try:
            if self.fd is None:
                os.link(staged, output)
            else:
                os.link(staged.name, output.name, src_dir_fd=self.fd, dst_dir_fd=self.fd)
        except FileExistsError as exc:
            raise HealthRepairApplyError(f"output already exists: {output}") from exc
        except OSError as exc:
            raise HealthRepairApplyError(f"could not publish repaired workbook: {exc}") from exc

    def identity(self, path: Path) -> tuple[int, int]:
        if self.fd is None:
            return _identity(path)
        details = os.stat(path.name, dir_fd=self.fd, follow_symlinks=False)
        return details.st_dev, details.st_ino

    def stable_path(self, path: Path) -> Path:
        if self.fd is None:
            return path
        for root in (Path("/proc/self/fd"), Path("/dev/fd")):
            candidate = root / str(self.fd) / path.name
            try:
                details = candidate.stat(follow_symlinks=False)
            except OSError:
                continue
            expected = self.identity(path)
            if (details.st_dev, details.st_ino) == expected:
                return candidate
        raise HealthRepairApplyError(
            "platform cannot expose a stable output-directory path to the workbook library"
        )

    def ensure_visible(self) -> None:
        try:
            visible = _identity(self.path)
        except OSError as exc:
            raise HealthRepairApplyError("output directory path changed during repair") from exc
        if visible != self.expected:
            raise HealthRepairApplyError("output directory path changed during repair")

    def unlink(self, path: Path, expected: tuple[int, int] | None) -> None:
        if expected is None:
            return
        try:
            if self.identity(path) != expected:
                return
            if self.fd is None:
                path.unlink()
            else:
                os.unlink(path.name, dir_fd=self.fd)
        except (FileNotFoundError, OSError):
            pass

    def close(self) -> None:
        if self.fd is not None:
            os.close(self.fd)
            self.fd = None
        if self.handle is not None:
            import ctypes
            from ctypes import wintypes
            close_handle = ctypes.windll.kernel32.CloseHandle
            close_handle.argtypes = [wintypes.HANDLE]
            close_handle.restype = wintypes.BOOL
            close_handle(self.handle)
            self.handle = None

    def __exit__(self, *_: Any) -> None:
        self.close()


def _opened_digest(path: Path, expected_identity: tuple[int, int]) -> str:
    flags = os.O_RDONLY | getattr(os, "O_BINARY", 0) | getattr(os, "O_NOFOLLOW", 0)
    descriptor = os.open(path, flags)
    try:
        details = os.fstat(descriptor)
        if (details.st_dev, details.st_ino) != expected_identity:
            raise HealthRepairApplyError("source workbook changed during repair")
        checksum = sha256()
        while chunk := os.read(descriptor, 1024 * 1024):
            checksum.update(chunk)
        return checksum.hexdigest()
    finally:
        os.close(descriptor)


def _descriptor_digest(descriptor: int, expected_identity: tuple[int, int]) -> str:
    duplicated = os.dup(descriptor)
    try:
        details = os.fstat(duplicated)
        if (details.st_dev, details.st_ino) != expected_identity:
            raise HealthRepairApplyError("staged output identity changed during hashing")
        os.lseek(duplicated, 0, os.SEEK_SET)
        checksum = sha256()
        while chunk := os.read(duplicated, 1024 * 1024):
            checksum.update(chunk)
        return checksum.hexdigest()
    finally:
        os.close(duplicated)


def _stable_source_copy(
    source: Path, guard: _DirectoryGuard, expected_identity: tuple[int, int]
) -> tuple[Path, tuple[int, int], str]:
    flags = os.O_RDONLY | getattr(os, "O_BINARY", 0) | getattr(os, "O_NOFOLLOW", 0)
    source_fd = os.open(source, flags)
    copied, copy_fd, copy_identity = guard.create_temp(
        prefix=f".{source.stem}-snapshot-", suffix=source.suffix
    )
    checksum = sha256()
    try:
        details = os.fstat(source_fd)
        if (details.st_dev, details.st_ino) != expected_identity:
            raise HealthRepairApplyError("source workbook changed before snapshot")
        while chunk := os.read(source_fd, 1024 * 1024):
            checksum.update(chunk)
            view = memoryview(chunk)
            while view:
                view = view[os.write(copy_fd, view):]
        os.fsync(copy_fd)
    except Exception:
        os.close(source_fd)
        os.close(copy_fd)
        guard.unlink(copied, copy_identity)
        raise
    else:
        os.close(source_fd)
        os.close(copy_fd)
    return copied, copy_identity, checksum.hexdigest()


def _reject_symlink(path: Path, label: str) -> None:
    current = path.absolute()
    while True:
        try:
            mode = current.lstat().st_mode
        except FileNotFoundError:
            pass
        else:
            if stat.S_ISLNK(mode):
                raise HealthRepairApplyError(f"{label} must not contain a symlink: {current}")
        if current.parent == current:
            return
        current = current.parent


def _validate_paths(source: Path, output: Path) -> tuple[tuple[int, int], tuple[int, int]]:
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise HealthRepairApplyError("source must be an .xlsx or .xlsm workbook")
    if output.suffix.lower() != source.suffix.lower():
        raise HealthRepairApplyError("output must preserve the source workbook extension")
    if not source.is_file():
        raise HealthRepairApplyError(f"source workbook does not exist: {source}")
    if not output.parent.is_dir():
        raise HealthRepairApplyError(f"output directory does not exist: {output.parent}")
    _reject_symlink(source, "source")
    _reject_symlink(output.parent, "output parent")
    if source.resolve() == output.resolve(strict=False):
        raise HealthRepairApplyError("output must be a new workbook path")
    try:
        output.lstat()
    except FileNotFoundError:
        return _identity(source), _identity(output.parent)
    raise HealthRepairApplyError(f"output already exists: {output}")


def _identity(path: Path) -> tuple[int, int]:
    details = path.stat(follow_symlinks=False)
    return details.st_dev, details.st_ino


def _require_identity(path: Path, expected: tuple[int, int], label: str) -> None:
    try:
        actual = _identity(path)
    except OSError as exc:
        raise HealthRepairApplyError(f"{label} changed during repair") from exc
    if actual != expected:
        raise HealthRepairApplyError(f"{label} changed during repair")


def _baseline_for(action: HealthRepairAction) -> dict[str, Any] | None:
    target = dict(action.proposed_state.get("target") or {})
    if action.operation == "declare_style_baseline":
        return {"target_signature_digest": target.get("target_signature_digest")}
    if action.operation == "declare_date_format":
        return {key: target.get(key) for key in ("target_number_format", "locale")}
    if action.operation == "declare_blank_zero_policy":
        return {key: target.get(key) for key in ("blank_action", "zero_action")}
    if action.operation == "rename_duplicate_header":
        return {"target_header": target.get("target_header")}
    if action.operation == "convert_numeric_text" and "target_representation" in target:
        return {"target_representation": target["target_representation"]}
    return None


def _canonical_plan(source: Path, supplied: HealthRepairPlan) -> HealthRepairPlan:
    dispositions = {
        action.action_id: action.decision
        for action in supplied.actions
        if action.decision != "proposed"
    }
    baselines = {
        action.action_id: baseline
        for action in supplied.actions
        if (baseline := _baseline_for(action)) is not None
    }
    try:
        canonical = plan_workbook_health_repairs(
            source, dispositions=dispositions, baselines=baselines
        )
    except Exception as exc:
        raise HealthRepairApplyError("repair plan does not match the current source") from exc
    canonical_actions = [action.to_dict() for action in canonical.actions]
    supplied_actions = [action.to_dict() for action in supplied.actions]
    if canonical.source_sha256 != supplied.source_sha256 or canonical_actions != supplied_actions:
        raise HealthRepairApplyError("repair plan was forged, mutated, or is stale")
    return canonical


def _selected(plan: HealthRepairPlan, selected_action_ids: Iterable[str]) -> tuple[HealthRepairAction, ...]:
    requested = tuple(selected_action_ids)
    if not requested:
        raise HealthRepairApplyError("selected action allowlist must not be empty")
    if len(requested) != len(set(requested)):
        raise HealthRepairApplyError("selected action allowlist contains duplicates")
    known = {action.action_id: action for action in plan.actions}
    unknown = sorted(set(requested) - set(known))
    if unknown:
        raise HealthRepairApplyError(f"unknown action id(s): {unknown!r}")
    actions = tuple(known[action_id] for action_id in requested)
    invalid = [
        action.action_id
        for action in actions
        if action.decision != "selected" or action.proposed_state.get("status") != "ready"
    ]
    if invalid:
        raise HealthRepairApplyError(
            f"allowlisted action(s) are not selected and ready in the plan: {invalid!r}"
        )
    dangerous = [action.action_id for action in actions if action.operation not in _APPLICABLE_OPERATIONS]
    if dangerous:
        raise HealthRepairApplyError(
            f"selected action(s) cannot be proven to preserve workbook structure: {dangerous!r}"
        )
    return actions


def _numeric_value(candidate: Any) -> int | float:
    if isinstance(candidate, int) and not isinstance(candidate, bool):
        return candidate
    if isinstance(candidate, dict) and candidate.get("kind") == "decimal":
        return float(Decimal(candidate["canonical"]))
    raise HealthRepairApplyError("numeric target is not a supported exact candidate")


def _change(action: HealthRepairAction, before: Any, after: Any) -> HealthRepairChange:
    return HealthRepairChange(
        action_id=action.action_id,
        operation=action.operation,
        sheet=action.sheet,
        location=action.location,
        before=before,
        after=after,
    )


def _apply_cell_action(sheet: Any, action: HealthRepairAction, target: dict[str, Any]) -> HealthRepairChange:
    cell = sheet[action.location]
    if is_merged_non_anchor(cell):
        raise HealthRepairApplyError(f"cannot mutate merged-range follower: {action.impact_scope}")
    if is_protected_formula_value(cell.value):
        raise HealthRepairApplyError(f"cannot mutate protected formula: {action.impact_scope}")
    before = cell.value
    if action.operation in {"normalize_whitespace", "remove_invisible_characters"}:
        cell.value = target["candidate_value"]
    elif action.operation == "convert_numeric_text":
        if target.get("target_representation") == "preserve_text":
            cell.value = target["value"]
        else:
            cell.value = _numeric_value(target["candidate_value"])
    elif action.operation == "rename_duplicate_header":
        cell.value = target["target_header"]
    else:
        raise HealthRepairApplyError(f"unsupported cell operation: {action.operation}")
    return _change(action, before, cell.value)


def _apply_range_action(sheet: Any, action: HealthRepairAction, target: dict[str, Any]) -> HealthRepairChange:
    if action.operation == "align_filter_range":
        before = sheet.auto_filter.ref
        sheet.auto_filter.ref = target["candidate_ref"]
        return _change(action, before, sheet.auto_filter.ref)
    if action.operation == "align_freeze_panes":
        before = str(sheet.freeze_panes) if sheet.freeze_panes else None
        sheet.freeze_panes = target["candidate_coordinate"]
        return _change(action, before, str(sheet.freeze_panes))
    if action.operation == "align_print_area":
        before = str(sheet.print_area)
        sheet.print_area = target["candidate_refs"]
        return _change(action, before, list(target["candidate_refs"]))
    if action.operation == "declare_date_format":
        coordinates = sorted(
            coordinate
            for values in action.current_state["summary"]["formats"].values()
            for coordinate in values
        )
        before = {coordinate: sheet[coordinate].number_format for coordinate in coordinates}
        for coordinate in coordinates:
            sheet[coordinate].number_format = target["target_number_format"]
        after = {coordinate: sheet[coordinate].number_format for coordinate in coordinates}
        return _change(action, before, after)
    raise HealthRepairApplyError(f"unsupported range operation: {action.operation}")


def _apply_semantic_action(sheet: Any, action: HealthRepairAction, target: dict[str, Any]) -> HealthRepairChange:
    summary = action.current_state["summary"]
    before: dict[str, Any] = {}
    after: dict[str, Any] = {}
    for coordinate in summary["blank_locations"]:
        before[coordinate] = sheet[coordinate].value
        if target["blank_action"] == "convert_to_zero":
            sheet[coordinate] = 0
        after[coordinate] = sheet[coordinate].value
    for coordinate in summary["zero_locations"]:
        before[coordinate] = sheet[coordinate].value
        if target["zero_action"] == "convert_to_blank":
            sheet[coordinate] = None
        after[coordinate] = sheet[coordinate].value
    return _change(action, before, after)


def _apply_action(workbook: Any, action: HealthRepairAction) -> HealthRepairChange:
    if action.sheet not in workbook.sheetnames:
        raise HealthRepairApplyError(f"worksheet is stale or missing: {action.sheet}")
    sheet = workbook[action.sheet]
    target = dict(action.proposed_state["target"])
    if action.operation in {
        "normalize_whitespace", "remove_invisible_characters", "convert_numeric_text",
        "rename_duplicate_header",
    }:
        return _apply_cell_action(sheet, action, target)
    if action.operation in {
        "align_filter_range", "align_freeze_panes", "align_print_area",
        "declare_date_format",
    }:
        return _apply_range_action(sheet, action, target)
    if action.operation == "declare_blank_zero_policy":
        return _apply_semantic_action(sheet, action, target)
    raise HealthRepairApplyError(f"operation is not supported for application: {action.operation}")


def _meaningful_cells(sheet: Any) -> dict[str, tuple[Any, str, str, int]]:
    result: dict[str, tuple[Any, str, str, int]] = {}
    for cell in getattr(sheet, "_cells", {}).values():
        if cell.value is None and not cell.has_style and cell.comment is None and cell.hyperlink is None:
            continue
        result[cell.coordinate] = (
            cell.value,
            cell.data_type,
            cell.number_format,
            cell.style_id,
        )
    return result


def _validations(sheet: Any) -> tuple[str, ...]:
    return tuple(
        sorted(
            repr(
                (
                    str(item.sqref), item.type, item.operator, item.formula1, item.formula2,
                    item.allow_blank, item.showErrorMessage, item.error, item.errorTitle,
                )
            )
            for item in sheet.data_validations.dataValidation
        )
    )


def _print_refs(value: Any) -> tuple[str, ...]:
    if not value:
        return ()
    if hasattr(value, "ranges"):
        return tuple(str(item) for item in value.ranges)
    refs: list[str] = []
    current: list[str] = []
    quoted = False
    for character in str(value):
        if character == "'":
            quoted = not quoted
        if character == "," and not quoted:
            refs.append("".join(current))
            current = []
        else:
            current.append(character)
    if current:
        refs.append("".join(current))
    return tuple(ref for ref in refs if ref)


def _range_bounds(reference: str) -> tuple[int, int, int, int]:
    if "!" in reference:
        _, bounds = range_to_tuple(reference)
        return bounds
    return range_boundaries(reference)


def _workbook_snapshot(workbook: Any) -> dict[str, Any]:
    sheets: dict[str, Any] = {}
    for sheet in workbook.worksheets:
        sheets[sheet.title] = {
            "state": sheet.sheet_state,
            "cells": _meaningful_cells(sheet),
            "merged": tuple(sorted(str(item) for item in sheet.merged_cells.ranges)),
            "validations": _validations(sheet),
            "auto_filter": str(sheet.auto_filter.ref) if sheet.auto_filter.ref else None,
            "freeze": str(sheet.freeze_panes) if sheet.freeze_panes else None,
            "print_area": _print_refs(sheet.print_area),
            "protection": repr(sorted(sheet.protection.__dict__.items())),
            "hidden_rows": tuple(sorted(key for key, item in sheet.row_dimensions.items() if item.hidden)),
            "hidden_columns": tuple(sorted(key for key, item in sheet.column_dimensions.items() if item.hidden)),
        }
    return {"sheetnames": tuple(workbook.sheetnames), "sheets": sheets}


def _authorized_cells(action: HealthRepairAction) -> tuple[set[str], set[str]]:
    value_cells: set[str] = set()
    format_cells: set[str] = set()
    if action.operation in {
        "normalize_whitespace", "remove_invisible_characters", "convert_numeric_text",
        "rename_duplicate_header",
    }:
        value_cells.add(action.location)
    elif action.operation == "declare_blank_zero_policy":
        summary = action.current_state["summary"]
        value_cells.update(summary["blank_locations"])
        value_cells.update(summary["zero_locations"])
    elif action.operation == "declare_date_format":
        for coordinates in action.current_state["summary"]["formats"].values():
            format_cells.update(coordinates)
    return value_cells, format_cells


def _neutralize_authorized(
    before: dict[str, Any], after: dict[str, Any], actions: tuple[HealthRepairAction, ...]
) -> None:
    for action in actions:
        before_sheet = before["sheets"][action.sheet]
        after_sheet = after["sheets"][action.sheet]
        values, formats = _authorized_cells(action)
        for coordinate in values | formats:
            old = before_sheet["cells"].get(coordinate, (None, "n", "General", 0))
            new = after_sheet["cells"].get(coordinate, (None, "n", "General", 0))
            if coordinate in values:
                new = (old[0], old[1], new[2], new[3])
            if coordinate in formats:
                new = (new[0], new[1], old[2], old[3])
            after_sheet["cells"][coordinate] = new
            before_sheet["cells"].setdefault(coordinate, old)
        field = {
            "align_filter_range": "auto_filter",
            "align_freeze_panes": "freeze",
            "align_print_area": "print_area",
        }.get(action.operation)
        if field:
            after_sheet[field] = before_sheet[field]


def _verify_unauthorized_unchanged(
    before: dict[str, Any], after: dict[str, Any], actions: tuple[HealthRepairAction, ...]
) -> None:
    _neutralize_authorized(before, after, actions)
    if after != before:
        raise HealthRepairApplyError("persisted output changed unauthorized workbook content or structure")


def _observation(workbook: Any, action: HealthRepairAction) -> Any:
    sheet = workbook[action.sheet]
    values, formats = _authorized_cells(action)
    if values:
        observed = {coordinate: sheet[coordinate].value for coordinate in sorted(values)}
        return next(iter(observed.values())) if len(observed) == 1 else observed
    if formats:
        return {coordinate: sheet[coordinate].number_format for coordinate in sorted(formats)}
    if action.operation == "align_filter_range":
        return sheet.auto_filter.ref
    if action.operation == "align_freeze_panes":
        return str(sheet.freeze_panes) if sheet.freeze_panes else None
    if action.operation == "align_print_area":
        return _print_refs(sheet.print_area)
    raise HealthRepairApplyError(f"cannot observe operation: {action.operation}")


def _persisted_ledger(source: Any, output: Any, actions: tuple[HealthRepairAction, ...]) -> tuple[HealthRepairChange, ...]:
    return tuple(
        _change(action, _observation(source, action), _observation(output, action))
        for action in actions
    )


def _assert_persisted_targets(workbook: Any, actions: tuple[HealthRepairAction, ...]) -> None:
    for action in actions:
        observed = _observation(workbook, action)
        target = action.proposed_state["target"]
        if action.operation in {"normalize_whitespace", "remove_invisible_characters"}:
            expected = target["candidate_value"]
        elif action.operation == "convert_numeric_text":
            expected = (
                target["value"]
                if "value" in target
                else _numeric_value(target["candidate_value"])
            )
        elif action.operation == "rename_duplicate_header":
            expected = target["target_header"]
        elif action.operation == "align_filter_range":
            expected = target["candidate_ref"]
        elif action.operation == "align_freeze_panes":
            expected = target["candidate_coordinate"]
        elif action.operation == "align_print_area":
            expected_bounds = {_range_bounds(item) for item in target["candidate_refs"]}
            if {_range_bounds(item) for item in observed} != expected_bounds:
                raise HealthRepairApplyError(f"persisted target mismatch: {action.action_id}")
            continue
        elif action.operation == "declare_date_format":
            expected = {coordinate: target["target_number_format"] for coordinate in observed}
        elif action.operation == "declare_blank_zero_policy":
            summary = action.current_state["summary"]
            expected = {coordinate: None for coordinate in summary["blank_locations"]}
            expected.update({coordinate: 0 for coordinate in summary["zero_locations"]})
            if target["blank_action"] == "convert_to_zero":
                expected.update({coordinate: 0 for coordinate in summary["blank_locations"]})
            if target["zero_action"] == "convert_to_blank":
                expected.update({coordinate: None for coordinate in summary["zero_locations"]})
        else:
            raise HealthRepairApplyError(f"cannot verify operation: {action.operation}")
        if observed != expected:
            raise HealthRepairApplyError(f"persisted target mismatch: {action.action_id}")


def _publish_without_overwrite(
    guard: _DirectoryGuard, staged: Path, output: Path
) -> None:
    guard.link(staged, output)


def _verify_xlsm_package(source: Path, staged: Any) -> None:
    if source.suffix.lower() != ".xlsm":
        return
    try:
        with ZipFile(source) as original, ZipFile(staged) as repaired:
            if repaired.testzip() is not None:
                raise HealthRepairApplyError("repaired XLSM package contains a corrupt member")
            macro_members = [name for name in original.namelist() if name.endswith("vbaProject.bin")]
            for name in macro_members:
                if name not in repaired.namelist() or repaired.read(name) != original.read(name):
                    raise HealthRepairApplyError("repaired XLSM did not preserve its VBA project")
    except BadZipFile as exc:
        raise HealthRepairApplyError("source or repaired XLSM is not a valid ZIP package") from exc


def _close_workbook(workbook: Any) -> None:
    workbook.close()
    vba_archive = getattr(workbook, "vba_archive", None)
    if vba_archive is not None:
        vba_archive.close()


def _duplicate_binary_stream(descriptor: int) -> Any:
    duplicated = os.dup(descriptor)
    os.lseek(duplicated, 0, os.SEEK_SET)
    return os.fdopen(duplicated, "rb")


def apply_workbook_health_repairs(
    source: str | Path,
    output: str | Path,
    plan: HealthRepairPlan,
    *,
    selected_action_ids: Iterable[str],
) -> HealthRepairResult:
    """Apply only explicitly allowlisted, selected, ready actions to a new copy."""

    source_path = Path(source).absolute()
    output_path = Path(output).absolute()
    source_identity, parent_identity = _validate_paths(source_path, output_path)
    if Path(plan.source).resolve() != source_path.resolve():
        raise HealthRepairApplyError("repair plan belongs to a different source path")
    _require_identity(source_path, source_identity, "source workbook")
    with _DirectoryGuard(output_path.parent, parent_identity) as guard:
        return _apply_with_directory_guard(
            source_path, output_path, plan, tuple(selected_action_ids),
            source_identity, guard,
        )


def _apply_with_directory_guard(
    source_path: Path,
    output_path: Path,
    plan: HealthRepairPlan,
    selected_action_ids: tuple[str, ...],
    source_identity: tuple[int, int],
    guard: _DirectoryGuard,
) -> HealthRepairResult:
    snapshot, snapshot_identity, snapshot_digest = _stable_source_copy(
        source_path, guard, source_identity
    )
    if snapshot_digest != plan.source_sha256:
        guard.unlink(snapshot, snapshot_identity)
        raise HealthRepairApplyError("source workbook changed after repair planning")
    try:
        canonical = _canonical_plan(guard.stable_path(snapshot), plan)
        actions = _selected(canonical, selected_action_ids)
        staged, descriptor, staged_identity = guard.create_temp(
            prefix=f".{output_path.stem}-", suffix=output_path.suffix
        )
    except Exception:
        guard.unlink(snapshot, snapshot_identity)
        raise
    published = False
    try:
        snapshot_path = guard.stable_path(snapshot)
        with os.fdopen(descriptor, "w+b") as staged_stream:
            source_book = load_workbook(
                snapshot_path, data_only=False, keep_vba=snapshot.suffix.lower() == ".xlsm"
            )
            try:
                before = _workbook_snapshot(source_book)
                for action in actions:
                    _apply_action(source_book, action)
                staged_stream.seek(0)
                staged_stream.truncate()
                source_book.save(staged_stream)
                staged_stream.flush()
                os.fsync(staged_stream.fileno())
            finally:
                _close_workbook(source_book)
            if guard.identity(staged) != staged_identity:
                raise HealthRepairApplyError("staged output changed during repair")
            with _duplicate_binary_stream(staged_stream.fileno()) as package_stream:
                _verify_xlsm_package(snapshot_path, package_stream)
            original = load_workbook(
                snapshot_path, data_only=False, keep_vba=snapshot.suffix.lower() == ".xlsm"
            )
            with _duplicate_binary_stream(staged_stream.fileno()) as repaired_stream:
                repaired = load_workbook(
                    repaired_stream, data_only=False,
                    keep_vba=staged.suffix.lower() == ".xlsm",
                )
                try:
                    after = _workbook_snapshot(repaired)
                    _verify_unauthorized_unchanged(before, after, actions)
                    _assert_persisted_targets(repaired, actions)
                    changes = _persisted_ledger(original, repaired, actions)
                finally:
                    _close_workbook(repaired)
            _close_workbook(original)
            if _opened_digest(source_path, source_identity) != plan.source_sha256:
                raise HealthRepairApplyError("source workbook changed before repaired copy publication")
            _require_identity(source_path, source_identity, "source workbook")
            if guard.identity(staged) != staged_identity:
                raise HealthRepairApplyError("staged output changed during repair")
            guard.ensure_visible()
            _publish_without_overwrite(guard, staged, output_path)
            published = True
            if guard.identity(output_path) != staged_identity:
                raise HealthRepairApplyError("published output identity does not match staged output")
            guard.ensure_visible()
            output_digest = _descriptor_digest(staged_stream.fileno(), staged_identity)
            return HealthRepairResult(
                source=str(source_path), source_sha256=plan.source_sha256,
                output=str(output_path), output_sha256=output_digest, changes=changes,
            )
    except Exception:
        if published:
            guard.unlink(output_path, staged_identity)
        raise
    finally:
        guard.unlink(staged, staged_identity)
        guard.unlink(snapshot, snapshot_identity)
