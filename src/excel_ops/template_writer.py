"""Safe, declarative writes into existing Excel workbook templates."""

from __future__ import annotations

import json
import re
import shutil
from copy import copy
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .cells import column_number, is_merged_non_anchor, is_protected_formula_value
from .naming import ResolvedOutput, resolve_output_path
from .number_formats import (
    FormatPolicy,
    NumberFormatPolicyError,
    ResolvedFieldFormat,
    policy_manifest,
    resolve_format_policy,
)


class TemplateWriteError(ValueError):
    """Raised when a requested template write would be unsafe."""


@dataclass(frozen=True)
class TemplateMapping:
    """The only workbook cells a template delivery is allowed to change."""

    sheet: str
    field_columns: Mapping[str, str | int]
    header_row: int = 1
    data_start_row: int = 2
    template_type: str = "table"
    style_source_row: int | None = None
    max_rows: int | None = None
    overwrite_formulas: bool = False
    format_policy: FormatPolicy | None = None

    def __post_init__(self) -> None:
        if not self.sheet.strip():
            raise TemplateWriteError("mapping sheet must not be empty")
        if self.header_row < 1 or self.data_start_row <= self.header_row:
            raise TemplateWriteError("data_start_row must be after header_row")
        if not self.field_columns:
            raise TemplateWriteError("field_columns must not be empty")
        if self.max_rows is not None and self.max_rows < 0:
            raise TemplateWriteError("max_rows must be zero or greater")
        columns = [_column_number(value) for value in self.field_columns.values()]
        if len(columns) != len(set(columns)):
            raise TemplateWriteError("field_columns must map to distinct columns")
        if self.format_policy is not None:
            missing = sorted(set(self.format_policy.fields) - set(self.field_columns))
            if missing:
                raise TemplateWriteError(
                    "format policy fields must be mapped columns: " + ", ".join(missing)
                )


@dataclass(frozen=True)
class TemplateChange:
    sheet: str
    cell: str
    field: str
    row_index: int
    previous_value: Any
    new_value: Any


@dataclass(frozen=True)
class TemplateSkip:
    record_index: int
    field: str
    reason: str
    cell: str | None = None


@dataclass(frozen=True)
class TemplateFormatChange:
    sheet: str
    cell: str
    field: str
    row_index: int
    previous_format: str
    new_format: str


@dataclass(frozen=True)
class TemplateWriteResult:
    template_path: Path
    output_path: Path
    change_log_path: Path
    mapping: TemplateMapping
    changes: tuple[TemplateChange, ...] = field(default_factory=tuple)
    skipped: tuple[TemplateSkip, ...] = field(default_factory=tuple)
    naming: ResolvedOutput | None = None
    format_changes: tuple[TemplateFormatChange, ...] = field(default_factory=tuple)
    format_policy: Mapping[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "template_path": str(self.template_path),
            "output_path": str(self.output_path),
            "change_log_path": str(self.change_log_path),
            "template_type": self.mapping.template_type,
            "sheet": self.mapping.sheet,
            "written_cells": len(self.changes),
            "skipped_items": len(self.skipped),
            "changes": [asdict(item) for item in self.changes],
            "skipped": [asdict(item) for item in self.skipped],
            "formatted_cells": len(self.format_changes),
            "format_changes": [asdict(item) for item in self.format_changes],
            "format_policy": dict(self.format_policy),
            "naming": self.naming.to_dict() if self.naming else None,
        }


def write_template(
    template_path: str | Path,
    records: Iterable[Mapping[str, Any] | object],
    mapping: TemplateMapping,
    *,
    output_path: str | Path | None = None,
    output_dir: str | Path | None = None,
    output_pattern: str | None = None,
    period: object | None = None,
    naming_values: Mapping[str, object] | None = None,
) -> TemplateWriteResult:
    """Copy a template, modify mapped cells only, and emit a JSON change log.

    Use either ``output_path`` or the #17-compatible
    ``output_dir``/``output_pattern``/``period`` naming inputs. The returned
    ``output_path`` is the actual file written and is the path downstream
    verification must consume.
    """

    source = Path(template_path).resolve()
    if not source.is_file():
        raise TemplateWriteError(f"template does not exist: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise TemplateWriteError("template must be an .xlsx or .xlsm workbook")

    materialized = list(records)
    try:
        resolved_formats = (
            resolve_format_policy(mapping.format_policy, materialized)
            if mapping.format_policy is not None
            else {}
        )
    except NumberFormatPolicyError as error:
        raise TemplateWriteError(f"{error.code}: {error}") from error

    destination, naming = _destination(
        source,
        output_path=output_path,
        output_dir=output_dir,
        output_pattern=output_pattern,
        period=period,
        naming_values=naming_values,
    )
    if destination.resolve() == source:
        raise TemplateWriteError("output path must not overwrite the source template")
    if destination.exists():
        raise TemplateWriteError(f"output path already exists: {destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)

    keep_vba = source.suffix.lower() == ".xlsm"
    workbook = load_workbook(destination, keep_vba=keep_vba)
    if mapping.sheet not in workbook.sheetnames:
        workbook.close()
        destination.unlink(missing_ok=True)
        raise TemplateWriteError(f"mapped sheet does not exist: {mapping.sheet}")

    worksheet = workbook[mapping.sheet]
    changes: list[TemplateChange] = []
    skipped: list[TemplateSkip] = []
    format_changes: list[TemplateFormatChange] = []
    if mapping.max_rows is not None and len(materialized) > mapping.max_rows:
        workbook.close()
        destination.unlink(missing_ok=True)
        raise TemplateWriteError(
            f"record count {len(materialized)} exceeds mapped max_rows {mapping.max_rows}"
        )

    style_row = mapping.style_source_row or mapping.data_start_row
    for record_index, raw_record in enumerate(materialized):
        values = _record_values(raw_record)
        row = mapping.data_start_row + record_index
        for field_name, column in mapping.field_columns.items():
            column_number = _column_number(column)
            coordinate = f"{get_column_letter(column_number)}{row}"
            if field_name not in values:
                skipped.append(TemplateSkip(record_index, field_name, "missing_field", coordinate))
                continue
            cell = worksheet.cell(row, column_number)
            if is_merged_non_anchor(cell):
                skipped.append(TemplateSkip(record_index, field_name, "merged_non_anchor", coordinate))
                continue
            if is_protected_formula_value(cell.value, overwrite_formulas=mapping.overwrite_formulas):
                skipped.append(TemplateSkip(record_index, field_name, "protected_formula", coordinate))
                continue
            if row != style_row:
                _copy_style(worksheet.cell(style_row, column_number), cell)
            previous = cell.value
            new_value = values[field_name]
            cell.value = new_value
            if field_name in resolved_formats:
                previous_format = cell.number_format
                cell.number_format = resolved_formats[field_name].number_format
                format_changes.append(
                    TemplateFormatChange(
                        mapping.sheet,
                        coordinate,
                        field_name,
                        record_index,
                        previous_format,
                        cell.number_format,
                    )
                )
            changes.append(
                TemplateChange(mapping.sheet, coordinate, field_name, record_index, previous, new_value)
            )

    workbook.save(destination)
    workbook.close()
    try:
        _verify_number_formats(destination, mapping, materialized, resolved_formats)
    except TemplateWriteError:
        destination.unlink(missing_ok=True)
        raise
    result = TemplateWriteResult(
        template_path=source,
        output_path=destination.resolve(),
        change_log_path=destination.with_suffix(destination.suffix + ".changes.json").resolve(),
        mapping=mapping,
        changes=tuple(changes),
        skipped=tuple(skipped),
        format_changes=tuple(format_changes),
        format_policy=policy_manifest(mapping.format_policy, resolved_formats),
        naming=naming,
    )
    _write_change_log(result)
    return result


def _destination(
    source: Path,
    *,
    output_path: str | Path | None,
    output_dir: str | Path | None,
    output_pattern: str | None,
    period: object | None,
    naming_values: Mapping[str, object] | None,
) -> tuple[Path, ResolvedOutput | None]:
    naming_requested = any(value is not None for value in (output_dir, output_pattern, period))
    if output_path is not None and naming_requested:
        raise TemplateWriteError("use output_path or period-aware naming inputs, not both")
    if output_path is not None:
        return Path(output_path), None
    if output_dir is None or period is None:
        raise TemplateWriteError("output_path or output_dir plus period is required")
    pattern = output_pattern or _period_pattern_from_filename(source.stem)
    resolved = resolve_output_path(
        pattern,
        output_dir=output_dir,
        extension=source.suffix,
        period=period,
        values=naming_values,
    )
    return resolved.path, resolved


def _period_pattern_from_filename(stem: str) -> str:
    replaced = re.sub(r"(?<!\d)\d{4}[-_.]\d{2}[-_.]\d{2}(?!\d)", "YYYY-MM-DD", stem)
    replaced = re.sub(r"(?<!\d)\d{8}(?!\d)", "YYYYMMDD", replaced)
    return replaced if replaced != stem else f"{stem}-YYYY-MM-DD"


def _record_values(record: Mapping[str, Any] | object) -> Mapping[str, Any]:
    if isinstance(record, Mapping):
        return record
    try:
        return vars(record)
    except TypeError as exc:
        raise TemplateWriteError("records must be mappings or objects with fields") from exc


def _column_number(value: str | int) -> int:
    """Convert the shared parser's ``ValueError`` into this module's contract."""

    try:
        return column_number(value)
    except ValueError as exc:
        raise TemplateWriteError(str(exc)) from exc


def _copy_style(source: Any, destination: Any) -> None:
    if source.has_style:
        destination._style = copy(source._style)
    if source.number_format:
        destination.number_format = source.number_format


def _verify_number_formats(
    path: Path,
    mapping: TemplateMapping,
    records: list[Mapping[str, Any] | object],
    resolved: Mapping[str, ResolvedFieldFormat],
) -> None:
    """Reopen the persisted workbook and verify values and display formats."""

    if not resolved:
        return
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        worksheet = workbook[mapping.sheet]
        for row_index, raw_record in enumerate(records):
            values = _record_values(raw_record)
            row = mapping.data_start_row + row_index
            for field_name, rule in resolved.items():
                if field_name not in values or values[field_name] is None:
                    continue
                column = _column_number(mapping.field_columns[field_name])
                cell = worksheet.cell(row, column)
                if cell.number_format != rule.number_format:
                    raise TemplateWriteError(
                        f"number_format_verification_failed: {cell.coordinate} persisted "
                        f"{cell.number_format!r}, expected {rule.number_format!r}"
                    )
                if not _same_numeric_value(values[field_name], cell.value):
                    raise TemplateWriteError(
                        f"numeric_value_verification_failed: {cell.coordinate} changed while formatting"
                    )
    finally:
        workbook.close()


def _same_numeric_value(expected: Any, actual: Any) -> bool:
    if isinstance(expected, bool) or isinstance(actual, bool):
        return expected == actual
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        return float(expected) == float(actual)
    return expected == actual


def _write_change_log(result: TemplateWriteResult) -> None:
    payload = result.to_dict()
    payload["created_at_utc"] = datetime.now(timezone.utc).isoformat()
    result.change_log_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8"
    )
