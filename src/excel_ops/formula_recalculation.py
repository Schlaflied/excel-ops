"""Static formula checks plus fail-closed external-engine recalculation."""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable, Literal, Sequence

from openpyxl import load_workbook

from .delivery_verification import VerificationFinding
from .formula_verification import FormulaRegion, FormulaVerifier


RecalculationStatus = Literal["verified", "unverified", "failed"]
_FORMULA_ERRORS = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
}


class FormulaRecalculationError(ValueError):
    """Raised when an available calculation engine fails unexpectedly."""


class FormulaEngineUnavailable(FormulaRecalculationError):
    """Raised when no requested spreadsheet calculation engine is installed."""


@dataclass(frozen=True)
class FormulaValueExpectation:
    sheet: str
    cell: str
    expected: Any
    tolerance: Decimal = Decimal("0")


@dataclass(frozen=True)
class FormulaRecalculationResult:
    status: RecalculationStatus
    engine: str | None
    recalculated_path: Path | None
    findings: tuple[VerificationFinding, ...]

    @property
    def passed(self) -> bool:
        return self.status == "verified"

    def to_dict(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "engine": self.engine,
            "recalculated_path": (
                str(self.recalculated_path) if self.recalculated_path else None
            ),
            "findings": [asdict(item) for item in self.findings],
        }


Recalculator = Callable[[Path, Path], str]


def verify_formula_recalculation(
    source_path: str | Path,
    destination_path: str | Path,
    *,
    expectations: Sequence[FormulaValueExpectation],
    formula_regions: Sequence[FormulaRegion] = (),
    engine: Literal["auto", "excel", "libreoffice"] = "auto",
    recalculator: Recalculator | None = None,
) -> FormulaRecalculationResult:
    """Recalculate a workbook copy and verify declared results independently."""

    source = Path(source_path).resolve()
    destination = Path(destination_path).resolve()
    _validate_paths(source, destination)
    static_findings = _static_findings(source, formula_regions)
    if any(item.severity == "error" for item in static_findings):
        return FormulaRecalculationResult("failed", None, None, static_findings)
    if not expectations:
        finding = VerificationFinding(
            "recalculation_expectations_missing",
            "No independently derived formula expectations were supplied.",
            "Declare sampled or full expected results before claiming recalculation verification.",
            severity="warning",
        )
        return FormulaRecalculationResult(
            "unverified", None, None, (*static_findings, finding)
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    staged = _temporary_path(destination)
    try:
        try:
            used_engine = (recalculator or _engine_recalculator(engine))(source, staged)
        except FormulaEngineUnavailable as error:
            finding = VerificationFinding(
                "recalculation_not_verified",
                str(error),
                "Install Microsoft Excel on Windows or LibreOffice, then rerun verification.",
                severity="warning",
            )
            return FormulaRecalculationResult(
                "unverified", None, None, (*static_findings, finding)
            )
        if not staged.is_file():
            raise FormulaRecalculationError(
                f"{used_engine} reported success without producing a workbook"
            )
        findings = [*static_findings, *_verify_values(staged, expectations)]
        if any(item.severity == "error" for item in findings):
            return FormulaRecalculationResult("failed", used_engine, None, tuple(findings))
        os.replace(staged, destination)
        return FormulaRecalculationResult(
            "verified", used_engine, destination, tuple(findings)
        )
    finally:
        staged.unlink(missing_ok=True)


def _validate_paths(source: Path, destination: Path) -> None:
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        raise FormulaRecalculationError("source must be an existing XLSX workbook")
    if source == destination:
        raise FormulaRecalculationError("source and recalculated paths must be different")
    if destination.suffix.lower() != ".xlsx":
        raise FormulaRecalculationError("recalculated workbook must end in .xlsx")
    if destination.exists():
        raise FormulaRecalculationError("recalculated workbook already exists")


def _static_findings(
    source: Path,
    regions: Sequence[FormulaRegion],
) -> tuple[VerificationFinding, ...]:
    workbook = load_workbook(source, data_only=False)
    try:
        verifier = FormulaVerifier(
            formula_regions=tuple(regions),
            report_recalculation_limit=False,
        )
        return tuple(verifier(workbook, None, None))  # type: ignore[arg-type]
    finally:
        workbook.close()


def _verify_values(
    recalculated: Path,
    expectations: Sequence[FormulaValueExpectation],
) -> tuple[VerificationFinding, ...]:
    workbook = load_workbook(recalculated, data_only=True)
    findings: list[VerificationFinding] = []
    error_cells: set[tuple[str, str]] = set()
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.upper() in _FORMULA_ERRORS:
                        error_cells.add((worksheet.title, cell.coordinate))
                        findings.append(
                            VerificationFinding(
                                "recalculated_formula_error",
                                f"The recalculated cell contains {cell.value}.",
                                "Repair the formula or its inputs before delivery.",
                                sheet=worksheet.title,
                                row=cell.row,
                                cell=cell.coordinate,
                            )
                        )
        for expectation in expectations:
            if expectation.sheet not in workbook.sheetnames:
                findings.append(
                    VerificationFinding(
                        "recalculation_sheet_missing",
                        f"Expected worksheet {expectation.sheet!r} is missing after recalculation.",
                        "Restore the worksheet and rerun recalculation.",
                        sheet=expectation.sheet,
                    )
                )
                continue
            cell = workbook[expectation.sheet][expectation.cell]
            actual = cell.value
            if (expectation.sheet, cell.coordinate) in error_cells:
                continue
            if not _matches(actual, expectation.expected, expectation.tolerance):
                findings.append(
                    VerificationFinding(
                        "recalculation_mismatch",
                        "The recalculated value does not match the independent expectation.",
                        "Review the business rule, inputs, and generated formula.",
                        sheet=expectation.sheet,
                        row=cell.row,
                        cell=cell.coordinate,
                    )
                )
    finally:
        workbook.close()
    return tuple(findings)


def _matches(actual: Any, expected: Any, tolerance: Decimal) -> bool:
    if actual == expected:
        return True
    if isinstance(actual, bool) or isinstance(expected, bool):
        return False
    try:
        return abs(Decimal(str(actual)) - Decimal(str(expected))) <= tolerance
    except (InvalidOperation, TypeError, ValueError):
        return False


def _engine_recalculator(
    engine: Literal["auto", "excel", "libreoffice"],
) -> Recalculator:
    if engine not in {"auto", "excel", "libreoffice"}:
        raise FormulaRecalculationError("engine must be auto, excel, or libreoffice")

    def recalculate(source: Path, destination: Path) -> str:
        if engine in {"auto", "excel"} and os.name == "nt":
            try:
                return _recalculate_with_excel(source, destination)
            except FormulaEngineUnavailable:
                if engine == "excel":
                    raise
        if engine in {"auto", "libreoffice"}:
            return _recalculate_with_libreoffice(source, destination)
        raise FormulaEngineUnavailable("Microsoft Excel recalculation requires Windows")

    return recalculate


_EXCEL_RECALCULATE = r"""
Option Explicit
Dim args, excel, workbook, errorNumber, errorDescription
Set args = WScript.Arguments
If args.Count <> 2 Then WScript.Quit 2
On Error Resume Next
Set excel = CreateObject("Excel.Application")
If Err.Number <> 0 Then WScript.Quit 3
excel.Visible = False
excel.DisplayAlerts = False
Set workbook = excel.Workbooks.Open(args(0), 0, False)
If Err.Number = 0 Then
    excel.CalculateFullRebuild
    workbook.SaveAs args(1), 51
End If
errorNumber = Err.Number
errorDescription = Err.Description
If Not workbook Is Nothing Then workbook.Close False
excel.Quit
If errorNumber <> 0 Then
    WScript.Echo errorDescription
    WScript.Quit 4
End If
""".strip()


def _recalculate_with_excel(source: Path, destination: Path) -> str:
    cscript = shutil.which("cscript.exe")
    if not cscript:
        raise FormulaEngineUnavailable("Windows Script Host is unavailable")
    with tempfile.TemporaryDirectory(prefix="excel-ops-formula-excel-") as name:
        script = Path(name) / "recalculate.vbs"
        script.write_text(_EXCEL_RECALCULATE, encoding="ascii")
        try:
            result = subprocess.run(
                [cscript, "//NoLogo", "//B", str(script), str(source), str(destination)],
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as error:
            raise FormulaRecalculationError(
                f"Microsoft Excel recalculation could not complete: {error}"
            ) from error
    if result.returncode == 3:
        raise FormulaEngineUnavailable("Microsoft Excel is unavailable")
    if result.returncode != 0 or not destination.is_file():
        raise FormulaRecalculationError(
            f"Microsoft Excel recalculation failed with exit code {result.returncode}"
        )
    return "excel"


def _recalculate_with_libreoffice(source: Path, destination: Path) -> str:
    soffice = shutil.which("soffice")
    if not soffice:
        raise FormulaEngineUnavailable("LibreOffice soffice is unavailable")
    with tempfile.TemporaryDirectory(prefix="excel-ops-formula-libreoffice-") as name:
        temporary = Path(name)
        input_dir = temporary / "input"
        output_dir = temporary / "output"
        input_dir.mkdir()
        output_dir.mkdir()
        prepared = input_dir / "recalculate.xlsx"
        shutil.copy2(source, prepared)
        try:
            result = subprocess.run(
                [
                    soffice,
                    f"-env:UserInstallation={(temporary / 'profile').as_uri()}",
                    "--headless",
                    "--convert-to",
                    "xlsx",
                    "--outdir",
                    str(output_dir),
                    str(prepared),
                ],
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as error:
            raise FormulaRecalculationError(
                f"LibreOffice recalculation could not complete: {error}"
            ) from error
        rendered = output_dir / prepared.name
        if result.returncode != 0 or not rendered.is_file():
            raise FormulaRecalculationError(
                f"LibreOffice recalculation failed with exit code {result.returncode}"
            )
        shutil.copy2(rendered, destination)
    return "libreoffice"


def _temporary_path(destination: Path) -> Path:
    descriptor, name = tempfile.mkstemp(
        prefix=f".{destination.stem}-",
        suffix=destination.suffix,
        dir=destination.parent,
    )
    os.close(descriptor)
    path = Path(name)
    path.unlink()
    return path
