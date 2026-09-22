"""Per-delivery source-and-verification Manifest (issue #20).

One delivered workbook gets one Manifest.  A Manifest is an *evidence artifact
generated from a completed run*: it proves which inputs were read, which
template and Recipe version produced the file, how many records reached each
tab, what the file that is actually on disk hashes to, and whether #4's
independent verification passed.  It computes nothing new.

.. rubric:: Read back, never assumed

Every field that describes the persisted artifact is taken from the file on
disk *after* the write and after ``verify_and_deliver`` published it:

* ``output_hash`` hashes the delivered file's bytes with
  :func:`excel_ops.idempotency.file_content_digest`;
* ``template_version`` hashes the template file that was actually used;
* each tab's ``rows`` is counted by reopening the delivered workbook and
  reading its mapped record-ID column.

``rows`` is what makes the Manifest falsifiable.  When the rows really in the
file disagree with the run's ``written`` counters, ``reconciled`` is ``False``
and the mismatch is named in ``discrepancies`` instead of being smoothed over.

.. rubric:: What this is not

* **Not** #19's idempotency fingerprint.  ``.excel-ops/idempotency.json``
  answers "may this run be skipped?" and is keyed by an opaque task digest.  A
  Manifest answers "what did this run actually deliver?" and sits next to the
  delivered file.  The two share hashing helpers (this module reuses
  :mod:`excel_ops.idempotency`'s, rather than hashing anything itself) and
  nothing else; neither file format is readable as the other.
* **Not** a cross-run audit or evidence package (issue #12 and the later
  audit-delivery roadmap item).  A Manifest describes exactly one delivery.
* **Not** a second verification pass.  The ``verification`` block reports the
  result :mod:`excel_ops.delivery_verification` already produced.

.. rubric:: Counts, hashes and metadata only

No cell value is ever copied into a Manifest.  Not a location, not an
identifier, not a category, and not a verification finding's message -- a
finding message may quote the very cell that failed (for example
``written_value_mismatch``), so only finding **codes**, **severities** and
counts are carried over.  Record IDs are not copied either: they are derived
from source values.  What a Manifest holds is file names, file content hashes,
sheet names, declared field names, and integer counts.
"""

from __future__ import annotations

import json
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import TYPE_CHECKING, Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

from .cells import column_number
from .idempotency import file_content_digest, recipe_version_digest
from .periods import PeriodResult

if TYPE_CHECKING:  # pragma: no cover - typing only
    # Imported lazily: ``delivery`` imports this module to attach Manifests to
    # its own result, so a runtime import here would be circular.
    from .delivery import DeliveryRun, DeliveryTarget, RecordOutcome, TargetOutcome


MANIFEST_FORMAT = "excel-ops-delivery-manifest-v1"

#: ``<delivered file>.manifest.json`` / ``.manifest.txt``, following the sibling
#: idiom already used by ``template_writer``'s ``.changes.json`` and
#: ``delivery_verification``'s ``.verification.json``.
MANIFEST_SUFFIX = ".manifest.json"
READABLE_SUFFIX = ".manifest.txt"

PASSED = "passed"
FAILED = "failed"

_STATUSES = ("written", "accepted", "review", "rejected", "skipped_existing")


@dataclass(frozen=True)
class ManifestSource:
    """One input file's contribution, identified by content rather than name."""

    file: str
    path: str
    hash: str
    #: Records ingested from this file across the whole run.
    records: int
    #: Records from this file written into *this* delivered output.
    written: int
    #: This file's run-level record statuses, so the run reconciles per source.
    statuses: Mapping[str, int] = field(default_factory=dict)
    #: The distinct provenance labels the records declared (a sheet-bearing
    #: workbook reports its own name, an extraction JSON reports the image it
    #: came from).  File names only -- never a cell value.
    origins: tuple[str, ...] = ()

    def to_dict(self) -> dict[str, Any]:
        return {
            "file": self.file,
            "path": self.path,
            "hash": self.hash,
            "records": self.records,
            "written": self.written,
            "statuses": dict(self.statuses),
            "origins": list(self.origins),
        }


@dataclass(frozen=True)
class ManifestTab:
    """One written worksheet: where its rows came from, and how many are real."""

    sheet: str
    #: Records the run believes it wrote into this sheet.
    written: int
    #: Rows actually found in the persisted file, counted by reopening it.
    rows: int
    #: Per input file, how many records this sheet received.
    sources: Mapping[str, int] = field(default_factory=dict)

    @property
    def reconciled(self) -> bool:
        return self.written == self.rows

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "written": self.written,
            "rows": self.rows,
            "reconciled": self.reconciled,
            "sources": dict(self.sources),
        }


@dataclass(frozen=True)
class ManifestVerification:
    """#4's verdict for this delivery, reduced to codes and counts."""

    status: str
    findings: int = 0
    codes: tuple[str, ...] = ()
    severities: Mapping[str, int] = field(default_factory=dict)
    report_path: str | None = None

    @property
    def passed(self) -> bool:
        return self.status == PASSED

    def to_dict(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "passed": self.passed,
            "findings": self.findings,
            "codes": list(self.codes),
            "severities": dict(self.severities),
            "report_path": self.report_path,
        }


@dataclass(frozen=True)
class DeliveryManifest:
    """The machine-readable evidence record for one delivered workbook."""

    output: str
    output_name: str
    output_hash: str
    destination_key: str
    template: str
    template_version: str
    recipe_version: str
    verification: ManifestVerification
    sources: tuple[ManifestSource, ...] = ()
    tabs: tuple[ManifestTab, ...] = ()
    period_start: str | None = None
    period_end: str | None = None
    period_display_text: str | None = None
    recipe_path: str | None = None
    #: Run-level record states, so ``input`` reconciles with the five terminal
    #: statuses exactly as :class:`~excel_ops.delivery.DeliveryRun` reports them.
    run_counts: Mapping[str, int] = field(default_factory=dict)
    #: The same five states narrowed to this output's destination.
    target_counts: Mapping[str, int] = field(default_factory=dict)
    #: Codes for anything the Manifest could not honestly claim.
    discrepancies: tuple[str, ...] = ()
    generated_at: str = ""
    #: #19's run fingerprint, when whole-run idempotency was enabled -- a
    #: cross-reference only; the Manifest never reads or writes that record.
    run_fingerprint: str | None = None
    format: str = MANIFEST_FORMAT

    @property
    def accepted(self) -> int:
        return int(self.run_counts.get("accepted", 0))

    @property
    def review(self) -> int:
        return int(self.run_counts.get("review", 0))

    @property
    def rejected(self) -> int:
        return int(self.run_counts.get("rejected", 0))

    @property
    def written(self) -> int:
        """Records written into *this* output."""

        return int(self.target_counts.get("written", 0))

    @property
    def rows(self) -> int:
        """Rows actually present in the persisted output."""

        return sum(item.rows for item in self.tabs)

    @property
    def reconciled(self) -> bool:
        """True when nothing in this Manifest contradicts the real file."""

        return not self.discrepancies

    def to_dict(self) -> dict[str, Any]:
        return {
            "format": self.format,
            "generated_at": self.generated_at,
            "output": self.output,
            "output_name": self.output_name,
            "output_hash": self.output_hash,
            "destination_key": self.destination_key,
            "period_start": self.period_start,
            "period_end": self.period_end,
            "period_display_text": self.period_display_text,
            "template": self.template,
            "template_version": self.template_version,
            "recipe_path": self.recipe_path,
            "recipe_version": self.recipe_version,
            "sources": [item.to_dict() for item in self.sources],
            "accepted": self.accepted,
            "review": self.review,
            "rejected": self.rejected,
            "written": self.written,
            "rows": self.rows,
            "run_counts": dict(self.run_counts),
            "target_counts": dict(self.target_counts),
            "tabs": {item.sheet: item.to_dict() for item in self.tabs},
            "verification": self.verification.to_dict(),
            "reconciled": self.reconciled,
            "discrepancies": list(self.discrepancies),
            "run_fingerprint": self.run_fingerprint,
        }

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, indent=2, default=str) + "\n"

    def manifest_path(self) -> Path:
        output = Path(self.output)
        return output.with_suffix(output.suffix + MANIFEST_SUFFIX)

    def readable_path(self) -> Path:
        output = Path(self.output)
        return output.with_suffix(output.suffix + READABLE_SUFFIX)


# --------------------------------------------------------------------------- #
# Building a Manifest from a completed run
# --------------------------------------------------------------------------- #


def build_delivery_manifests(
    run: "DeliveryRun",
    targets: Sequence["DeliveryTarget"],
    *,
    period: PeriodResult | None = None,
    recipe_decisions: Iterable[Any] | Mapping[str, Any] | None = None,
    record_input_paths: Sequence[str] = (),
    generated_at: str | None = None,
) -> tuple[DeliveryManifest, ...]:
    """Build one Manifest per *delivered* output of a completed run.

    A target that was planned, skipped, blocked or failed verification produced
    no delivered file, so it gets no Manifest: a Manifest exists only for a file
    that is really on disk.

    ``record_input_paths`` is aligned with ``run.records`` -- entry *i* is the
    input path record *i* was ingested from.  It is what lets a source be
    identified by the file the pipeline read rather than by the provenance label
    the record declared (an extraction JSON's records name the *image*, not the
    JSON).  When it is absent, sources fall back to that declared label.
    """

    by_key = {target.key: target for target in targets}
    stamp = generated_at or datetime.now(timezone.utc).isoformat()
    attribution = _attribute(run, record_input_paths)
    manifests: list[DeliveryManifest] = []
    for outcome in run.targets:
        if not outcome.delivered or not outcome.delivery_path:
            continue
        target = by_key.get(outcome.destination_key)
        if target is None:
            continue
        manifests.append(
            _manifest_for(
                run,
                target,
                outcome,
                attribution,
                period=period,
                recipe_decisions=recipe_decisions,
                generated_at=stamp,
            )
        )
    return tuple(manifests)


@dataclass(frozen=True)
class _Attribution:
    """Which input file each record came from, and its run-level statuses."""

    #: Input path per record position in ``run.records``.
    per_record: tuple[str, ...]
    #: Ordered input paths, as the run declared them.
    inputs: tuple[str, ...]
    #: Input path -> status -> count, across the whole run.
    statuses: Mapping[str, Mapping[str, int]]
    #: Input path -> declared provenance labels.
    origins: Mapping[str, tuple[str, ...]]

    def label(self, index: int) -> str:
        if 0 <= index < len(self.per_record):
            return self.per_record[index]
        return ""


def _attribute(run: "DeliveryRun", record_input_paths: Sequence[str]) -> _Attribution:
    records = run.records
    if len(record_input_paths) == len(records):
        per_record = tuple(str(item) for item in record_input_paths)
    else:
        # No aligned attribution was supplied (for example a Manifest rebuilt
        # from a stored run): fall back to the provenance the record declares.
        per_record = tuple(item.source.source_file for item in records)

    declared = tuple(str(item) for item in run.plan.inputs)
    listed = list(declared)
    for label in per_record:
        if label and label not in listed:
            listed.append(label)

    statuses: dict[str, dict[str, int]] = {
        label: {name: 0 for name in _STATUSES} for label in listed
    }
    origins: dict[str, list[str]] = {label: [] for label in listed}
    for index, record in enumerate(records):
        label = per_record[index] if index < len(per_record) else ""
        bucket = statuses.setdefault(label, {name: 0 for name in _STATUSES})
        bucket[record.status] = bucket.get(record.status, 0) + 1
        seen = origins.setdefault(label, [])
        if record.source.source_file and record.source.source_file not in seen:
            seen.append(record.source.source_file)
    return _Attribution(
        per_record,
        tuple(listed),
        {key: dict(value) for key, value in statuses.items()},
        {key: tuple(value) for key, value in origins.items()},
    )


def _manifest_for(
    run: "DeliveryRun",
    target: "DeliveryTarget",
    outcome: "TargetOutcome",
    attribution: _Attribution,
    *,
    period: PeriodResult | None,
    recipe_decisions: Iterable[Any] | Mapping[str, Any] | None,
    generated_at: str,
) -> DeliveryManifest:
    delivered = Path(str(outcome.delivery_path))
    written_positions = [
        index
        for index, item in enumerate(run.records)
        if item.status == "written" and item.destination_key == target.key
    ]
    written_records = [run.records[index] for index in written_positions]

    tabs = _tabs(delivered, target, written_positions, written_records, attribution)
    sources = _sources(attribution, written_positions)
    target_counts = _target_counts(run.records, target.key)
    discrepancies = _discrepancies(delivered, tabs, target_counts)

    return DeliveryManifest(
        output=str(delivered),
        output_name=delivered.name,
        # Hash the bytes that are on disk now, after the write and after
        # verification published the file -- never the in-memory workbook.
        output_hash=file_content_digest(delivered),
        destination_key=target.key,
        template=str(target.template_path),
        template_version=file_content_digest(target.template_path),
        recipe_version=recipe_version_digest(recipe_decisions),
        verification=_verification(outcome),
        sources=sources,
        tabs=tabs,
        period_start=period.period_start.isoformat() if period else None,
        period_end=period.period_end.isoformat() if period else None,
        period_display_text=period.display_text if period else None,
        recipe_path=run.recipe_path,
        run_counts=dict(run.counts),
        target_counts=target_counts,
        discrepancies=discrepancies,
        generated_at=generated_at,
        run_fingerprint=run.fingerprint,
    )


def _target_counts(records: Sequence["RecordOutcome"], key: str) -> dict[str, int]:
    """Narrow the five terminal states to one destination.

    ``written``, ``accepted`` and ``skipped_existing`` are attributed by the
    destination the record was routed to.  ``review`` is attributed by
    *candidacy*: a record held back for review has no destination yet, so the
    honest statement is "this many records could have landed here".  ``rejected``
    records never reached matching and belong to the run, not to a destination.
    """

    counts = {"written": 0, "accepted": 0, "review": 0, "skipped_existing": 0}
    for record in records:
        if record.status == "review":
            if key in record.candidates:
                counts["review"] += 1
        elif record.status in counts and record.destination_key == key:
            counts[record.status] += 1
    return counts


def _tabs(
    delivered: Path,
    target: "DeliveryTarget",
    written_positions: Sequence[int],
    written_records: Sequence["RecordOutcome"],
    attribution: _Attribution,
) -> tuple[ManifestTab, ...]:
    """Per worksheet: the run's written count, and the file's real row count."""

    per_sheet: dict[str, Counter[str]] = {}
    counted: dict[str, set[int]] = {}
    for offset, record in enumerate(written_records):
        label = attribution.label(written_positions[offset])
        sheets = {cell.sheet for cell in record.cells} or {target.mapping.sheet}
        for sheet in sheets:
            per_sheet.setdefault(sheet, Counter())[label] += 1
            counted.setdefault(sheet, set()).add(written_positions[offset])
    if not per_sheet:
        per_sheet[target.mapping.sheet] = Counter()
        counted[target.mapping.sheet] = set()

    real_rows = _persisted_row_counts(delivered, target, tuple(per_sheet))
    return tuple(
        ManifestTab(
            sheet=sheet,
            written=len(counted.get(sheet, ())),
            rows=real_rows.get(sheet, 0),
            sources={key: value for key, value in sorted(contributions.items()) if key},
        )
        for sheet, contributions in sorted(per_sheet.items())
    )


def _persisted_row_counts(
    delivered: Path, target: "DeliveryTarget", sheets: Sequence[str]
) -> dict[str, int]:
    """Count the rows really written in the delivered file, by reopening it.

    A row counts when its mapped record-ID cell holds a non-empty value.  This
    is the one number in a Manifest that comes from the file rather than from
    the run, which is what allows the two to be compared at all.
    """

    column = target.mapping.field_columns.get(target.record_id_field)
    if column is None or not delivered.is_file():
        return {}
    try:
        index = column_number(column)
    except ValueError:
        return {}
    try:
        workbook = load_workbook(delivered, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError):
        return {}
    try:
        counts: dict[str, int] = {}
        for sheet in sheets:
            if sheet not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet]
            total = 0
            for row in worksheet.iter_rows(
                min_row=target.mapping.data_start_row,
                min_col=index,
                max_col=index,
                values_only=True,
            ):
                if str(row[0] or "").strip():
                    total += 1
            counts[sheet] = total
        return counts
    finally:
        workbook.close()


def _sources(
    attribution: _Attribution, written_positions: Sequence[int]
) -> tuple[ManifestSource, ...]:
    written_per_input: Counter[str] = Counter(
        attribution.label(index) for index in written_positions
    )
    sources: list[ManifestSource] = []
    for label in attribution.inputs:
        if not label:
            continue
        statuses = attribution.statuses.get(label, {})
        path = Path(label)
        sources.append(
            ManifestSource(
                file=path.name,
                path=label,
                hash=file_content_digest(path),
                records=sum(statuses.values()),
                written=written_per_input.get(label, 0),
                statuses=statuses,
                origins=attribution.origins.get(label, ()),
            )
        )
    return tuple(sources)


def _verification(outcome: "TargetOutcome") -> ManifestVerification:
    """Report #4's result: codes and counts only, never a finding's message.

    A finding message can quote the cell that failed, so only the machine code
    and the severity cross into the Manifest.  The full report stays in the
    verification JSON that ``verify_and_deliver`` already wrote.
    """

    severities: Counter[str] = Counter(item.severity for item in outcome.findings)
    return ManifestVerification(
        status=PASSED if outcome.delivered else FAILED,
        findings=len(outcome.findings),
        codes=tuple(sorted({item.code for item in outcome.findings})),
        severities=dict(sorted(severities.items())),
        report_path=outcome.verification_report_path,
    )


def _discrepancies(
    delivered: Path, tabs: Sequence[ManifestTab], target_counts: Mapping[str, int]
) -> tuple[str, ...]:
    """Name everything the Manifest cannot honestly claim about the real file."""

    codes: list[str] = []
    if not delivered.is_file():
        codes.append("missing_output")
    for tab in tabs:
        if not tab.reconciled:
            codes.append(f"row_count_mismatch:{tab.sheet}")
    total_rows = sum(tab.rows for tab in tabs)
    if total_rows != target_counts.get("written", 0):
        codes.append("written_count_mismatch")
    return tuple(codes)


# --------------------------------------------------------------------------- #
# Persisting and rendering
# --------------------------------------------------------------------------- #


def write_delivery_manifest(
    manifest: DeliveryManifest,
    *,
    path: str | Path | None = None,
    readable: bool = True,
) -> Path:
    """Write ``<delivered file>.manifest.json`` and return its path.

    ``readable=True`` also writes the human-readable ``.manifest.txt`` beside
    it, so the same evidence can be read without a JSON tool.
    """

    target = Path(path) if path is not None else manifest.manifest_path()
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(manifest.to_json(), encoding="utf-8")
    if readable:
        companion = (
            target.with_suffix(".txt") if path is not None else manifest.readable_path()
        )
        companion.write_text(format_manifest(manifest) + "\n", encoding="utf-8")
    return target


def write_delivery_manifests(
    manifests: Sequence[DeliveryManifest], *, readable: bool = True
) -> tuple[Path, ...]:
    return tuple(write_delivery_manifest(item, readable=readable) for item in manifests)


def load_delivery_manifest(path: str | Path) -> dict[str, Any]:
    """Read a Manifest back, refusing any other format."""

    target = Path(path)
    payload = json.loads(target.read_text(encoding="utf-8"))
    if not isinstance(payload, Mapping) or payload.get("format") != MANIFEST_FORMAT:
        raise ValueError(f"{target.name}: unsupported delivery-manifest format")
    return dict(payload)


def format_manifest(manifest: DeliveryManifest) -> str:
    """Render one Manifest for a human reviewer.

    Mirrors ``workdir.format_dry_run`` and ``idempotency.format_decision``: the
    JSON stays the machine contract, this is the same facts as readable lines.
    """

    period = (
        f"{manifest.period_start} .. {manifest.period_end}"
        if manifest.period_start or manifest.period_end
        else "not declared"
    )
    lines = [
        f"delivery manifest: {manifest.output_name}",
        f"  output           {manifest.output}",
        f"  output hash      {manifest.output_hash}",
        f"  destination      {manifest.destination_key}",
        f"  period           {period}",
        f"  template         {Path(manifest.template).name} ({manifest.template_version})",
        f"  recipe           {manifest.recipe_path or 'none'} ({manifest.recipe_version})",
        f"  verification     {manifest.verification.status}"
        f" ({manifest.verification.findings} finding(s))",
    ]
    if manifest.verification.codes:
        lines.append(f"  finding codes    {', '.join(manifest.verification.codes)}")
    lines.append(
        "  run counts       "
        + ", ".join(f"{name}={manifest.run_counts.get(name, 0)}" for name in ("input", *_STATUSES))
    )
    lines.append(
        f"  this output      written={manifest.written}, rows in file={manifest.rows}, "
        f"review candidates={manifest.target_counts.get('review', 0)}, "
        f"skipped as existing={manifest.target_counts.get('skipped_existing', 0)}"
    )
    lines.append("  sources")
    for source in manifest.sources:
        lines.append(
            f"    - {source.file}: {source.records} record(s), "
            f"{source.written} written here, {source.hash}"
        )
    lines.append("  tabs")
    for tab in manifest.tabs:
        contributions = ", ".join(
            f"{Path(name).name}={count}" for name, count in tab.sources.items()
        )
        lines.append(
            f"    - {tab.sheet}: written={tab.written}, rows in file={tab.rows}"
            + (f", from {contributions}" if contributions else "")
        )
    if manifest.discrepancies:
        lines.append(f"  DISCREPANCIES    {', '.join(manifest.discrepancies)}")
    else:
        lines.append("  reconciled       yes (run counters match the persisted file)")
    lines.append(
        "  note             counts, hashes and metadata only; no cell value is recorded here"
    )
    return "\n".join(lines)


def format_manifests(manifests: Sequence[DeliveryManifest]) -> str:
    if not manifests:
        return "delivery manifest: nothing was delivered, so no manifest was generated"
    return "\n\n".join(format_manifest(item) for item in manifests)
