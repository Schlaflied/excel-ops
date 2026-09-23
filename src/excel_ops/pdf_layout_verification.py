"""Conservative page-level verification for rendered workbook PDFs."""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Sequence

from openpyxl import load_workbook
from pypdf import PdfReader


MAX_CONTENT_STREAM_BYTES = 50 * 1024 * 1024


@dataclass(frozen=True)
class PdfLayoutFinding:
    """One machine-readable reason a PDF layout could not be verified."""

    code: str
    message: str
    sheet: str | None = None
    page: int | None = None

    def to_dict(self) -> dict[str, Any]:
        """Return JSON-safe evidence without workbook cell values."""

        return asdict(self)


@dataclass(frozen=True)
class PdfLayoutVerification:
    """Page-level verification result for one selected-sheet PDF."""

    passed: bool
    pages: int
    findings: tuple[PdfLayoutFinding, ...]


def verify_pdf_layout(
    workbook_path: str | Path,
    pdf_path: str | Path,
    *,
    sheets: Sequence[str],
) -> PdfLayoutVerification:
    """Check pagination, page bounds, readable text, headers, and edge content."""

    findings: list[PdfLayoutFinding] = []
    try:
        reader = PdfReader(str(pdf_path))
        pages = tuple(reader.pages)
    except Exception as error:
        return PdfLayoutVerification(
            False,
            0,
            (PdfLayoutFinding("pdf_unreadable", f"PDF could not be parsed: {error}"),),
        )

    if len(pages) < len(sheets):
        findings.append(
            PdfLayoutFinding(
                "page_count_too_small",
                "The PDF has fewer pages than selected worksheets.",
            )
        )

    page_text: list[str] = []
    for index, page in enumerate(pages, start=1):
        try:
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)
        except Exception as error:
            findings.append(
                PdfLayoutFinding(
                    "invalid_page_bounds",
                    f"The rendered page dimensions could not be read: {error}",
                    page=index,
                )
            )
        else:
            if width <= 0 or height <= 0 or width > 2000 or height > 2000:
                findings.append(
                    PdfLayoutFinding(
                        "invalid_page_bounds",
                        "The rendered page has invalid or unreasonable dimensions.",
                        page=index,
                    )
                )
        try:
            contents = page.get_contents()
            content_size = len(contents.get_data()) if contents is not None else 0
            if content_size > MAX_CONTENT_STREAM_BYTES:
                findings.append(
                    PdfLayoutFinding(
                        "content_stream_too_large",
                        "The page content stream exceeds the safe text-extraction budget.",
                        page=index,
                    )
                )
                page_text.append("")
                continue
            text = page.extract_text() or ""
        except Exception as error:
            findings.append(
                PdfLayoutFinding(
                    "unreadable_page",
                    f"Text could not be extracted from the rendered page: {error}",
                    page=index,
                )
            )
            page_text.append("")
            continue
        page_text.append(_normalize(text))
        if not text.strip():
            findings.append(
                PdfLayoutFinding(
                    "unreadable_page",
                    "No readable text could be extracted from the rendered page.",
                    page=index,
                )
            )

    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        rendered_text = " ".join(page_text)
        for sheet_name in sheets:
            worksheet = workbook[sheet_name]
            if worksheet.max_row > 50 and not worksheet.print_title_rows:
                findings.append(
                    PdfLayoutFinding(
                        "repeated_header_not_configured",
                        "A long worksheet has no repeated print-title rows configured.",
                        sheet=sheet_name,
                    )
                )
            if not _has_safe_scaling(worksheet):
                findings.append(
                    PdfLayoutFinding(
                        "scaling_not_bounded",
                        "Worksheet print scaling is not bounded to a supported range.",
                        sheet=sheet_name,
                    )
                )
            boundary_tokens, boundary_unverifiable = _boundary_tokens(worksheet)
            if boundary_unverifiable:
                findings.append(
                    PdfLayoutFinding(
                        "boundary_content_unverifiable",
                        "A printable boundary cell is not literal text and cannot be matched safely.",
                        sheet=sheet_name,
                    )
                )
            for token in boundary_tokens:
                if token not in rendered_text:
                    findings.append(
                        PdfLayoutFinding(
                            "boundary_content_missing",
                            "Printable boundary content was not found in the rendered PDF.",
                            sheet=sheet_name,
                        )
                    )
                    break
            if len(sheets) == 1 and len(pages) > 1 and worksheet.print_title_rows:
                title_tokens = _title_tokens(worksheet)
                if title_tokens and any(
                    any(token not in text for token in title_tokens) for text in page_text
                ):
                    findings.append(
                        PdfLayoutFinding(
                            "repeated_header_missing",
                            "Configured print-title content was not found on every page.",
                            sheet=sheet_name,
                        )
                    )
    finally:
        workbook.close()

    return PdfLayoutVerification(not findings, len(pages), tuple(findings))


def _has_safe_scaling(worksheet) -> bool:
    """Accept fit-to-page or an explicit, bounded percentage scale."""

    setup = worksheet.page_setup
    if setup.fitToWidth is not None:
        return 1 <= int(setup.fitToWidth) <= 2
    if setup.scale is not None:
        return 10 <= int(setup.scale) <= 400
    return True


def _boundary_tokens(worksheet) -> tuple[tuple[str, ...], bool]:
    """Return literal edge tokens and whether a non-text edge is unverifiable."""

    first = None
    last = None
    for row in worksheet.iter_rows():
        populated = [cell for cell in row if cell.value not in (None, "")]
        if populated:
            first = first or populated[0]
            last = populated[-1]
    edges = (first, last)
    tokens = tuple(
        dict.fromkeys(
            token
            for cell in edges
            if cell is not None and cell.data_type in {"s", "inlineStr"}
            for token in (_normalize(str(cell.value)),)
            if len(token) >= 3
        )
    )
    unverifiable = any(
        cell is not None and cell.data_type not in {"s", "inlineStr"}
        for cell in edges
    )
    return tokens, unverifiable


def _normalize(value: str) -> str:
    """Normalize renderer whitespace for conservative text-presence checks."""

    return re.sub(r"\s+", " ", value).strip().casefold()


def _title_tokens(worksheet) -> tuple[str, ...]:
    """Return stable text tokens from configured repeated title rows."""

    title_range = worksheet.print_title_rows
    if not title_range:
        return ()
    bounds = title_range.replace("$", "").split(":", maxsplit=1)
    start, end = (int(bounds[0]), int(bounds[-1]))
    tokens = []
    for row in worksheet.iter_rows(min_row=start, max_row=end):
        for cell in row:
            token = _normalize(str(cell.value or ""))
            if len(token) >= 3:
                tokens.append(token)
    return tuple(dict.fromkeys(tokens))
