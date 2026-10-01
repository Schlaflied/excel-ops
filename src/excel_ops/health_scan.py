"""Read-only XLSX/XLSM health scanning.

The scanner never saves a workbook. It verifies the source digest before and
after inspection so a completed result is evidence about one unchanged file.
"""

from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
import unicodedata
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from .health_models import (
    HealthFinding,
    HealthFixability,
    HealthScanError,
    HealthScanResult,
    HealthSeverity,
)
from .health_scan_extended import extended_sheet_findings

COVERAGE = (
    "abnormal_merged_cells",
    "blank_data_rows",
    "blank_zero_scan_budget_exceeded",
    "blank_zero_semantics",
    "content_beyond_working_area",
    "default_empty_sheet",
    "duplicate_headers",
    "extra_whitespace",
    "filter_range_mismatch",
    "freeze_panes_mismatch",
    "inconsistent_cell_style",
    "invisible_characters",
    "mixed_date_formats",
    "numeric_text",
    "print_area_mismatch",
)
_NUMERIC_TEXT = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?%?$")
_DEFAULT_SHEET = re.compile(r"^(?:sheet|工作表)\d*$", re.IGNORECASE)
_ROW_GAP = 50
_COLUMN_GAP = 10


def _digest(path: Path) -> str:
    checksum = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            checksum.update(chunk)
    return checksum.hexdigest()


def _populated_cells(sheet: Any) -> list[Any]:
    # A normal-mode worksheet exposes instantiated cells without expanding a
    # malicious or accidental A1:XFD1048576 dimension into a full grid scan.
    cells = getattr(sheet, "_cells", {}).values()
    return sorted(
        (
            cell
            for cell in cells
            if not isinstance(cell, MergedCell) and cell.value not in (None, "")
        ),
        key=lambda cell: (cell.row, cell.column),
    )


def _header_row(cells: list[Any]) -> int | None:
    by_row: dict[int, list[Any]] = defaultdict(list)
    for cell in cells:
        by_row[cell.row].append(cell)
    for row in sorted(by_row):
        values = by_row[row]
        if sum(isinstance(cell.value, str) for cell in values) >= 2:
            return row
    return min(by_row) if by_row else None


def _primary_limit(indices: Iterable[int], gap: int) -> int | None:
    ordered = sorted(set(indices))
    if not ordered:
        return None
    for previous, current in zip(ordered, ordered[1:]):
        if current - previous > gap:
            return previous
    return ordered[-1]


def _finding(
    code: str,
    sheet: str,
    location: str,
    evidence: str,
    suggestion: str,
    *,
    severity: HealthSeverity = "warning",
    fixability: HealthFixability = "review",
    details: Mapping[str, Any] | None = None,
) -> HealthFinding:
    return HealthFinding(
        code=code,
        severity=severity,
        sheet=sheet,
        location=location,
        evidence=evidence,
        suggestion=suggestion,
        fixability=fixability,
        details=dict(details or {}),
    )


def _text_findings(sheet: str, cells: list[Any], header_row: int | None) -> list[HealthFinding]:
    findings: list[HealthFinding] = []
    for cell in cells:
        value = cell.value
        if not isinstance(value, str) or value.startswith("="):
            continue
        invisible = [char for char in value if _is_invisible(char)]
        if invisible:
            codes = ", ".join(sorted({f"U+{ord(char):04X}" for char in invisible}))
            findings.append(_finding(
                "invisible_characters", sheet, cell.coordinate,
                f"Value contains invisible character(s): {codes}.",
                "Review the source meaning before removing invisible characters.",
                details={
                    "current_value": value,
                    "candidate_value": "".join(char for char in value if not _is_invisible(char)),
                    "removed_codepoints": sorted({f"U+{ord(char):04X}" for char in invisible}),
                },
            ))
        leading = len(value) - len(value.lstrip())
        trailing = len(value) - len(value.rstrip())
        repeated = bool(re.search(r"(?<=\S) {2,}(?=\S)", value))
        if leading or trailing or repeated:
            findings.append(_finding(
                "extra_whitespace", sheet, cell.coordinate,
                f"Whitespace profile: leading={leading}, trailing={trailing}, repeated_internal={repeated}.",
                "Review the value before normalizing whitespace.",
                details={
                    "current_value": value,
                    "candidate_value": re.sub(r"(?<=\S) {2,}(?=\S)", " ", value.strip()),
                },
            ))
        stripped = value.strip()
        if cell.row != header_row and _looks_like_numeric_text(stripped):
            candidate = _numeric_candidate(stripped)
            details: dict[str, Any] = {"current_value": value}
            if candidate is None:
                details.update({
                    "required_baseline": "target_representation",
                    "allowed_representations": ["preserve_text"],
                    "precision_reason": "exceeds_excel_15_significant_digits",
                })
            else:
                details["candidate_value"] = candidate
            findings.append(_finding(
                "numeric_text", sheet, cell.coordinate,
                "A numeric-looking value is stored as text.",
                "Confirm that the value is not an identifier before converting its type.",
                details=details,
            ))
    return findings


def _is_invisible(char: str) -> bool:
    return char == "\u00a0" or (unicodedata.category(char) in {"Cc", "Cf"} and char not in "\t\n\r")


def _looks_like_numeric_text(value: str) -> bool:
    if not _NUMERIC_TEXT.fullmatch(value):
        return False
    unsigned = value.lstrip("+-").replace(",", "").rstrip("%")
    integer = unsigned.split(".", 1)[0]
    return not (len(integer) > 1 and integer.startswith("0"))


def _significant_digits(value: str) -> int:
    digits = value.lstrip("+-").replace(",", "").rstrip("%").replace(".", "")
    return len(digits.lstrip("0") or "0")


def _canonical_decimal(value: Decimal) -> str:
    rendered = format(value, "f")
    if "." in rendered:
        rendered = rendered.rstrip("0").rstrip(".")
    return rendered or "0"


def _numeric_candidate(value: str) -> int | dict[str, str] | None:
    if _significant_digits(value) > 15:
        return None
    cleaned = value.replace(",", "")
    percent = cleaned.endswith("%")
    numeric = Decimal(cleaned.rstrip("%"))
    if percent:
        numeric /= Decimal(100)
    if not percent and "." not in cleaned:
        return int(numeric)
    return {"kind": "decimal", "canonical": _canonical_decimal(numeric)}


def _duplicate_header_findings(sheet: str, cells: list[Any], header_row: int | None) -> list[HealthFinding]:
    if header_row is None:
        return []
    headers = [
        cell
        for cell in cells
        if cell.row == header_row and isinstance(cell.value, str) and cell.value.strip()
    ]
    catalog = [
        {
            "coordinate": cell.coordinate,
            "value": cell.value,
            "normalized": " ".join(cell.value.split()).casefold(),
        }
        for cell in headers
    ]
    seen: dict[str, str] = {}
    findings: list[HealthFinding] = []
    for cell in headers:
        normalized = " ".join(cell.value.split()).casefold()
        if normalized in seen:
            findings.append(_finding(
                "duplicate_header", sheet, cell.coordinate,
                f"Header duplicates {seen[normalized]} after whitespace and case normalization.",
                "Rename or disambiguate the duplicate field before using the table.",
                details={
                    "current_header": cell.value,
                    "current_normalized": normalized,
                    "duplicates_coordinate": seen[normalized],
                    "header_row": header_row,
                    "headers": catalog,
                    "required_baseline": "target_header",
                },
            ))
        else:
            seen[normalized] = cell.coordinate
    return findings


def _blank_row_findings(
    sheet: str,
    cells: list[Any],
    header_row: int | None,
    max_row: int | None,
    max_column: int | None,
) -> list[HealthFinding]:
    if header_row is None or max_row is None or max_column is None or max_row <= header_row + 1:
        return []
    occupied = {cell.row for cell in cells if cell.row <= max_row and cell.column <= max_column}
    blank = [row for row in range(header_row + 1, max_row + 1) if row not in occupied]
    findings: list[HealthFinding] = []
    for start, end in _contiguous_ranges(blank):
        location = f"A{start}:{get_column_letter(max_column)}{end}"
        findings.append(_finding(
            "blank_data_rows", sheet, location,
            f"Rows {start}-{end} are blank inside the detected data region.",
            "Review whether the blank rows should remain as intentional separators.",
            details={"blank_rows": list(range(start, end + 1)), "candidate_operation": "remove_rows"},
        ))
    return findings


def _contiguous_ranges(values: list[int]) -> Iterable[tuple[int, int]]:
    if not values:
        return
    start = previous = values[0]
    for value in values[1:]:
        if value != previous + 1:
            yield start, previous
            start = value
        previous = value
    yield start, previous


def _outlier_findings(
    sheet: str, cells: list[Any], max_row: int | None, max_column: int | None
) -> list[HealthFinding]:
    if max_row is None or max_column is None:
        return []
    return [
        _finding(
            "content_beyond_working_area", sheet, cell.coordinate,
            f"Content lies beyond the primary area ending at {get_column_letter(max_column)}{max_row}.",
            "Confirm whether this isolated content belongs to the workbook's active table.",
            fixability="unsupported",
            details={"coordinate": cell.coordinate, "value_type": type(cell.value).__name__},
        )
        for cell in cells
        if cell.row > max_row or cell.column > max_column
    ]


def _merged_findings(sheet: Any, header_row: int | None) -> list[HealthFinding]:
    findings: list[HealthFinding] = []
    for merged in sorted(sheet.merged_cells.ranges, key=lambda item: str(item)):
        area = (merged.max_row - merged.min_row + 1) * (merged.max_col - merged.min_col + 1)
        enters_data = header_row is not None and merged.max_row > header_row
        if enters_data or area > 100:
            findings.append(_finding(
                "abnormal_merged_cells", sheet.title, str(merged),
                "Merged range intersects the detected data region." if enters_data else f"Merged range spans {area} cells.",
                "Review the merge because it can make row and column meaning ambiguous.",
                fixability="unsupported",
                details={"merged_range": str(merged), "candidate_operation": "unmerge"},
            ))
    return findings


def _sheet_findings(sheet: Any) -> list[HealthFinding]:
    cells = _populated_cells(sheet)
    if not cells:
        if _DEFAULT_SHEET.fullmatch(sheet.title.strip()):
            return [_finding(
                "default_empty_sheet", sheet.title, "A1",
                "The default-named worksheet contains no values.",
                "Confirm whether the empty default worksheet is intentional.",
                severity="info", fixability="review",
                details={"sheet": sheet.title, "candidate_operation": "remove_sheet"},
            )]
        return []
    header_row = _header_row(cells)
    max_row = _primary_limit((cell.row for cell in cells), _ROW_GAP)
    max_column = _primary_limit((cell.column for cell in cells), _COLUMN_GAP)
    findings = _text_findings(sheet.title, cells, header_row)
    findings.extend(_duplicate_header_findings(sheet.title, cells, header_row))
    findings.extend(_blank_row_findings(sheet.title, cells, header_row, max_row, max_column))
    findings.extend(_outlier_findings(sheet.title, cells, max_row, max_column))
    findings.extend(_merged_findings(sheet, header_row))
    findings.extend(extended_sheet_findings(sheet, header_row, max_row, max_column))
    return findings


def scan_workbook_health(source: str | Path) -> HealthScanResult:
    """Inspect one XLSX/XLSM without saving it or executing embedded code."""

    path = Path(source)
    if not path.is_file():
        raise HealthScanError(f"workbook does not exist: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise HealthScanError("health scan supports only .xlsx and .xlsm workbooks")
    try:
        snapshot = path.read_bytes()
    except OSError as exc:
        raise HealthScanError(f"workbook could not be read: {exc}") from exc
    before_size = len(snapshot)
    before = sha256(snapshot).hexdigest()
    try:
        workbook = load_workbook(
            BytesIO(snapshot),
            read_only=False,
            data_only=False,
            keep_vba=path.suffix.lower() == ".xlsm",
            keep_links=False,
        )
    except Exception as exc:
        raise HealthScanError(f"workbook could not be scanned: {exc}") from exc
    try:
        findings = [item for sheet in workbook.worksheets for item in _sheet_findings(sheet)]
        if not any(_populated_cells(sheet) for sheet in workbook.worksheets):
            first_sheet = workbook.sheetnames[0] if workbook.sheetnames else "(workbook)"
            findings.append(_finding(
                "empty_workbook", first_sheet, "A1",
                "No populated cells were found in any worksheet.",
                "Confirm that an empty workbook is the intended input.",
                severity="warning", fixability="unsupported",
                details={"sheet_count": len(workbook.sheetnames)},
            ))
        sheets = tuple(workbook.sheetnames)
    finally:
        workbook.close()
        vba_archive = getattr(workbook, "vba_archive", None)
        if vba_archive is not None:
            vba_archive.close()
    after = _digest(path)
    after_size = path.stat().st_size
    if after != before or after_size != before_size:
        raise HealthScanError("source_changed_during_scan")
    ordered = tuple(sorted(
        findings,
        key=lambda item: (item.sheet.casefold(), item.location, item.code, item.evidence),
    ))
    return HealthScanResult(
        source=str(path),
        source_sha256=before,
        source_size=before_size,
        sheets=sheets,
        coverage=COVERAGE,
        findings=ordered,
    )
