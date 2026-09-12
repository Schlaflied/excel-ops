from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import ReviewedRecord

HEADERS = ["Location", "Event date", "Identifier", "Category", "Confidence", "Source", "Review reasons"]


def _append_record(ws, reviewed: ReviewedRecord) -> None:
    record = reviewed.record
    ws.append([
        record.location,
        record.event_date,
        record.identifier,
        record.category,
        record.confidence,
        record.source,
        "; ".join(reviewed.reasons),
    ])


def write_workbook(records: list[ReviewedRecord], output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    accepted = wb.active
    accepted.title = "Accepted"
    review = wb.create_sheet("Review")
    audit = wb.create_sheet("Audit")

    for ws in (accepted, review):
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:G1"

    for reviewed in records:
        _append_record(accepted if reviewed.status == "accepted" else review, reviewed)

    accepted_count = sum(item.status == "accepted" for item in records)
    review_count = len(records) - accepted_count
    audit.append(["Metric", "Value"])
    audit.append(["Processed at (UTC)", datetime.now(timezone.utc).isoformat()])
    audit.append(["Total records", len(records)])
    audit.append(["Accepted records", accepted_count])
    audit.append(["Review records", review_count])
    for cell in audit[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="548235")

    for ws in wb.worksheets:
        for column in ws.columns:
            width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            ws.column_dimensions[column[0].column_letter].width = width

    wb.save(output)
    return output


def verify_workbook(path: str | Path, expected_total: int) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    required = {"Accepted", "Review", "Audit"}
    if not required.issubset(wb.sheetnames):
        raise ValueError("Workbook is missing required delivery sheets")
    actual_total = max(0, wb["Accepted"].max_row - 1) + max(0, wb["Review"].max_row - 1)
    wb.close()
    if actual_total != expected_total:
        raise ValueError(f"Workbook count mismatch: expected {expected_total}, got {actual_total}")
