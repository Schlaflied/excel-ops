"""Native PDF export through a locally installed Microsoft Excel instance."""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
import ctypes
from pathlib import Path
from typing import Callable, Sequence

from openpyxl import load_workbook


class ExcelPdfExportError(ValueError):
    """Raised when native Microsoft Excel PDF export cannot complete safely."""


_VBSCRIPT_EXPORT = r"""
Option Explicit
Dim args, excel, workbook, names(), index, exportError, exportDescription
Set args = WScript.Arguments
If args.Count < 3 Then WScript.Quit 2
On Error Resume Next
Set excel = CreateObject("Excel.Application")
If Err.Number <> 0 Then
    WScript.Echo "Excel unavailable: " & Err.Description
    WScript.Quit 3
End If
excel.Visible = False
excel.DisplayAlerts = False
WScript.StdOut.WriteLine CStr(excel.Hwnd)
Set workbook = excel.Workbooks.Open(args(0), 0, True)
If Err.Number <> 0 Then
    exportError = Err.Number
    exportDescription = Err.Description
Else
    ReDim names(args.Count - 3)
    For index = 2 To args.Count - 1
        names(index - 2) = args(index)
    Next
    Err.Clear
    If args.Count = 3 Then
        workbook.Sheets.Item(names(0)).ExportAsFixedFormat 0, args(1)
    Else
        workbook.Sheets(names).Select
        excel.ActiveSheet.ExportAsFixedFormat 0, args(1)
    End If
    exportError = Err.Number
    exportDescription = Err.Description
    Err.Clear
    workbook.Close False
End If
excel.Quit
If exportError <> 0 Then
    WScript.Echo "Excel export failed: " & exportDescription
    WScript.Quit 4
End If
""".strip()


def export_pdf_with_excel(
    source: str | Path,
    destination: str | Path,
    *,
    sheets: Sequence[str] | None = None,
    cscript: str | Path | None = None,
    popen: Callable[..., subprocess.Popen[str]] = subprocess.Popen,
) -> Path:
    """Export selected sheets with native Excel on Windows and verify the PDF."""

    if not _is_windows():
        raise ExcelPdfExportError("Microsoft Excel PDF export requires Windows")
    source_path = Path(source).resolve()
    destination_path = Path(destination).resolve()
    if source_path.suffix.lower() != ".xlsx":
        raise ExcelPdfExportError("Microsoft Excel PDF export requires an XLSX source")
    if destination_path.suffix.lower() != ".pdf":
        raise ExcelPdfExportError("Microsoft Excel PDF destination must end in .pdf")
    if not source_path.is_file():
        raise ExcelPdfExportError(f"source workbook does not exist: {source_path}")
    executable = str(cscript) if cscript is not None else _find_cscript()
    if not executable:
        raise ExcelPdfExportError("Windows Script Host is unavailable")
    selected = _validated_sheets(source_path, sheets)

    destination_path.parent.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(
        dir=destination_path.parent,
        prefix=f".{destination_path.stem}-",
        suffix=".pdf",
        delete=False,
    )
    staged = Path(handle.name)
    handle.close()
    staged.unlink(missing_ok=True)
    try:
        with tempfile.TemporaryDirectory(prefix="excel-ops-ms-excel-") as temporary_name:
            script_path = Path(temporary_name) / "export.vbs"
            script_path.write_text(_VBSCRIPT_EXPORT, encoding="ascii")
            process = None
            try:
                process = popen(
                    [executable, "//NoLogo", "//B", str(script_path), str(source_path),
                     str(staged), *selected],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                )
                if process.stdout is None:
                    raise ExcelPdfExportError("Windows Script Host stdout is unavailable")
                window_line = process.stdout.readline().strip()
                try:
                    excel_window = int(window_line)
                except ValueError as error:
                    process.communicate()
                    raise ExcelPdfExportError(
                        "Microsoft Excel did not provide a process window handle"
                    ) from error
                try:
                    process.communicate(timeout=120)
                except subprocess.TimeoutExpired as error:
                    process.kill()
                    process.communicate()
                    cleaned = _terminate_excel_window(excel_window)
                    detail = "terminated" if cleaned else "cleanup could not be confirmed"
                    raise ExcelPdfExportError(
                        f"Microsoft Excel PDF export timed out; Excel process {detail}"
                    ) from error
            except OSError as error:
                raise ExcelPdfExportError(f"Microsoft Excel could not run: {error}") from error
            if process.returncode != 0:
                raise ExcelPdfExportError(
                    f"Microsoft Excel PDF export failed with exit code {process.returncode}"
                )
        if not staged.is_file() or staged.stat().st_size < 5:
            raise ExcelPdfExportError("Microsoft Excel produced no usable PDF")
        with staged.open("rb") as pdf:
            if pdf.read(5) != b"%PDF-":
                raise ExcelPdfExportError("Microsoft Excel output is not a PDF")
        staged.replace(destination_path)
        return destination_path
    finally:
        staged.unlink(missing_ok=True)


def _validated_sheets(source: Path, sheets: Sequence[str] | None) -> tuple[str, ...]:
    """Resolve an explicit, duplicate-free sheet selection."""

    workbook = load_workbook(source, read_only=True)
    try:
        visible = tuple(
            name
            for name in workbook.sheetnames
            if getattr(workbook[name], "sheet_state", "visible") == "visible"
        )
        selected = tuple(visible if sheets is None else sheets)
        if not selected:
            raise ExcelPdfExportError("Microsoft Excel PDF export requires at least one sheet")
        if len(set(selected)) != len(selected):
            raise ExcelPdfExportError("Microsoft Excel PDF sheet selection contains duplicates")
        missing = [name for name in selected if name not in workbook.sheetnames]
        if missing:
            raise ExcelPdfExportError(f"unknown worksheet(s): {', '.join(missing)}")
        hidden = [name for name in selected if name not in visible]
        if hidden:
            raise ExcelPdfExportError(
                f"hidden worksheet(s) cannot be exported: {', '.join(hidden)}"
            )
        return selected
    finally:
        workbook.close()


def _find_cscript() -> str | None:
    """Locate the Windows Script Host command-line executable."""

    return shutil.which("cscript.exe")


def _terminate_excel_window(window: int) -> bool:
    """Terminate only the Excel process associated with the emitted window handle."""

    process_id = ctypes.c_ulong()
    if not ctypes.windll.user32.GetWindowThreadProcessId(
        ctypes.c_void_p(window), ctypes.byref(process_id)
    ):
        return False
    process = ctypes.windll.kernel32.OpenProcess(0x0001, False, process_id.value)
    if not process:
        return False
    try:
        return bool(ctypes.windll.kernel32.TerminateProcess(process, 1))
    finally:
        ctypes.windll.kernel32.CloseHandle(process)


def _is_windows() -> bool:
    """Return whether native Microsoft Excel automation is available in principle."""

    return os.name == "nt"
