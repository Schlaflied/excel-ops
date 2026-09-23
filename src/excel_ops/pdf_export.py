"""Bounded PDF export through a locally installed LibreOffice renderer."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Callable, Sequence

from openpyxl import load_workbook


class PdfExportError(ValueError):
    """Raised when PDF export cannot be performed or verified safely."""


def export_pdf(
    source: str | Path,
    destination: str | Path,
    *,
    sheets: Sequence[str] | None = None,
    soffice: str | Path | None = None,
    runner: Callable[..., subprocess.CompletedProcess[str]] = subprocess.run,
) -> Path:
    """Render selected workbook sheets to one verified PDF via LibreOffice."""

    source_path = Path(source)
    destination_path = Path(destination)
    if source_path.suffix.lower() != ".xlsx":
        raise PdfExportError("PDF export requires an XLSX source")
    if destination_path.suffix.lower() != ".pdf":
        raise PdfExportError("PDF export destination must end in .pdf")
    if not source_path.is_file():
        raise PdfExportError(f"source workbook does not exist: {source_path}")
    executable = str(soffice) if soffice is not None else shutil.which("soffice")
    if not executable:
        raise PdfExportError("PDF renderer unavailable: install LibreOffice soffice")

    selected = _validated_sheets(source_path, sheets)
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="excel-ops-pdf-") as temporary_name:
        temporary = Path(temporary_name)
        prepared = temporary / "prepared.xlsx"
        _prepare_workbook(source_path, prepared, selected)
        try:
            result = runner(
                [
                    executable,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(temporary),
                    str(prepared),
                ],
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as error:
            raise PdfExportError(f"PDF renderer could not run: {error}") from error
        if result.returncode != 0:
            raise PdfExportError(f"PDF renderer failed with exit code {result.returncode}")
        rendered = temporary / "prepared.pdf"
        if not rendered.is_file() or rendered.stat().st_size < 5:
            raise PdfExportError("PDF renderer produced no usable file")
        if rendered.read_bytes()[:5] != b"%PDF-":
            raise PdfExportError("rendered file is not a PDF")
        handle = tempfile.NamedTemporaryFile(
            dir=destination_path.parent,
            prefix=f".{destination_path.stem}-",
            suffix=".tmp",
            delete=False,
        )
        staged = Path(handle.name)
        try:
            with handle:
                with rendered.open("rb") as source_handle:
                    shutil.copyfileobj(source_handle, handle)
            staged.replace(destination_path)
        except Exception:
            staged.unlink(missing_ok=True)
            raise
    return destination_path


def _validated_sheets(source: Path, sheets: Sequence[str] | None) -> tuple[str, ...]:
    """Resolve all or explicit sheets without silently choosing one."""

    workbook = load_workbook(source, read_only=True)
    try:
        selected = tuple(workbook.sheetnames if sheets is None else sheets)
        if not selected:
            raise PdfExportError("PDF export requires at least one sheet")
        if len(set(selected)) != len(selected):
            raise PdfExportError("PDF export sheet selection contains duplicates")
        missing = [name for name in selected if name not in workbook.sheetnames]
        if missing:
            raise PdfExportError(f"unknown worksheet(s): {', '.join(missing)}")
        return selected
    finally:
        workbook.close()


def _prepare_workbook(source: Path, destination: Path, selected: Sequence[str]) -> None:
    """Create a temporary workbook containing only visible selected sheets."""

    workbook = load_workbook(source)
    try:
        for worksheet in list(workbook.worksheets):
            if worksheet.title not in selected:
                workbook.remove(worksheet)
            else:
                worksheet.sheet_state = "visible"
        workbook.active = 0
        workbook.save(destination)
    finally:
        workbook.close()
