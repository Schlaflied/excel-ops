"""Static formula integrity checks for the delivery verification pipeline.

openpyxl preserves formula text but is not a calculation engine.  This module
therefore reports recalculation as an explicit warning and never treats cached
values as proof that a formula still computes correctly.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable

from openpyxl.formula import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import range_boundaries

from .delivery_verification import DeliveryContract, VerificationFinding
from .template_writer import TemplateWriteResult


_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
_REF = re.compile(
    r"(?:(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_. ]*))!)?"
    r"(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?"
)
_EXTERNAL = re.compile(r"\[[^\]]+\][^!]*!")
_DYNAMIC = re.compile(r"(?:^|[^A-Z_])(FILTER|UNIQUE|SORTBY?|SEQUENCE|RANDARRAY)\s*\(|[A-Z]+\d+#", re.I)


@dataclass(frozen=True)
class FormulaRegion:
    """A declared range whose cells should inherit one translated formula."""

    sheet: str
    cell_range: str
    anchor_cell: str | None = None


@dataclass(frozen=True)
class SummaryReconciliation:
    """Independently sum source values and compare them with a stored summary."""

    source_sheet: str
    source_range: str
    summary_sheet: str
    summary_cell: str
    tolerance: Decimal = Decimal("0")


@dataclass(frozen=True)
class ExpectedErrorMarker:
    sheet: str
    cell_range: str
    values: tuple[str, ...] = ("#N/A",)


@dataclass(frozen=True)
class FormulaVerifier:
    formula_regions: tuple[FormulaRegion, ...] = field(default_factory=tuple)
    reconciliations: tuple[SummaryReconciliation, ...] = field(default_factory=tuple)
    expected_error_markers: tuple[ExpectedErrorMarker, ...] = field(default_factory=tuple)
    report_recalculation_limit: bool = True

    def __call__(
        self,
        workbook: Any,
        write_result: TemplateWriteResult,
        contract: DeliveryContract,
    ) -> Iterable[VerificationFinding]:
        del write_result, contract
        findings = list(_scan_formulas(workbook, self.expected_error_markers))
        for region in self.formula_regions:
            findings.extend(_check_region(workbook, region))
        for rule in self.reconciliations:
            findings.extend(_reconcile(workbook, rule))
        if self.report_recalculation_limit and any(
            _is_formula(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
        ):
            findings.append(
                VerificationFinding(
                    "recalculation_not_verified",
                    "Formula text was checked statically, but openpyxl cannot calculate formulas or prove cached results are current.",
                    "Recalculate the workbook in Excel or LibreOffice and validate the recalculated artifact in that engine.",
                    severity="warning",
                )
            )
        return findings


def _scan_formulas(workbook: Any, markers: tuple[ExpectedErrorMarker, ...]) -> Iterable[VerificationFinding]:
    graph: dict[tuple[str, str], set[tuple[str, str]]] = {}
    marker_cells = _marker_cells(workbook, markers)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                location = (sheet.title, cell.coordinate)
                if isinstance(value, str) and value in _ERRORS and location not in marker_cells:
                    yield _finding("formula_error_literal", f"Cell contains formula error {value}.", sheet.title, cell)
                if not _is_formula(value):
                    continue
                if _EXTERNAL.search(value):
                    yield _finding("external_reference", f"Formula references an external workbook: {value}", sheet.title, cell)
                if "#REF!" in value.upper():
                    yield _finding("broken_reference", f"Formula contains #REF!: {value}", sheet.title, cell)
                try:
                    error_tokens = [token.value for token in Tokenizer(value).items if token.subtype == "ERROR"]
                except ValueError:
                    error_tokens = []
                if error_tokens:
                    yield _finding("formula_error_token", f"Formula contains error token(s) {error_tokens}: {value}", sheet.title, cell)
                if _DYNAMIC.search(value):
                    yield _finding("unsupported_dynamic_array", f"Formula uses spill or dynamic-array behavior: {value}", sheet.title, cell)
                dependencies: set[tuple[str, str]] = set()
                for match in _REF.finditer(value):
                    raw_sheet = match.group(1) or match.group(2)
                    target_sheet = (raw_sheet.replace("''", "'") if raw_sheet else sheet.title).strip()
                    if target_sheet not in workbook.sheetnames:
                        yield _finding("missing_formula_sheet", f"Formula references missing sheet {target_sheet!r}: {value}", sheet.title, cell)
                        continue
                    start, end = match.group(3), match.group(4)
                    try:
                        min_col, min_row, max_col, max_row = range_boundaries(f"{start}:{end}" if end else start)
                        if max_col > 16384 or max_row > 1048576 or min_col < 1 or min_row < 1:
                            raise ValueError("outside Excel worksheet bounds")
                    except ValueError:
                        yield _finding("invalid_formula_range", f"Formula contains invalid range {match.group(0)!r}.", sheet.title, cell)
                        continue
                    if not end:
                        dependencies.add((target_sheet, start.replace("$", "")))
                graph[location] = dependencies
    yield from _cycles(graph)


def _check_region(workbook: Any, region: FormulaRegion) -> list[VerificationFinding]:
    if region.sheet not in workbook.sheetnames:
        return [VerificationFinding("missing_formula_region_sheet", f"Formula region sheet {region.sheet!r} is missing.", "Restore the declared sheet.", sheet=region.sheet)]
    sheet = workbook[region.sheet]
    cells = [cell for row in sheet[region.cell_range] for cell in row]
    if not cells:
        return []
    anchor = sheet[region.anchor_cell] if region.anchor_cell else next((cell for cell in cells if _is_formula(cell.value)), cells[0])
    if not _is_formula(anchor.value):
        return [_finding("missing_formula_anchor", "The declared formula region has no formula anchor.", region.sheet, anchor)]
    findings: list[VerificationFinding] = []
    for cell in cells:
        expected = Translator(anchor.value, origin=anchor.coordinate).translate_formula(cell.coordinate)
        if not _is_formula(cell.value):
            code = "formula_fill_gap" if cell.value is None else "formula_overwritten_by_constant"
            findings.append(_finding(code, f"Expected formula {expected!r}, got {cell.value!r}.", region.sheet, cell))
        elif cell.value != expected:
            findings.append(_finding("formula_pattern_mismatch", f"Expected translated formula {expected!r}, got {cell.value!r}.", region.sheet, cell))
    return findings


def _reconcile(workbook: Any, rule: SummaryReconciliation) -> list[VerificationFinding]:
    if rule.source_sheet not in workbook.sheetnames or rule.summary_sheet not in workbook.sheetnames:
        return [VerificationFinding("reconciliation_sheet_missing", "A declared reconciliation sheet is missing.", "Restore both source and summary sheets before delivery.")]
    source_cells = [cell for row in workbook[rule.source_sheet][rule.source_range] for cell in row]
    summary = workbook[rule.summary_sheet][rule.summary_cell]
    if any(_is_formula(cell.value) for cell in source_cells) or _is_formula(summary.value):
        return [_finding("reconciliation_not_computable", "Independent reconciliation needs stored numeric values; one or more declared cells contain formulas that openpyxl cannot calculate.", rule.summary_sheet, summary, severity="warning")]
    try:
        detail_total = sum((_decimal(cell.value) for cell in source_cells if cell.value not in (None, "")), Decimal("0"))
        summary_total = _decimal(summary.value)
    except (InvalidOperation, TypeError, ValueError):
        return [_finding("reconciliation_non_numeric", "The declared reconciliation contains a non-numeric value.", rule.summary_sheet, summary)]
    difference = abs(detail_total - summary_total)
    if difference > rule.tolerance:
        return [_finding("summary_reconciliation_mismatch", f"Independent detail total {detail_total} does not match summary {summary_total}; difference {difference} exceeds tolerance {rule.tolerance}.", rule.summary_sheet, summary)]
    return []


def _marker_cells(workbook: Any, markers: tuple[ExpectedErrorMarker, ...]) -> set[tuple[str, str]]:
    expected: set[tuple[str, str]] = set()
    for marker in markers:
        if marker.sheet not in workbook.sheetnames:
            continue
        for row in workbook[marker.sheet][marker.cell_range]:
            for cell in row:
                if str(cell.value) in marker.values:
                    expected.add((marker.sheet, cell.coordinate))
    return expected


def _cycles(graph: dict[tuple[str, str], set[tuple[str, str]]]) -> Iterable[VerificationFinding]:
    state: dict[tuple[str, str], int] = {}
    stack: list[tuple[str, str]] = []
    reported: set[tuple[str, str]] = set()

    def visit(node: tuple[str, str]) -> None:
        state[node] = 1
        stack.append(node)
        for child in graph.get(node, ()):
            if child not in graph:
                continue
            if state.get(child, 0) == 0:
                visit(child)
            elif state[child] == 1:
                reported.update(stack[stack.index(child) :])
        stack.pop()
        state[node] = 2

    for node in graph:
        if state.get(node, 0) == 0:
            visit(node)
    for sheet, coordinate in sorted(reported):
        yield VerificationFinding("circular_reference", "Formula participates in a circular reference.", "Break the dependency cycle before delivery.", sheet=sheet, cell=coordinate)


def _finding(code: str, message: str, sheet: str, cell: Any, *, severity: str = "error") -> VerificationFinding:
    return VerificationFinding(code, message, "Repair the formula or declare the intentional exception before delivery.", sheet, cell.row, cell=cell.coordinate, severity=severity)


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _decimal(value: Any) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise TypeError(value)
    return Decimal(str(value))
