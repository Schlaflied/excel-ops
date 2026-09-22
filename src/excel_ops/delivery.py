"""End-to-end delivery orchestration for declared workbook templates.

This module contains no new matching, ambiguity, write-back or verification
logic.  It only sequences the already-merged modules into one agent-callable
pipeline and returns a single JSON-serializable result:

``ingest -> validate -> strict match -> ambiguity/Recipe -> plan ->
staged template write -> independent verification of the persisted file ->
delivery``

A successful ``write_template`` call is never reported as a delivery.  Only a
target whose reopened file passed :func:`verify_and_deliver` is published.

Passing ``idempotency=IdempotencyOptions(...)`` additionally turns on the
whole-run short-circuit from :mod:`excel_ops.idempotency`: the run fingerprint is
computed *before* any matching, writing or verification, and a fingerprint that
matches a previously successful run returns a ``no_op`` result immediately.  That
is a different mechanism from the per-record deduplication below
(``_existing_record_ids`` / ``_build_plan`` / ``_mark_existing``), which stops one
row being appended twice inside a target workbook; both stay in force.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field, replace
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .ambiguity import (
    UNKNOWN,
    Ambiguity,
    ConfirmationBatch,
    RecipeDecision,
    build_confirmation_batch,
    load_project_recipe,
    save_project_recipe,
)
from .cells import column_number, is_merged_non_anchor, is_protected_formula_value
from .delivery_manifest import (
    DeliveryManifest,
    build_delivery_manifests,
    write_delivery_manifests,
)
from .delivery_verification import (
    DeliveryContract,
    DeliveryVerificationError,
    PeriodExpectation,
    StageCounts,
    VerificationFinding,
    verify_and_deliver,
)
from .formula_verification import FormulaVerifier
from .idempotency import (
    CHANGED,
    FAILED,
    NO_OP,
    STARTED,
    SUCCEEDED,
    ConnectorTarget,
    IdempotencyError,
    IdempotencyOptions,
    RunDecision,
    RunFingerprint,
    RunRecord,
    compute_fingerprint,
    evaluate_run,
    load_run_record,
    record_run,
)
from .ingestion import LayoutDetectionError, load_input_records
from .matching import Destination, MatchResult, preallocate_records, stable_record_id
from .models import ExtractedRecord
from .periods import PeriodResult
from .review import review_record
from .template_writer import (
    TemplateMapping,
    TemplateWriteError,
    TemplateWriteResult,
    write_template,
)


#: Every field a delivery row may map.  A mapping naming anything else is a
#: blocking planning error rather than a silent empty cell.
CONTRACT_FIELDS = (
    "record_id",
    "location",
    "event_date",
    "identifier",
    "category",
    "confidence",
    "source",
    "source_file",
    "source_sheet",
    "source_row",
    "source_region",
)

WRITTEN = "written"
ACCEPTED = "accepted"
REVIEW = "review"
REJECTED = "rejected"
SKIPPED_EXISTING = "skipped_existing"

_REJECT_REASON_PREFIXES = ("missing:", "invalid:")
_AMBIGUITY_FIELD = "destination"


class DeliveryPlanError(ValueError):
    """Raised when a declared delivery target cannot be planned at all."""


@dataclass(frozen=True)
class DeliveryTarget:
    """One declared destination plus the template it is written into."""

    destination: Destination
    template_path: str | Path
    mapping: TemplateMapping
    required_fields: tuple[str, ...] = ()
    record_id_field: str = "record_id"
    period_expectations: tuple[PeriodExpectation, ...] = field(default_factory=tuple)
    formula_verifier: FormulaVerifier | None = None
    delivery_name: str | None = None

    @property
    def key(self) -> str:
        return self.destination.key


@dataclass(frozen=True)
class SourceTrace:
    """Where a record came from, before any workbook was touched."""

    locator: str
    source_file: str
    source_sheet: str
    source_row: int | None
    source_region: str

    @classmethod
    def of(cls, record: ExtractedRecord) -> "SourceTrace":
        return cls(
            record.source,
            record.source_file,
            record.source_sheet,
            record.source_row,
            record.source_region,
        )


@dataclass(frozen=True)
class WrittenCell:
    """A delivered cell traced back to its source and stable record ID."""

    record_id: str
    sheet: str
    cell: str
    field: str
    value: Any
    source: SourceTrace

    def to_dict(self) -> dict[str, Any]:
        return {
            "record_id": self.record_id,
            "sheet": self.sheet,
            "cell": self.cell,
            "field": self.field,
            "value": self.value,
            "source": asdict(self.source),
        }


@dataclass(frozen=True)
class RecordOutcome:
    """The single terminal state of one input record."""

    record_id: str
    status: str
    destination_key: str | None
    match_status: str
    rule: str | None
    confidence: float
    reasons: tuple[str, ...]
    source: SourceTrace
    values: Mapping[str, Any]
    candidates: tuple[str, ...] = field(default_factory=tuple)
    ambiguity_key: str | None = None
    cells: tuple[WrittenCell, ...] = field(default_factory=tuple)

    def to_dict(self) -> dict[str, Any]:
        return {
            "record_id": self.record_id,
            "status": self.status,
            "destination_key": self.destination_key,
            "match_status": self.match_status,
            "rule": self.rule,
            "confidence": self.confidence,
            "reasons": list(self.reasons),
            "source": asdict(self.source),
            "values": dict(self.values),
            "candidates": list(self.candidates),
            "ambiguity_key": self.ambiguity_key,
            "cells": [item.to_dict() for item in self.cells],
        }


@dataclass(frozen=True)
class PlannedTarget:
    """What a target is expected to do, produced before any file is touched."""

    destination_key: str
    template_path: str
    sheet: str
    field_mapping: Mapping[str, str | int]
    staged_path: str
    delivery_path: str
    expected_written: int
    expected_review: int
    expected_skipped_existing: int
    #: Rows the writer is predicted to refuse because a mapped target cell
    #: already holds a formula or is a merged-range follower.  Those rows are
    #: excluded from ``expected_written`` so the plan never overstates the write.
    expected_skipped_protected: int = 0
    protected_cells: tuple[str, ...] = field(default_factory=tuple)
    blocking_items: tuple[str, ...] = field(default_factory=tuple)

    def to_dict(self) -> dict[str, Any]:
        return {
            "destination_key": self.destination_key,
            "template_path": self.template_path,
            "sheet": self.sheet,
            "field_mapping": {name: str(column) for name, column in self.field_mapping.items()},
            "staged_path": self.staged_path,
            "delivery_path": self.delivery_path,
            "expected_written": self.expected_written,
            "expected_review": self.expected_review,
            "expected_skipped_existing": self.expected_skipped_existing,
            "expected_skipped_protected": self.expected_skipped_protected,
            "protected_cells": list(self.protected_cells),
            "blocking_items": list(self.blocking_items),
        }


@dataclass(frozen=True)
class DeliveryPlan:
    """The checkable dry-run contract for a run."""

    inputs: tuple[str, ...]
    counts: Mapping[str, int]
    targets: tuple[PlannedTarget, ...]
    unresolved_ambiguities: tuple[str, ...] = field(default_factory=tuple)
    blocking_items: tuple[str, ...] = field(default_factory=tuple)

    @property
    def writable(self) -> bool:
        return not self.blocking_items

    def to_dict(self) -> dict[str, Any]:
        return {
            "inputs": list(self.inputs),
            "counts": dict(self.counts),
            "targets": [item.to_dict() for item in self.targets],
            "unresolved_ambiguities": list(self.unresolved_ambiguities),
            "blocking_items": list(self.blocking_items),
            "writable": self.writable,
        }


@dataclass(frozen=True)
class DeliveryFailure:
    """A machine-readable reason a target or the run did not deliver."""

    code: str
    message: str
    suggestion: str
    destination_key: str | None = None
    findings: tuple[VerificationFinding, ...] = field(default_factory=tuple)

    def to_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "message": self.message,
            "suggestion": self.suggestion,
            "destination_key": self.destination_key,
            "findings": [asdict(item) for item in self.findings],
        }


@dataclass(frozen=True)
class TargetOutcome:
    """What actually happened to one target, judged by the persisted file."""

    destination_key: str
    template_path: str
    delivered: bool
    staged_path: str | None = None
    delivery_path: str | None = None
    change_log_path: str | None = None
    verification_report_path: str | None = None
    written_record_ids: tuple[str, ...] = field(default_factory=tuple)
    written_cells: int = 0
    skipped_writes: tuple[Mapping[str, Any], ...] = field(default_factory=tuple)
    findings: tuple[VerificationFinding, ...] = field(default_factory=tuple)
    status: str = "planned"

    def to_dict(self) -> dict[str, Any]:
        return {
            "destination_key": self.destination_key,
            "template_path": self.template_path,
            "delivered": self.delivered,
            "status": self.status,
            "staged_path": self.staged_path,
            "delivery_path": self.delivery_path,
            "change_log_path": self.change_log_path,
            "verification_report_path": self.verification_report_path,
            "written_record_ids": list(self.written_record_ids),
            "written_cells": self.written_cells,
            "skipped_writes": [dict(item) for item in self.skipped_writes],
            "findings": [asdict(item) for item in self.findings],
        }


@dataclass(frozen=True)
class AmbiguityOutcome:
    key: str
    field_name: str
    question: str
    status: str
    selected: str | None
    decision_source: str | None
    candidates: tuple[str, ...]
    record_ids: tuple[str, ...]

    def to_dict(self) -> dict[str, Any]:
        return {
            "key": self.key,
            "field": self.field_name,
            "question": self.question,
            "status": self.status,
            "selected": self.selected,
            "decision_source": self.decision_source,
            "candidates": list(self.candidates),
            "record_ids": list(self.record_ids),
        }


@dataclass(frozen=True)
class DeliveryRun:
    """The single JSON-serializable result of one delivery attempt."""

    delivered: bool
    dry_run: bool
    plan: DeliveryPlan
    counts: Mapping[str, int]
    records: tuple[RecordOutcome, ...]
    targets: tuple[TargetOutcome, ...]
    ambiguities: tuple[AmbiguityOutcome, ...] = field(default_factory=tuple)
    failures: tuple[DeliveryFailure, ...] = field(default_factory=tuple)
    recipe_path: str | None = None
    #: The live confirmation batch, so callers can build reusable Recipe
    #: decisions with ``ambiguity.decide(...)``. Not part of ``to_dict()``.
    confirmation_batch: ConfirmationBatch | None = field(default=None, repr=False)
    #: True when whole-run idempotency short-circuited this call: nothing
    #: changed since a previously successful run, so no work was redone.
    no_op: bool = False
    #: The idempotency verdict for this call, when ``idempotency`` was supplied.
    run_decision: RunDecision | None = field(default=None, repr=False)
    #: One source-and-verification Manifest per *delivered* output (issue #20),
    #: generated from this run's own outcomes and from the persisted file.  A run
    #: that delivered nothing carries none.
    manifests: tuple[DeliveryManifest, ...] = field(default_factory=tuple)

    @property
    def manifest_paths(self) -> tuple[str, ...]:
        return tuple(str(item.manifest_path()) for item in self.manifests)

    @property
    def fingerprint(self) -> str | None:
        """The run fingerprint, when whole-run idempotency was enabled."""

        return self.run_decision.fingerprint.digest if self.run_decision else None

    @property
    def delivery_paths(self) -> tuple[str, ...]:
        return tuple(
            item.delivery_path for item in self.targets if item.delivered and item.delivery_path
        )

    def records_by_status(self, status: str) -> tuple[RecordOutcome, ...]:
        return tuple(item for item in self.records if item.status == status)

    def to_dict(self) -> dict[str, Any]:
        return {
            "delivered": self.delivered,
            "dry_run": self.dry_run,
            "plan": self.plan.to_dict(),
            "counts": dict(self.counts),
            "records": [item.to_dict() for item in self.records],
            "targets": [item.to_dict() for item in self.targets],
            "ambiguities": [item.to_dict() for item in self.ambiguities],
            "failures": [item.to_dict() for item in self.failures],
            "recipe_path": self.recipe_path,
            "delivery_paths": list(self.delivery_paths),
            "no_op": self.no_op,
            "run_decision": self.run_decision.to_dict() if self.run_decision else None,
            "manifests": [item.to_dict() for item in self.manifests],
            "manifest_paths": list(self.manifest_paths),
        }


def run_delivery(
    inputs: Sequence[str | Path],
    targets: Sequence[DeliveryTarget],
    *,
    staging_dir: str | Path,
    delivery_dir: str | Path,
    confidence_threshold: float = 0.85,
    recipe_path: str | Path | None = None,
    decisions: Sequence[RecipeDecision] = (),
    dry_run: bool = False,
    post_stage_hook: Callable[[Path], None] | None = None,
    idempotency: IdempotencyOptions | None = None,
    period: PeriodResult | None = None,
    write_manifest: bool = True,
) -> DeliveryRun:
    """Run ingest to verified delivery for declared targets and return one result.

    ``post_stage_hook`` receives each staged workbook path after the write and
    before verification.  It exists so callers and tests can inspect or corrupt
    the staged artifact and prove that verification, not the writer, decides
    whether a file is delivered.

    ``idempotency`` opts this call into the whole-run short-circuit.  The run
    fingerprint is computed before any matching, write or verification; when it
    matches a previously *successful* run the call returns a ``no_op`` result
    without redoing the work.  A dry run never short-circuits -- it still returns
    the full plan -- but it does report the verdict in ``run_decision``.

    ``period`` is the resolved report period this run delivers for.  It is
    recorded in the #20 Manifest as ``period_start``/``period_end``; the period
    values written *into* a workbook are still declared per target as
    ``period_expectations`` and verified by #4.

    Every delivered output gets one Manifest in ``result.manifests``, built from
    this run's own outcomes and from the persisted file.  ``write_manifest=True``
    (the default) also writes ``<delivered file>.manifest.json`` and
    ``.manifest.txt`` beside it.  A dry run and a no-op produce no Manifest,
    because no file was delivered.
    """

    if not targets:
        raise DeliveryPlanError("at least one delivery target is required")
    _validate_targets(targets)
    staging = Path(staging_dir)
    delivery = Path(delivery_dir)

    failures: list[DeliveryFailure] = []
    project_recipe, recipe_failure = _load_recipe(recipe_path)
    guard = _RunGuard.of(inputs, targets, delivery, project_recipe, decisions, idempotency)
    if guard is not None and guard.decision.no_op and not dry_run:
        return _no_op_run(inputs, targets, delivery, guard.decision, recipe_path)

    records, record_inputs, ingest_failures = _ingest(inputs)
    failures.extend(ingest_failures)

    outcomes, match_results = _classify(records, targets, confidence_threshold)
    ambiguities, ambiguity_records = _ambiguities_from_matches(match_results)
    if recipe_failure is not None:
        failures.append(recipe_failure)
    run_decisions = {item.key: item for item in decisions if item.scope == "this-run"}
    project_decisions = {item.key: item for item in decisions if item.scope == "project"}
    batch = build_confirmation_batch(
        ambiguities,
        run_decisions=run_decisions,
        project_recipe={**project_recipe, **project_decisions},
    )
    confirmations = _confirmations(batch, ambiguity_records, targets)
    if confirmations:
        # Re-run matching with the reusable human decisions. The first batch is
        # kept as the provenance of those decisions.
        outcomes, _ = _classify(
            records, targets, confidence_threshold, confirmations=confirmations
        )

    outcomes = _block_unresolved(outcomes, batch, ambiguity_records)
    # A dry run must not touch a single file, and a real run must never drop a
    # decision an earlier run already recorded: merge the loaded Recipe with
    # this call's decisions so the file only ever grows or updates in place.
    if not dry_run and recipe_path is not None and project_decisions:
        save_project_recipe({**project_recipe, **project_decisions}.values(), recipe_path)

    plan, existing_ids = _build_plan(inputs, outcomes, targets, staging, delivery, batch)
    outcomes = _mark_existing(outcomes, existing_ids)
    plan = replace(plan, counts=_counts(outcomes))

    if dry_run or not plan.writable or failures:
        if not plan.writable:
            failures.append(
                DeliveryFailure(
                    "plan_blocked",
                    f"The delivery plan is blocked by {len(plan.blocking_items)} item(s).",
                    "Resolve every blocking item, then re-run the delivery.",
                )
            )
        if guard is not None and not dry_run:
            # A blocked or failed run is recorded as a failure, never as a
            # baseline a later identical run could short-circuit against.
            guard.finish(FAILED, False, _counts(outcomes), failures, ())
        return DeliveryRun(
            False,
            dry_run,
            plan,
            _counts(outcomes),
            tuple(outcomes),
            tuple(_planned_outcome(item, plan) for item in targets),
            tuple(_ambiguity_outcomes(batch, ambiguity_records)),
            tuple(failures),
            str(recipe_path) if recipe_path else None,
            batch,
            False,
            guard.decision if guard else None,
        )

    if guard is not None:
        # Mark the attempt before the first write, so a run interrupted between
        # here and the end is found as ``started`` and retried rather than being
        # mistaken for a completed no-op baseline.
        guard.start()
    outcomes, target_outcomes, write_failures = _write_and_verify(
        outcomes, targets, plan, post_stage_hook
    )
    failures.extend(write_failures)
    delivered = bool(target_outcomes) and all(
        item.delivered or item.status == "no_new_records" for item in target_outcomes
    ) and not failures
    if guard is not None:
        guard.finish(
            SUCCEEDED if delivered else FAILED,
            delivered,
            _counts(outcomes),
            failures,
            target_outcomes,
        )
    run = DeliveryRun(
        delivered,
        False,
        plan,
        _counts(outcomes),
        tuple(outcomes),
        tuple(target_outcomes),
        tuple(_ambiguity_outcomes(batch, ambiguity_records)),
        tuple(failures),
        str(recipe_path) if recipe_path else None,
        batch,
        False,
        guard.decision if guard else None,
    )
    # The Manifest is generated last, from the run that just finished and from
    # the file that is now on disk.  It is evidence about a completed delivery,
    # never a parallel bookkeeping system that could drift from it.
    manifests = build_delivery_manifests(
        run,
        targets,
        period=period,
        recipe_decisions={**project_recipe, **project_decisions},
        record_input_paths=record_inputs,
    )
    if write_manifest:
        write_delivery_manifests(manifests)
    return replace(run, manifests=manifests)


def plan_delivery(
    inputs: Sequence[str | Path],
    targets: Sequence[DeliveryTarget],
    *,
    staging_dir: str | Path,
    delivery_dir: str | Path,
    confidence_threshold: float = 0.85,
    recipe_path: str | Path | None = None,
    decisions: Sequence[RecipeDecision] = (),
    idempotency: IdempotencyOptions | None = None,
) -> DeliveryRun:
    """Return the checkable plan without touching a single file."""

    return run_delivery(
        inputs,
        targets,
        staging_dir=staging_dir,
        delivery_dir=delivery_dir,
        confidence_threshold=confidence_threshold,
        recipe_path=recipe_path,
        decisions=decisions,
        dry_run=True,
        idempotency=idempotency,
    )


def _delivery_path(target: DeliveryTarget, delivery: Path) -> Path:
    """The single place a target's delivered filename is derived."""

    template = Path(target.template_path)
    return delivery / (
        target.delivery_name or f"{template.stem}-{target.key}{template.suffix}"
    )


def _period_payload(targets: Sequence[DeliveryTarget], period: Any) -> dict[str, Any]:
    """The declared report period: the explicit one plus the in-workbook banners."""

    return {
        "declared": period,
        "expectations": sorted(
            f"{item.field}@{item.sheet}!{item.cell}={item.expected}"
            for target in targets
            for item in target.period_expectations
        ),
    }


def _mapping_payload(targets: Sequence[DeliveryTarget]) -> list[dict[str, Any]]:
    """The declared mapping version: what each target promises to write, where."""

    return [
        {
            "key": target.key,
            "aliases": sorted(target.destination.aliases),
            "sheet": target.mapping.sheet,
            "field_columns": {
                name: str(column) for name, column in sorted(target.mapping.field_columns.items())
            },
            "header_row": target.mapping.header_row,
            "data_start_row": target.mapping.data_start_row,
            "template_type": target.mapping.template_type,
            "style_source_row": target.mapping.style_source_row,
            "max_rows": target.mapping.max_rows,
            "overwrite_formulas": target.mapping.overwrite_formulas,
            "record_id_field": target.record_id_field,
            "required_fields": sorted(target.required_fields),
        }
        for target in sorted(targets, key=lambda item: item.key)
    ]


@dataclass
class _RunGuard:
    """The whole-run idempotency state for one ``run_delivery`` call."""

    options: IdempotencyOptions
    state_path: Path
    task_key: str
    connector: ConnectorTarget
    decision: RunDecision
    previous: RunRecord | None
    #: Recomputes the fingerprint against the files as they are *now*, so the
    #: record left behind describes the delivered state rather than the
    #: pre-delivery one.  Without this, the output content hash recorded before
    #: the write would never match the next run and no run could ever be a no-op.
    recompute: Callable[[], RunFingerprint]

    @classmethod
    def of(
        cls,
        inputs: Sequence[str | Path],
        targets: Sequence[DeliveryTarget],
        delivery: Path,
        project_recipe: Mapping[str, RecipeDecision],
        decisions: Sequence[RecipeDecision],
        options: IdempotencyOptions | None,
    ) -> "_RunGuard | None":
        if options is None:
            return None
        outputs = [_delivery_path(target, delivery) for target in targets]
        connector = options.connector or ConnectorTarget.local(outputs)
        templates = [target.template_path for target in targets]
        mapping = _mapping_payload(targets)
        period = _period_payload(targets, options.period)

        def recompute() -> RunFingerprint:
            return compute_fingerprint(
                inputs=inputs,
                templates=templates,
                template_profile_version=options.template_profile_version,
                mapping=mapping,
                recipe_decisions=project_recipe,
                confirmations=decisions,
                period=period,
                connector=connector,
                revision_source=options.revision_source,
            )

        fingerprint = recompute()
        path = options.resolved_state_path(delivery)
        task_key = options.resolved_task_key(connector)
        try:
            previous = load_run_record(path, task_key)
        except IdempotencyError:
            # An unreadable or foreign run-state file must never authorize a
            # no-op; the run proceeds and rewrites the record.
            return cls(
                options,
                path,
                task_key,
                connector,
                RunDecision(CHANGED, "unreadable_run_state", fingerprint),
                None,
                recompute,
            )
        decision = evaluate_run(
            fingerprint,
            previous,
            connector=connector,
            revision_source=options.revision_source,
        )
        return cls(options, path, task_key, connector, decision, previous, recompute)

    def start(self) -> None:
        record_run(
            self.decision.fingerprint,
            self.state_path,
            task_key=self.task_key,
            status=STARTED,
            delivered=False,
            detail={"stage": "write_and_verify"},
            previous=self.previous,
        )

    def finish(
        self,
        status: str,
        delivered: bool,
        counts: Mapping[str, int],
        failures: Sequence[DeliveryFailure],
        target_outcomes: Sequence[TargetOutcome],
    ) -> RunRecord:
        # Only aggregate, non-identifying detail is persisted: counts, failure
        # codes, and how many targets were delivered.  No path, no destination
        # name, no cell value ever reaches the run record.
        detail = {
            "counts": dict(counts),
            "failure_codes": sorted({item.code for item in failures}),
            "targets": len(target_outcomes),
            "delivered_targets": sum(1 for item in target_outcomes if item.delivered),
        }
        return record_run(
            self.recompute(),
            self.state_path,
            task_key=self.task_key,
            status=status,
            delivered=delivered,
            detail=detail,
            previous=self.previous,
        )


def _no_op_run(
    inputs: Sequence[str | Path],
    targets: Sequence[DeliveryTarget],
    delivery: Path,
    decision: RunDecision,
    recipe_path: str | Path | None,
) -> DeliveryRun:
    """Return "nothing changed" without redoing matching, writing, or verifying."""

    planned: list[PlannedTarget] = []
    outcomes: list[TargetOutcome] = []
    for target in targets:
        template = Path(target.template_path)
        delivered_path = _delivery_path(target, delivery)
        planned.append(
            PlannedTarget(
                destination_key=target.key,
                template_path=str(template),
                sheet=target.mapping.sheet,
                field_mapping=dict(target.mapping.field_columns),
                staged_path="",
                delivery_path=str(delivered_path),
                expected_written=0,
                expected_review=0,
                expected_skipped_existing=0,
            )
        )
        outcomes.append(
            TargetOutcome(
                target.key,
                str(template),
                False,
                delivery_path=str(delivered_path) if delivered_path.is_file() else None,
                status=NO_OP,
            )
        )
    plan = DeliveryPlan(tuple(str(item) for item in inputs), _counts(()), tuple(planned))
    return DeliveryRun(
        False,
        False,
        plan,
        _counts(()),
        (),
        tuple(outcomes),
        (),
        (),
        str(recipe_path) if recipe_path else None,
        None,
        True,
        decision,
    )


def _validate_targets(targets: Sequence[DeliveryTarget]) -> None:
    keys = [item.key for item in targets]
    if len(keys) != len(set(keys)):
        raise DeliveryPlanError("delivery target destination keys must be unique")


def _ingest(
    inputs: Sequence[str | Path],
) -> tuple[list[ExtractedRecord], list[str], list[DeliveryFailure]]:
    """Ingest every input and remember which input each record came from.

    The third return value is aligned with the records: entry *i* is the input
    path record *i* was read from.  A record's own ``source_file`` names the
    provenance it *declares* — an extraction JSON's records name the image, not
    the JSON — so only the ingester knows which declared input produced it.  The
    #20 Manifest needs that to hash the right file per source.
    """

    records: list[ExtractedRecord] = []
    origins: list[str] = []
    failures: list[DeliveryFailure] = []
    for item in inputs:
        path = Path(item)
        try:
            loaded = load_input_records(path)
        except LayoutDetectionError as error:
            failures.append(
                DeliveryFailure(
                    "unknown_layout",
                    f"{path.name}: {error}",
                    "Confirm the header row or declare a mapping before re-running.",
                )
            )
        except (OSError, ValueError) as error:
            failures.append(
                DeliveryFailure(
                    "unreadable_input",
                    f"{path.name}: {error}",
                    "Supply a readable .xlsx, .xlsm, .csv, or extraction .json input.",
                )
            )
        else:
            records.extend(loaded)
            origins.extend(str(item) for _ in loaded)
    return records, origins, failures


def _row_values(record_id: str, record: ExtractedRecord) -> dict[str, Any]:
    return {
        "record_id": record_id,
        "location": record.location,
        "event_date": record.event_date,
        "identifier": record.identifier,
        "category": record.category,
        "confidence": record.confidence,
        "source": record.source,
        "source_file": record.source_file,
        "source_sheet": record.source_sheet,
        "source_row": record.source_row,
        "source_region": record.source_region,
    }


def _classify(
    records: Sequence[ExtractedRecord],
    targets: Sequence[DeliveryTarget],
    confidence_threshold: float,
    *,
    confirmations: Mapping[str, str] | None = None,
) -> tuple[list[RecordOutcome], list[MatchResult]]:
    """Route every record to exactly one non-written state before any write."""

    reviewed = [(record, review_record(record, confidence_threshold)) for record in records]
    rejected_reasons: dict[int, tuple[str, ...]] = {}
    matchable: list[ExtractedRecord] = []
    order: list[tuple[int, ExtractedRecord, tuple[str, ...]]] = []
    for index, (record, verdict) in enumerate(reviewed):
        if any(reason.startswith(_REJECT_REASON_PREFIXES) for reason in verdict.reasons):
            rejected_reasons[index] = verdict.reasons
        else:
            matchable.append(record)
        order.append((index, record, verdict.reasons))

    match_results = preallocate_records(
        matchable,
        [item.destination for item in targets],
        strict=False,
        confirmations=confirmations,
    )
    # preallocate_records returns exactly one result per input record, in order.
    by_record = dict(zip((index for index, _, _ in order if index not in rejected_reasons), match_results))

    outcomes: list[RecordOutcome] = []
    for index, record, reasons in order:
        if index in rejected_reasons:
            record_id = stable_record_id(record)
            outcomes.append(
                RecordOutcome(
                    record_id,
                    REJECTED,
                    None,
                    "not_matched",
                    None,
                    record.confidence,
                    rejected_reasons[index],
                    SourceTrace.of(record),
                    _row_values(record_id, record),
                )
            )
            continue
        match = by_record[index]
        status = ACCEPTED if match.status == "matched" and not reasons else REVIEW
        combined = tuple(dict.fromkeys((*reasons, *match.reasons)))
        outcomes.append(
            RecordOutcome(
                match.record_id,
                status,
                match.destination_key if status == ACCEPTED else None,
                match.status,
                match.rule,
                match.confidence,
                combined,
                SourceTrace.of(record),
                _row_values(match.record_id, record),
                tuple(item.destination_key for item in match.candidates),
                _ambiguity_key(match),
            )
        )
    return outcomes, match_results


def _ambiguity_shape(match: MatchResult) -> tuple[str, tuple[str, ...]] | None:
    """Return the record-independent shape of a match's open question.

    ``Ambiguity.key`` is ``field:question``, and ``group_ambiguities`` merges
    every ambiguity sharing a key into one confirmation item.  The key must
    therefore describe the *kind* of question — which candidates are competing
    for which field — and must not embed the individual record's source value,
    or nothing would ever group and a saved Recipe decision would only ever
    match the one run that produced it.
    """

    if match.status in {"review", "conflict"} and match.candidates:
        candidates = tuple(dict.fromkeys(item.destination_key for item in match.candidates))
        return match.status, candidates
    return None


def _ambiguity_question(status: str, candidates: Sequence[str]) -> str:
    joined = ", ".join(candidates)
    if status == "conflict":
        return (
            f"Which declared destination owns source values claimed by "
            f"more than one of [{joined}]?"
        )
    return f"Which declared destination owns source values fuzzily matching [{joined}]?"


def _ambiguity_key(match: MatchResult) -> str | None:
    shape = _ambiguity_shape(match)
    if shape is None:
        return None
    status, candidates = shape
    return f"{_AMBIGUITY_FIELD}:{_ambiguity_question(status, candidates)}"


def _ambiguities_from_matches(
    match_results: Sequence[MatchResult],
) -> tuple[tuple[Ambiguity, ...], dict[str, tuple[str, ...]]]:
    items: list[Ambiguity] = []
    records: dict[str, list[str]] = {}
    for match in match_results:
        shape = _ambiguity_shape(match)
        if shape is None:
            continue
        status, candidates = shape
        question = _ambiguity_question(status, candidates)
        key = f"{_AMBIGUITY_FIELD}:{question}"
        recommendation = candidates[0] if status == "review" and len(candidates) == 1 else None
        items.append(
            Ambiguity(
                _AMBIGUITY_FIELD,
                question,
                (match.record.location,),
                candidates,
                1,
                recommendation,
                "fuzzy candidate requires confirmation"
                if status == "review"
                else "several declared destinations claim this value",
                match.confidence,
                (match.record.source,),
            )
        )
        records.setdefault(key, []).append(match.record_id)
    return tuple(items), {key: tuple(value) for key, value in records.items()}


def _load_recipe(
    recipe_path: str | Path | None,
) -> tuple[dict[str, RecipeDecision], DeliveryFailure | None]:
    if recipe_path is None or not Path(recipe_path).is_file():
        return {}, None
    try:
        return load_project_recipe(recipe_path), None
    except (OSError, ValueError) as error:
        return {}, DeliveryFailure(
            "unreadable_recipe",
            f"{Path(recipe_path).name}: {error}",
            "Repair or remove the project Recipe before re-running the delivery.",
        )


def _confirmations(
    batch: ConfirmationBatch,
    ambiguity_records: Mapping[str, tuple[str, ...]],
    targets: Sequence[DeliveryTarget],
) -> dict[str, str]:
    keys = {item.key for item in targets}
    confirmations: dict[str, str] = {}
    for item in batch.items:
        if item.status != "resolved" or not item.selected or item.selected == UNKNOWN:
            continue
        if item.selected not in keys:
            continue
        for record_id in ambiguity_records.get(item.ambiguity.key, ()):
            confirmations[record_id] = item.selected
    return confirmations


def _block_unresolved(
    outcomes: Sequence[RecordOutcome],
    batch: ConfirmationBatch,
    ambiguity_records: Mapping[str, tuple[str, ...]],
) -> list[RecordOutcome]:
    """Never let a record with an unresolved question reach accepted."""

    blocked: dict[str, str] = {}
    for key in batch.unresolved_blockers:
        for record_id in ambiguity_records.get(key, ()):
            blocked[record_id] = key
    updated: list[RecordOutcome] = []
    for item in outcomes:
        if item.record_id in blocked and item.status != REJECTED:
            updated.append(
                replace(
                    item,
                    status=REVIEW,
                    destination_key=None,
                    reasons=tuple(dict.fromkeys((*item.reasons, "unresolved_ambiguity"))),
                )
            )
        else:
            updated.append(item)
    return updated


def _ambiguity_outcomes(
    batch: ConfirmationBatch, ambiguity_records: Mapping[str, tuple[str, ...]]
) -> list[AmbiguityOutcome]:
    return [
        AmbiguityOutcome(
            item.ambiguity.key,
            item.ambiguity.field,
            item.ambiguity.question,
            item.status,
            item.selected,
            item.decision_source,
            item.ambiguity.candidates,
            ambiguity_records.get(item.ambiguity.key, ()),
        )
        for item in batch.items
    ]


def _counts(outcomes: Sequence[RecordOutcome]) -> dict[str, int]:
    counts = {
        "input": len(outcomes),
        WRITTEN: 0,
        ACCEPTED: 0,
        REVIEW: 0,
        REJECTED: 0,
        SKIPPED_EXISTING: 0,
    }
    for item in outcomes:
        counts[item.status] += 1
    return counts


def _column_number(value: str | int) -> int:
    """Convert the shared parser's ``ValueError`` into this module's contract."""

    try:
        return column_number(value)
    except ValueError as exc:
        raise DeliveryPlanError(str(exc)) from exc


def _existing_record_ids(path: Path, target: DeliveryTarget) -> set[str]:
    """Read the record IDs already present in a previously delivered file."""

    column = target.mapping.field_columns.get(target.record_id_field)
    if not path.is_file() or column is None:
        return set()
    index = _column_number(column)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[target.mapping.sheet] if target.mapping.sheet in workbook.sheetnames else None
        if sheet is None:
            return set()
        found: set[str] = set()
        for row in sheet.iter_rows(
            min_row=target.mapping.data_start_row, min_col=index, max_col=index, values_only=True
        ):
            value = str(row[0] or "").strip()
            if value:
                found.add(value)
        return found
    finally:
        workbook.close()


def _unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    for revision in range(2, 1000):
        candidate = path.with_name(f"{path.stem}-r{revision}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise DeliveryPlanError(f"unable to allocate a staging path next to {path}")


def _target_blockers(target: DeliveryTarget) -> list[str]:
    blockers: list[str] = []
    template = Path(target.template_path)
    if not template.is_file():
        blockers.append(f"{target.key}: missing_template:{template.name}")
    unknown = [name for name in target.mapping.field_columns if name not in CONTRACT_FIELDS]
    for name in sorted(unknown):
        blockers.append(f"{target.key}: unmapped_contract_field:{name}")
    if target.record_id_field not in target.mapping.field_columns:
        blockers.append(f"{target.key}: record_id_not_mapped:{target.record_id_field}")
    missing_required = [
        name for name in target.required_fields if name not in target.mapping.field_columns
    ]
    for name in sorted(missing_required):
        blockers.append(f"{target.key}: required_field_not_mapped:{name}")
    if target.mapping.max_rows is not None and target.mapping.max_rows == 0:
        blockers.append(f"{target.key}: template_accepts_no_rows")
    return blockers


def _predicted_protected_cells(
    template: Path, target: DeliveryTarget, row_count: int
) -> tuple[tuple[str, ...], int]:
    """Predict the mapped cells ``write_template`` will refuse to touch.

    ``write_template`` only discovers a protected cell at write time, which
    would leave the plan promising rows that can never be written.  Reading the
    template's mapped cells here and applying the *same* shared predicates
    (:func:`cells.is_protected_formula_value`, :func:`cells.is_merged_non_anchor`)
    gives the plan the writer's own verdict ahead of time.

    Returns the predicted cell coordinates and the number of distinct rows they
    span, because one protected cell fails the whole row closed.
    """

    mapping = target.mapping
    if row_count <= 0 or not template.is_file():
        return (), 0
    try:
        workbook = load_workbook(template)
    except (OSError, ValueError, KeyError):
        # An unreadable template is already reported by ``_target_blockers`` and
        # by the writer; a failed pre-check must not invent a second verdict.
        return (), 0
    try:
        if mapping.sheet not in workbook.sheetnames:
            return (), 0
        worksheet = workbook[mapping.sheet]
        columns: list[int] = []
        for column in mapping.field_columns.values():
            try:
                columns.append(_column_number(column))
            except DeliveryPlanError:
                continue
        protected: list[str] = []
        affected_rows = 0
        for offset in range(row_count):
            row = mapping.data_start_row + offset
            hit = False
            for index in columns:
                cell = worksheet.cell(row, index)
                if is_merged_non_anchor(cell) or is_protected_formula_value(
                    cell.value, overwrite_formulas=mapping.overwrite_formulas
                ):
                    protected.append(f"{get_column_letter(index)}{row}")
                    hit = True
            if hit:
                affected_rows += 1
        return tuple(protected), affected_rows
    finally:
        workbook.close()


def _build_plan(
    inputs: Sequence[str | Path],
    outcomes: Sequence[RecordOutcome],
    targets: Sequence[DeliveryTarget],
    staging: Path,
    delivery: Path,
    batch: ConfirmationBatch,
) -> tuple[DeliveryPlan, dict[str, set[str]]]:
    planned: list[PlannedTarget] = []
    blocking: list[str] = []
    existing: dict[str, set[str]] = {}
    for target in targets:
        template = Path(target.template_path)
        accepted = [item for item in outcomes if item.destination_key == target.key]
        withheld = [
            item
            for item in outcomes
            if item.status == REVIEW and target.key in item.candidates
        ]
        delivery_path = _delivery_path(target, delivery)
        known = _existing_record_ids(delivery_path, target)
        existing[target.key] = known
        skipped = [item for item in accepted if item.record_id in known]
        blockers = _target_blockers(target)
        blocking.extend(blockers)
        writable_rows = len(accepted) - len(skipped)
        protected_cells, protected_rows = _predicted_protected_cells(
            template, target, writable_rows
        )
        planned.append(
            PlannedTarget(
                destination_key=target.key,
                template_path=str(template),
                sheet=target.mapping.sheet,
                field_mapping=dict(target.mapping.field_columns),
                staged_path=str(
                    _unique_path(staging / f"{template.stem}-{target.key}{template.suffix}")
                ),
                delivery_path=str(delivery_path),
                expected_written=writable_rows - protected_rows,
                expected_review=len(withheld),
                expected_skipped_existing=len(skipped),
                expected_skipped_protected=protected_rows,
                protected_cells=protected_cells,
                blocking_items=tuple(blockers),
            )
        )
    unresolved = batch.unresolved_blockers
    plan = DeliveryPlan(
        tuple(str(item) for item in inputs),
        _counts(outcomes),
        tuple(planned),
        unresolved,
        tuple(blocking),
    )
    return plan, existing


def _mark_existing(
    outcomes: Sequence[RecordOutcome], existing: Mapping[str, set[str]]
) -> list[RecordOutcome]:
    """Withhold records the target already holds instead of appending them twice."""

    updated: list[RecordOutcome] = []
    for item in outcomes:
        known = existing.get(item.destination_key or "", set())
        if item.status == ACCEPTED and item.record_id in known:
            updated.append(
                replace(
                    item,
                    status=SKIPPED_EXISTING,
                    reasons=tuple(dict.fromkeys((*item.reasons, "already_in_target"))),
                )
            )
        else:
            updated.append(item)
    return updated


def _planned_outcome(target: DeliveryTarget, plan: DeliveryPlan) -> TargetOutcome:
    planned = next(item for item in plan.targets if item.destination_key == target.key)
    return TargetOutcome(
        target.key,
        planned.template_path,
        False,
        status="planned",
    )


def _contract(target: DeliveryTarget, accepted: Sequence[RecordOutcome], withheld: int) -> DeliveryContract:
    verifiers = (target.formula_verifier,) if target.formula_verifier is not None else ()
    return DeliveryContract(
        counts=StageCounts(
            input=len(accepted) + withheld,
            accepted=len(accepted),
            review=withheld,
            written=len(accepted),
        ),
        required_fields=target.required_fields,
        expected_record_ids=tuple(item.record_id for item in accepted),
        record_id_field=target.record_id_field,
        period_expectations=target.period_expectations,
        verifiers=verifiers,
    )


def _write_and_verify(
    outcomes: Sequence[RecordOutcome],
    targets: Sequence[DeliveryTarget],
    plan: DeliveryPlan,
    post_stage_hook: Callable[[Path], None] | None,
) -> tuple[list[RecordOutcome], list[TargetOutcome], list[DeliveryFailure]]:
    # Keyed by position: a duplicate shares its stable record ID with the original.
    final = list(outcomes)
    target_outcomes: list[TargetOutcome] = []
    failures: list[DeliveryFailure] = []

    for target in targets:
        planned = next(item for item in plan.targets if item.destination_key == target.key)
        positions = [
            index
            for index, item in enumerate(outcomes)
            if item.status == ACCEPTED and item.destination_key == target.key
        ]
        accepted = [outcomes[index] for index in positions]
        withheld = planned.expected_review
        if not accepted:
            target_outcomes.append(
                TargetOutcome(
                    target.key,
                    planned.template_path,
                    False,
                    delivery_path=planned.delivery_path if Path(planned.delivery_path).is_file() else None,
                    status="no_new_records",
                )
            )
            continue

        rows = [dict(item.values) for item in accepted]
        try:
            write_result = write_template(
                target.template_path, rows, target.mapping, output_path=planned.staged_path
            )
        except TemplateWriteError as error:
            failures.append(
                DeliveryFailure(
                    "template_write_failed",
                    f"{target.key}: {error}",
                    "Correct the template mapping or the staging path, then re-run.",
                    target.key,
                )
            )
            target_outcomes.append(
                TargetOutcome(target.key, planned.template_path, False, status="write_failed")
            )
            continue

        staged = Path(write_result.output_path)
        if write_result.skipped:
            # A mapped cell the writer refused to touch means an accepted record
            # was not fully written. Fail closed instead of delivering a partial row.
            reasons = sorted({item.reason for item in write_result.skipped})
            failures.append(
                DeliveryFailure(
                    "incomplete_write",
                    f"{target.key}: the writer skipped {len(write_result.skipped)} mapped cell(s): {', '.join(reasons)}.",
                    "Adjust the template mapping or unprotect the declared region; never deliver a partially written row.",
                    target.key,
                )
            )
            target_outcomes.append(
                TargetOutcome(
                    target.key,
                    planned.template_path,
                    False,
                    str(staged),
                    None,
                    str(write_result.change_log_path),
                    None,
                    tuple(item.record_id for item in accepted),
                    len(write_result.changes),
                    tuple(asdict(item) for item in write_result.skipped),
                    (),
                    "incomplete_write",
                )
            )
            continue

        if post_stage_hook is not None:
            post_stage_hook(staged)

        contract = _contract(target, accepted, withheld)
        try:
            verification = verify_and_deliver(write_result, planned.delivery_path, contract)
        except DeliveryVerificationError as error:
            failures.append(
                DeliveryFailure(
                    "verification_failed",
                    f"{target.key}: {error}",
                    "Read the verification report, fix the cause, and deliver again.",
                    target.key,
                    error.result.findings,
                )
            )
            target_outcomes.append(
                TargetOutcome(
                    target.key,
                    planned.template_path,
                    False,
                    str(staged),
                    None,
                    str(write_result.change_log_path),
                    str(error.result.report_path),
                    tuple(item.record_id for item in accepted),
                    len(write_result.changes),
                    tuple(asdict(item) for item in write_result.skipped),
                    error.result.findings,
                    "verification_failed",
                )
            )
            continue

        traced = _trace(accepted, write_result)
        for offset, position in enumerate(positions):
            final[position] = replace(
                final[position], status=WRITTEN, cells=traced[offset]
            )
        target_outcomes.append(
            TargetOutcome(
                target.key,
                planned.template_path,
                True,
                str(staged),
                str(verification.delivery_path),
                str(write_result.change_log_path),
                str(verification.report_path),
                tuple(item.record_id for item in accepted),
                len(write_result.changes),
                tuple(asdict(item) for item in write_result.skipped),
                verification.findings,
                "delivered",
            )
        )

    return final, target_outcomes, failures


def _trace(
    accepted: Sequence[RecordOutcome], write_result: TemplateWriteResult
) -> dict[int, tuple[WrittenCell, ...]]:
    """Map each accepted row's position to the delivered cells it produced."""

    traced: dict[int, list[WrittenCell]] = {index: [] for index in range(len(accepted))}
    for change in write_result.changes:
        record = accepted[change.row_index]
        traced[change.row_index].append(
            WrittenCell(
                record.record_id,
                change.sheet,
                change.cell,
                change.field,
                change.new_value,
                record.source,
            )
        )
    return {key: tuple(value) for key, value in traced.items()}


def load_delivery_targets(
    payload: Mapping[str, Any], *, base_dir: str | Path = "."
) -> tuple[list[DeliveryTarget], dict[str, Any]]:
    """Build targets and run options from a declarative JSON configuration."""

    base = Path(base_dir)
    raw_targets = payload.get("targets")
    if not isinstance(raw_targets, list) or not raw_targets:
        raise DeliveryPlanError("configuration must contain a non-empty targets array")
    targets: list[DeliveryTarget] = []
    for raw in raw_targets:
        if not isinstance(raw, Mapping):
            raise DeliveryPlanError("each target must be an object")
        key = str(raw.get("key") or "").strip()
        if not key:
            raise DeliveryPlanError("each target needs a key")
        field_columns = raw.get("field_columns")
        if not isinstance(field_columns, Mapping) or not field_columns:
            raise DeliveryPlanError(f"target {key} needs field_columns")
        mapping = TemplateMapping(
            sheet=str(raw.get("sheet") or ""),
            field_columns=dict(field_columns),
            header_row=int(raw.get("header_row", 1)),
            data_start_row=int(raw.get("data_start_row", 2)),
            template_type=str(raw.get("template_type", "table")),
            style_source_row=raw.get("style_source_row"),
            max_rows=raw.get("max_rows"),
        )
        expectations = tuple(
            PeriodExpectation(
                str(item.get("sheet") or mapping.sheet),
                str(item.get("cell")),
                item.get("expected"),
                str(item.get("field", "report_period")),
            )
            for item in raw.get("period_expectations", ())
            if isinstance(item, Mapping)
        )
        targets.append(
            DeliveryTarget(
                destination=Destination(key, tuple(str(item) for item in raw.get("aliases", ()))),
                template_path=base / str(raw.get("template") or ""),
                mapping=mapping,
                required_fields=tuple(str(item) for item in raw.get("required_fields", ())),
                record_id_field=str(raw.get("record_id_field", "record_id")),
                period_expectations=expectations,
                delivery_name=raw.get("delivery_name"),
            )
        )
    inputs = payload.get("inputs")
    if not isinstance(inputs, list) or not inputs:
        raise DeliveryPlanError("configuration must contain a non-empty inputs array")
    options: dict[str, Any] = {
        "inputs": [base / str(item) for item in inputs],
        "staging_dir": base / str(payload.get("staging_dir") or "staging"),
        "delivery_dir": base / str(payload.get("delivery_dir") or "delivery"),
        "confidence_threshold": float(payload.get("confidence_threshold", 0.85)),
    }
    if payload.get("recipe"):
        options["recipe_path"] = base / str(payload["recipe"])
    return targets, options
