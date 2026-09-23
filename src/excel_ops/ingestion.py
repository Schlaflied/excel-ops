from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Sequence

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .extraction import load_extracted_json
from .models import ExtractedRecord


class LayoutDetectionError(ValueError):
    """Raised when an input cannot be mapped without guessing."""


FIELD_ALIASES = {
    "location": {"location", "site", "facility", "address"},
    "event_date": {
        "event date", "date", "inspection date", "date time", "datetime",
        "timestamp", "period", "cycle date", "reporting period",
    },
    "identifier": {"identifier", "id", "asset id", "record id", "code"},
    "category": {"category", "type", "inspection type", "event type"},
    "confidence": {"confidence", "score"},
    "amount": {"amount", "gross pay", "pay", "price", "cost", "total"},
    "currency": {"currency", "currency code", "iso currency"},
}
REQUIRED_FIELDS = ("location", "event_date", "identifier", "category")


def _header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _field_map(row: Sequence[Any]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for index, value in enumerate(row):
        normalized = _header(value)
        for field, aliases in FIELD_ALIASES.items():
            if normalized in aliases and field not in mapped:
                mapped[field] = index
    return mapped


def _normalize_date(value: Any, *, epoch=None) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch=epoch) if epoch is not None else from_excel(value)
            return converted.date().isoformat() if isinstance(converted, datetime) else converted.isoformat()
        except (TypeError, ValueError, OverflowError):
            return str(value)
    text = str(value or "").strip()
    if not text:
        return ""
    normalized = text.replace("/", "-")
    try:
        return datetime.fromisoformat(normalized).date().isoformat()
    except ValueError:
        for pattern in ("%m-%d-%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(normalized, pattern).date().isoformat()
            except ValueError:
                continue
    return text


def _records_from_rows(
    rows: Sequence[Sequence[Any]],
    source_file: str,
    source_sheet: str = "",
    *,
    epoch=None,
    header_limit: int = 20,
) -> tuple[list[ExtractedRecord], int]:
    candidates: list[tuple[int, int, dict[str, int]]] = []
    for index, row in enumerate(rows[:header_limit]):
        mapping = _field_map(row)
        required_count = sum(field in mapping for field in REQUIRED_FIELDS)
        if required_count == len(REQUIRED_FIELDS):
            candidates.append((len(mapping), index, mapping))
    if not candidates:
        raise LayoutDetectionError(f"Unable to identify a complete header in {source_file}")
    _, header_index, mapping = max(candidates, key=lambda item: (item[0], -item[1]))

    records: list[ExtractedRecord] = []
    for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not any(value not in (None, "") for value in row):
            continue
        values = {field: row[column] if column < len(row) else None for field, column in mapping.items()}
        confidence = values.get("confidence", 1.0)
        records.append(ExtractedRecord.from_dict({
            "location": values.get("location"),
            "event_date": _normalize_date(values.get("event_date"), epoch=epoch),
            "identifier": values.get("identifier"),
            "category": values.get("category"),
            "confidence": confidence,
            "amount": values.get("amount"),
            "currency": values.get("currency"),
            "source_file": source_file,
            "source_sheet": source_sheet,
            "source_row": row_index,
        }, source_file))
    return records, len(mapping)


def load_csv_records(path: str | Path) -> list[ExtractedRecord]:
    source = Path(path)
    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    records, _ = _records_from_rows(rows, source.name)
    return records


@dataclass(frozen=True)
class _SheetCandidate:
    records: list[ExtractedRecord]
    score: tuple[int, int]
    sheet_index: int


def load_xlsx_records(path: str | Path) -> list[ExtractedRecord]:
    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    candidates: list[_SheetCandidate] = []
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets):
            rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
            try:
                records, mapped_fields = _records_from_rows(
                    rows, source.name, worksheet.title, epoch=workbook.epoch
                )
            except LayoutDetectionError:
                continue
            candidates.append(_SheetCandidate(records, (mapped_fields, len(records)), sheet_index))
    finally:
        workbook.close()
    if not candidates:
        raise LayoutDetectionError(f"No worksheet in {source.name} has a recognized layout")
    return max(candidates, key=lambda item: (item.score, -item.sheet_index)).records


def load_input_records(path: str | Path) -> list[ExtractedRecord]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".json":
        return load_extracted_json(source)
    if suffix == ".csv":
        return load_csv_records(source)
    if suffix in {".xlsx", ".xlsm"}:
        return load_xlsx_records(source)
    raise ValueError(f"Unsupported input format: {suffix or '(none)'}")
