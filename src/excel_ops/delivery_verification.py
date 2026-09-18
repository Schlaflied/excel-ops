"""Independent, fail-closed verification for staged workbook deliveries."""

from __future__ import annotations

import json
import shutil
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Iterable, Protocol

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .cells import column_number
from .template_writer import TemplateWriteResult


@dataclass(frozen=True)
class StageCounts:
    input: int
    accepted: int
    review: int
    written: int


@dataclass(frozen=True)
class PeriodExpectation:
    sheet: str
    cell: str
    expected: Any
    field: str = "report_period"


@dataclass(frozen=True)
class VerificationFinding:
    code: str
    message: str
    suggestion: str
    sheet: str | None = None
    row: int | None = None
    field: str | None = None
    cell: str | None = None
    severity: str = "error"


class WorkbookVerifier(Protocol):
    """Extension check run against each independently reopened workbook."""

    def __call__(
        self,
        workbook: Any,
        write_result: TemplateWriteResult,
        contract: "DeliveryContract",
    ) -> Iterable[VerificationFinding]: ...


@dataclass(frozen=True)
class DeliveryContract:
    counts: StageCounts
    required_fields: tuple[str, ...]
    expected_record_ids: tuple[str, ...] = field(default_factory=tuple)
    record_id_field: str | None = None
    period_expectations: tuple[PeriodExpectation, ...] = field(default_factory=tuple)
    empty_marker_field: str | None = None
    empty_marker_value: Any = "No records"
    verifiers: tuple[WorkbookVerifier, ...] = field(default_factory=tuple)


@dataclass(frozen=True)
class DeliveryVerificationResult:
    staged_path: Path
    delivery_path: Path | None
    report_path: Path
    passed: bool
    counts: StageCounts
    findings: tuple[VerificationFinding, ...] = field(default_factory=tuple)

    def to_dict(self) -> dict[str, Any]:
        return {
            "staged_path": str(self.staged_path),
            "delivery_path": str(self.delivery_path) if self.delivery_path else None,
            "report_path": str(self.report_path),
            "passed": self.passed,
            "counts": asdict(self.counts),
            "findings": [asdict(item) for item in self.findings],
        }


class DeliveryVerificationError(ValueError):
    def __init__(self, result: DeliveryVerificationResult):
        self.result = result
        super().__init__(f"delivery verification failed with {len(result.findings)} finding(s); see {result.report_path}")


def verify_and_deliver(
    write_result: TemplateWriteResult,
    delivery_path: str | Path,
    contract: DeliveryContract,
    *,
    report_path: str | Path | None = None,
) -> DeliveryVerificationResult:
    """Reopen a writer's real output and publish it only after all checks pass."""

    staged = Path(write_result.output_path).resolve()
    destination = Path(delivery_path).resolve()
    report = Path(report_path).resolve() if report_path else staged.with_suffix(staged.suffix + ".verification.json")
    findings = _verify(staged, write_result, contract)
    if _has_blocking_findings(findings):
        result = DeliveryVerificationResult(staged, None, report, False, contract.counts, tuple(findings))
        _write_report(result)
        raise DeliveryVerificationError(result)

    if staged == destination:
        finding = VerificationFinding(
            "delivery_not_separate",
            "The staged workbook and delivery path are the same file.",
            "Write into a staging directory, then publish to a separate delivery directory.",
        )
        result = DeliveryVerificationResult(staged, None, report, False, contract.counts, (finding,))
        _write_report(result)
        raise DeliveryVerificationError(result)

    destination.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(dir=destination.parent, suffix=destination.suffix, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        shutil.copy2(staged, temporary)
        # Verify the persisted copy, not only the writer's staging artifact.
        delivery_findings = _verify(temporary, write_result, contract)
        if _has_blocking_findings(delivery_findings):
            result = DeliveryVerificationResult(staged, None, report, False, contract.counts, tuple(delivery_findings))
            _write_report(result)
            raise DeliveryVerificationError(result)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)

    result = DeliveryVerificationResult(staged, destination, report, True, contract.counts, tuple(delivery_findings))
    _write_report(result)
    return result


def _verify(path: Path, write_result: TemplateWriteResult, contract: DeliveryContract) -> list[VerificationFinding]:
    findings: list[VerificationFinding] = []
    if not path.is_file():
        return [VerificationFinding("missing_output", "The staged workbook does not exist.", "Use writer.output_path, not the template path.")]
    counts = contract.counts
    if min(counts.input, counts.accepted, counts.review, counts.written) < 0:
        findings.append(VerificationFinding("invalid_count", "Stage counts cannot be negative.", "Recompute the pipeline stage counts."))
    if counts.input != counts.accepted + counts.review:
        findings.append(VerificationFinding("stage_count_mismatch", f"input={counts.input}, accepted={counts.accepted}, review={counts.review}.", "Ensure every input is routed exactly once to accepted or review."))
    if counts.written != counts.accepted:
        findings.append(VerificationFinding("written_count_mismatch", f"accepted={counts.accepted}, written={counts.written}.", "Write every accepted record once and do not write review records."))

    keep_vba = path.suffix.lower() == ".xlsm"
    try:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_vba=keep_vba)
    except Exception as exc:
        findings.append(VerificationFinding("unreadable_workbook", str(exc), "Regenerate a readable XLSX/XLSM workbook."))
        return findings
    try:
        mapping = write_result.mapping
        if mapping.sheet not in workbook.sheetnames:
            findings.append(VerificationFinding("missing_sheet", f"Required sheet {mapping.sheet!r} is missing.", "Restore the mapped worksheet before delivery.", sheet=mapping.sheet))
            return findings
        sheet = workbook[mapping.sheet]
        columns = {name: _column_number(column) for name, column in mapping.field_columns.items()}
        template = load_workbook(write_result.template_path, read_only=True, data_only=False, keep_vba=keep_vba)
        try:
            template_sheet = template[mapping.sheet]
            for field_name, column in columns.items():
                expected_header = template_sheet.cell(mapping.header_row, column).value
                actual_header = sheet.cell(mapping.header_row, column).value
                if actual_header != expected_header:
                    findings.append(VerificationFinding("column_shift", f"Expected template header {expected_header!r}, got {actual_header!r}.", "Correct the template mapping or restore the expected column order.", mapping.sheet, mapping.header_row, field_name, f"{get_column_letter(column)}{mapping.header_row}"))
        finally:
            template.close()

        expected_rows = counts.written
        for offset in range(expected_rows):
            row = mapping.data_start_row + offset
            for field_name in contract.required_fields:
                column = columns.get(field_name)
                if column is None:
                    findings.append(VerificationFinding("unmapped_required_field", "Required field is absent from the template mapping.", "Add the field to TemplateMapping.field_columns.", mapping.sheet, row, field_name))
                    continue
                value = sheet.cell(row, column).value
                if value is None or (isinstance(value, str) and not value.strip()):
                    findings.append(VerificationFinding("missing_required_field", "A written record is missing a required value.", "Supply the field or keep the record in review.", mapping.sheet, row, field_name, f"{get_column_letter(column)}{row}"))

        if contract.record_id_field:
            id_column = columns.get(contract.record_id_field)
            if id_column is None:
                findings.append(VerificationFinding("unmapped_record_id", "The record ID field is not mapped.", "Map record_id_field so idempotency can be checked.", mapping.sheet, field=contract.record_id_field))
            else:
                seen: dict[str, int] = {}
                actual_ids: list[str] = []
                for offset in range(expected_rows):
                    row = mapping.data_start_row + offset
                    raw = sheet.cell(row, id_column).value
                    record_id = str(raw or "").strip()
                    actual_ids.append(record_id)
                    if record_id in seen:
                        findings.append(VerificationFinding("duplicate_record", f"Record ID {record_id!r} is also present at row {seen[record_id]}.", "Write each source record at most once.", mapping.sheet, row, contract.record_id_field, f"{get_column_letter(id_column)}{row}"))
                    else:
                        seen[record_id] = row
                if contract.expected_record_ids and tuple(actual_ids) != contract.expected_record_ids:
                    findings.append(VerificationFinding("record_id_mismatch", "Written record IDs do not match the accepted records in order.", "Re-run matching and write-back from the accepted set.", mapping.sheet, field=contract.record_id_field))

        if expected_rows == 0 and contract.empty_marker_field:
            marker_column = columns.get(contract.empty_marker_field)
            if marker_column is None or sheet.cell(mapping.data_start_row, marker_column).value != contract.empty_marker_value:
                findings.append(VerificationFinding("missing_empty_marker", "A zero-record target is not explicitly marked.", "Write the configured zero-record marker for the current period.", mapping.sheet, mapping.data_start_row, contract.empty_marker_field))

        for expectation in contract.period_expectations:
            if expectation.sheet not in workbook.sheetnames:
                findings.append(VerificationFinding("missing_period_sheet", f"Period sheet {expectation.sheet!r} is missing.", "Restore the declared period target.", expectation.sheet, field=expectation.field, cell=expectation.cell))
                continue
            actual = workbook[expectation.sheet][expectation.cell].value
            if not _same_value(actual, expectation.expected):
                findings.append(VerificationFinding("period_mismatch", f"Expected {expectation.expected!r}, got {actual!r}.", "Refresh this declared date slot to the resolved report period.", expectation.sheet, workbook[expectation.sheet][expectation.cell].row, expectation.field, expectation.cell))

        expected_changes = {(item.cell, item.field): item.new_value for item in write_result.changes}
        for (cell, field_name), expected in expected_changes.items():
            actual = sheet[cell].value
            if not _same_value(actual, expected):
                findings.append(VerificationFinding("written_value_mismatch", f"Expected {expected!r}, got {actual!r}.", "Re-run write-back and verify the writer's real output path.", mapping.sheet, sheet[cell].row, field_name, cell))
        for verifier in contract.verifiers:
            try:
                findings.extend(verifier(workbook, write_result, contract))
            except Exception as exc:
                findings.append(VerificationFinding("verifier_error", str(exc), "Fix the verifier or its declared inputs before delivery."))
    finally:
        workbook.close()
    return findings


def _same_value(left: Any, right: Any) -> bool:
    if isinstance(left, datetime) and isinstance(right, date) and not isinstance(right, datetime):
        return left.date() == right
    if isinstance(right, datetime) and isinstance(left, date) and not isinstance(left, datetime):
        return right.date() == left
    return left == right


def _has_blocking_findings(findings: Iterable[VerificationFinding]) -> bool:
    # Unknown severities fail closed; only an explicit warning is non-blocking.
    return any(item.severity != "warning" for item in findings)


def _column_number(value: str | int) -> int:
    return column_number(value)


def _write_report(result: DeliveryVerificationResult) -> None:
    result.report_path.parent.mkdir(parents=True, exist_ok=True)
    result.report_path.write_text(json.dumps(result.to_dict(), ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
