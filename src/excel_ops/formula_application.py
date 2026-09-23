"""Dry-run-first application of formula plans to workbook copies."""

from __future__ import annotations

import os
import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import range_boundaries

from .cells import is_merged_non_anchor, is_protected_formula_value
from .formula_planning import FormulaPlan


class FormulaApplicationError(ValueError):
    """Raised when a formula plan cannot be applied safely."""


@dataclass(frozen=True)
class FormulaChange:
    cell: str
    mode: str
    action: str = "write"

    def to_dict(self) -> dict[str, str]:
        return asdict(self)


@dataclass(frozen=True)
class FormulaSkip:
    cell: str
    reason: str

    def to_dict(self) -> dict[str, str]:
        return asdict(self)


@dataclass(frozen=True)
class FormulaApplicationResult:
    source_path: Path
    output_path: Path
    sheet: str
    target_range: str
    dry_run: bool
    verified: bool
    changes: tuple[FormulaChange, ...]
    skipped: tuple[FormulaSkip, ...]

    def to_dict(self) -> dict[str, Any]:
        return {
            "source_path": str(self.source_path),
            "output_path": str(self.output_path),
            "sheet": self.sheet,
            "target_range": self.target_range,
            "dry_run": self.dry_run,
            "verified": self.verified,
            "changes": [item.to_dict() for item in self.changes],
            "skipped": [item.to_dict() for item in self.skipped],
        }


def plan_formula_application(
    source_path: str | Path,
    output_path: str | Path,
    plan: FormulaPlan,
    *,
    sheet: str,
    target_range: str,
    overwrite_formulas: bool = False,
) -> FormulaApplicationResult:
    """Return a value-free change ledger without writing a workbook."""

    return _application(
        source_path,
        output_path,
        plan,
        sheet=sheet,
        target_range=target_range,
        overwrite_formulas=overwrite_formulas,
        write=False,
    )


def apply_formula_plan(
    source_path: str | Path,
    output_path: str | Path,
    plan: FormulaPlan,
    *,
    sheet: str,
    target_range: str,
    confirmed: bool = False,
    overwrite_formulas: bool = False,
) -> FormulaApplicationResult:
    """Write a confirmed formula plan to a new workbook and verify the copy."""

    if not confirmed:
        raise FormulaApplicationError(
            "formula application requires confirmed=True after reviewing the dry run"
        )
    return _application(
        source_path,
        output_path,
        plan,
        sheet=sheet,
        target_range=target_range,
        overwrite_formulas=overwrite_formulas,
        write=True,
    )


def _application(
    source_path: str | Path,
    output_path: str | Path,
    plan: FormulaPlan,
    *,
    sheet: str,
    target_range: str,
    overwrite_formulas: bool,
    write: bool,
) -> FormulaApplicationResult:
    source = Path(source_path).resolve()
    output = Path(output_path).resolve()
    if source == output:
        raise FormulaApplicationError("source and output paths must be different")
    if not source.is_file():
        raise FormulaApplicationError("source workbook does not exist")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise FormulaApplicationError("source workbook must be .xlsx or .xlsm")
    if output.suffix.lower() != source.suffix.lower():
        raise FormulaApplicationError("output workbook must preserve the source extension")
    if output.exists():
        raise FormulaApplicationError("output workbook already exists")

    bounds = _target_bounds(target_range)
    min_column, min_row, max_column, max_row = bounds
    if plan.output_mode == "static" and (min_column, min_row) != (max_column, max_row):
        raise FormulaApplicationError("static mode requires a single target cell")

    workbook = load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        if sheet not in workbook.sheetnames:
            raise FormulaApplicationError(f"worksheet {sheet!r} does not exist")
        worksheet = workbook[sheet]
        anchor = worksheet.cell(min_row, min_column).coordinate
        changes: list[FormulaChange] = []
        skipped: list[FormulaSkip] = []
        expected: dict[str, Any] = {}
        for row in worksheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_column,
            max_col=max_column,
        ):
            for cell in row:
                if is_merged_non_anchor(cell):
                    skipped.append(FormulaSkip(cell.coordinate, "merged_non_anchor"))
                    continue
                if is_protected_formula_value(
                    cell.value,
                    overwrite_formulas=overwrite_formulas,
                ):
                    skipped.append(FormulaSkip(cell.coordinate, "protected_formula"))
                    continue
                value = _planned_value(plan, anchor, cell.coordinate)
                changes.append(FormulaChange(cell.coordinate, plan.output_mode))
                expected[cell.coordinate] = value
                if write:
                    cell.value = value

        if write:
            output.parent.mkdir(parents=True, exist_ok=True)
            temporary = _temporary_output(output)
            try:
                workbook.save(temporary)
                _verify_written_values(temporary, sheet, expected, source.suffix.lower())
                os.replace(temporary, output)
            finally:
                temporary.unlink(missing_ok=True)
    finally:
        workbook.close()

    return FormulaApplicationResult(
        source,
        output,
        sheet,
        _normalized_range(bounds),
        not write,
        write,
        tuple(changes),
        tuple(skipped),
    )


def _planned_value(plan: FormulaPlan, anchor: str, target: str) -> Any:
    if plan.output_mode == "static":
        return plan.value
    if not isinstance(plan.value, str) or not plan.value.startswith("="):
        raise FormulaApplicationError("formula mode requires a formula value")
    return Translator(plan.value, origin=anchor).translate_formula(target)


def _target_bounds(target_range: str) -> tuple[int, int, int, int]:
    try:
        bounds = range_boundaries(target_range.strip().upper())
    except (TypeError, ValueError) as error:
        raise FormulaApplicationError("target_range must be one bounded A1 range") from error
    if any(value is None for value in bounds):
        raise FormulaApplicationError("target_range must be one bounded A1 range")
    return bounds


def _normalized_range(bounds: tuple[int, int, int, int]) -> str:
    from openpyxl.utils import get_column_letter

    min_column, min_row, max_column, max_row = bounds
    start = f"{get_column_letter(min_column)}{min_row}"
    end = f"{get_column_letter(max_column)}{max_row}"
    return start if start == end else f"{start}:{end}"


def _temporary_output(output: Path) -> Path:
    descriptor, name = tempfile.mkstemp(
        prefix=f".{output.stem}-",
        suffix=output.suffix,
        dir=output.parent,
    )
    os.close(descriptor)
    return Path(name)


def _verify_written_values(
    workbook_path: Path,
    sheet: str,
    expected: dict[str, Any],
    suffix: str,
) -> None:
    workbook = load_workbook(workbook_path, keep_vba=suffix == ".xlsm", data_only=False)
    try:
        worksheet = workbook[sheet]
        mismatches = [cell for cell, value in expected.items() if worksheet[cell].value != value]
    finally:
        workbook.close()
    if mismatches:
        raise FormulaApplicationError(
            "written workbook failed read-back verification at " + ", ".join(mismatches)
        )
