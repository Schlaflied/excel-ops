from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .inference import infer_value


def _canonical(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.casefold())


@dataclass(frozen=True)
class FieldProfile:
    name: str
    position: int
    types: tuple[str, ...]
    required: bool
    distinct_values: int
    sample_values: tuple[str, ...]
    enum_values: tuple[str, ...]


@dataclass(frozen=True)
class SheetProfile:
    name: str
    row_count: int
    fields: tuple[FieldProfile, ...]


def inspect_workbook(path: str | Path) -> dict[str, Any]:
    """Return a JSON-safe structural profile without changing the workbook."""
    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheets: list[SheetProfile] = []
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, ())
            names = [str(value).strip() if value is not None else "" for value in header]
            columns: list[list[Any]] = [[] for _ in names]
            row_count = 0
            for row in rows:
                if not any(value not in (None, "") for value in row):
                    continue
                row_count += 1
                for index in range(len(names)):
                    columns[index].append(row[index] if index < len(row) else None)

            fields = []
            for index, (name, values) in enumerate(zip(names, columns, strict=True), start=1):
                if not name:
                    continue
                populated = [value for value in values if value not in (None, "")]
                types = tuple(sorted({infer_value(name, value).inferred_type for value in populated})) or ("empty",)
                samples = tuple(dict.fromkeys(str(value) for value in populated[:20]))[:5]
                distinct = tuple(dict.fromkeys(str(value) for value in populated))
                fields.append(
                    FieldProfile(
                        name=name,
                        position=index,
                        types=types,
                        required=bool(values) and len(populated) == len(values),
                        distinct_values=len(distinct),
                        sample_values=samples,
                        enum_values=tuple(sorted(distinct)) if len(distinct) <= 20 else (),
                    )
                )
            sheets.append(SheetProfile(worksheet.title, row_count, tuple(fields)))
        return {"path": str(source), "sheets": [asdict(sheet) for sheet in sheets]}
    finally:
        workbook.close()


def load_confirmed_mappings(path: str | Path | None) -> dict[str, str]:
    if path is None:
        return {}
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    mappings = payload.get("confirmed_mappings", payload) if isinstance(payload, dict) else None
    if not isinstance(mappings, dict) or not all(isinstance(k, str) and isinstance(v, str) for k, v in mappings.items()):
        raise ValueError("Mapping file must contain an object named confirmed_mappings")
    return {_canonical(source): target for source, target in mappings.items()}


def save_confirmed_mapping(path: str | Path, source: str, target: str) -> Path:
    """Persist a human-confirmed candidate-to-baseline field mapping."""
    destination = Path(path)
    payload: dict[str, Any] = {"confirmed_mappings": {}}
    if destination.exists():
        loaded = json.loads(destination.read_text(encoding="utf-8"))
        if not isinstance(loaded, dict) or not isinstance(loaded.get("confirmed_mappings", {}), dict):
            raise ValueError("Mapping file must contain an object named confirmed_mappings")
        payload = loaded
        payload.setdefault("confirmed_mappings", {})
    payload["confirmed_mappings"][source] = target
    destination.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return destination


def _field_map(sheet: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {_canonical(field["name"]): field for field in sheet["fields"]}


def _semantic_tokens(name: str) -> set[str]:
    aliases = {
        "staff": "employee",
        "personnel": "employee",
        "worker": "employee",
        "number": "id",
        "no": "id",
        "code": "id",
    }
    return {aliases.get(token, token) for token in re.findall(r"[a-z0-9]+", name.casefold())}


def _suggestions(removed: Iterable[dict[str, Any]], added: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    suggestions = []
    for old in removed:
        old_tokens = _semantic_tokens(old["name"])
        for new in added:
            new_tokens = _semantic_tokens(new["name"])
            overlap = len(old_tokens & new_tokens) / max(len(old_tokens | new_tokens), 1)
            same_types = bool(set(old["types"]) & set(new["types"]))
            if overlap >= 0.5 or (same_types and old["position"] == new["position"]):
                suggestions.append(
                    {"from": old["name"], "to": new["name"], "status": "review", "reason": "possible_rename"}
                )
    return suggestions


def compare_workbooks(
    paths: Iterable[str | Path],
    *,
    mapping_path: str | Path | None = None,
    key_field: str | None = None,
) -> dict[str, Any]:
    """Compare two or more workbooks against the first workbook as baseline."""
    profiles = [inspect_workbook(path) for path in paths]
    if len(profiles) < 2:
        raise ValueError("At least two workbooks are required")
    confirmed = load_confirmed_mappings(mapping_path)
    baseline = profiles[0]
    comparisons = []

    baseline_sheets = {sheet["name"]: sheet for sheet in baseline["sheets"]}
    for candidate in profiles[1:]:
        candidate_sheets = {sheet["name"]: sheet for sheet in candidate["sheets"]}
        changes: list[dict[str, Any]] = []
        reviews: list[dict[str, Any]] = []

        for name in sorted(baseline_sheets.keys() - candidate_sheets.keys()):
            changes.append({"kind": "sheet_removed", "sheet": name})
        for name in sorted(candidate_sheets.keys() - baseline_sheets.keys()):
            changes.append({"kind": "sheet_added", "sheet": name})

        for name in sorted(baseline_sheets.keys() & candidate_sheets.keys()):
            old_sheet, new_sheet = baseline_sheets[name], candidate_sheets[name]
            old_fields, new_fields = _field_map(old_sheet), _field_map(new_sheet)
            removed = [field for key, field in old_fields.items() if key not in new_fields]
            added = [field for key, field in new_fields.items() if key not in old_fields]

            for new_key, new_field in list(new_fields.items()):
                target = confirmed.get(new_key)
                old_key = _canonical(target) if target else ""
                if old_key in old_fields and new_key not in old_fields:
                    changes.append(
                        {"kind": "field_renamed", "sheet": name, "from": old_fields[old_key]["name"], "to": new_field["name"], "mapping": "confirmed"}
                    )
                    removed = [field for field in removed if _canonical(field["name"]) != old_key]
                    added = [field for field in added if _canonical(field["name"]) != new_key]

            for key in sorted(old_fields.keys() & new_fields.keys()):
                old, new = old_fields[key], new_fields[key]
                if old["position"] != new["position"]:
                    changes.append({"kind": "field_reordered", "sheet": name, "field": new["name"], "from": old["position"], "to": new["position"]})
                if old["types"] != new["types"]:
                    changes.append({"kind": "field_type_changed", "sheet": name, "field": new["name"], "from": old["types"], "to": new["types"]})
                if old["required"] != new["required"]:
                    changes.append({"kind": "field_required_changed", "sheet": name, "field": new["name"], "from": old["required"], "to": new["required"]})
                if "text" in old["types"] and "text" in new["types"] and old["distinct_values"] <= 20 and new["distinct_values"] <= 20:
                    old_values, new_values = set(old["enum_values"]), set(new["enum_values"])
                    if old_values != new_values:
                        changes.append(
                            {
                                "kind": "field_enum_changed",
                                "sheet": name,
                                "field": new["name"],
                                "added": sorted(new_values - old_values),
                                "removed": sorted(old_values - new_values),
                            }
                        )

            changes.extend({"kind": "field_removed", "sheet": name, "field": field["name"]} for field in removed)
            changes.extend({"kind": "field_added", "sheet": name, "field": field["name"]} for field in added)
            reviews.extend({**item, "sheet": name} for item in _suggestions(removed, added))

            if key_field:
                old_key, new_key = old_fields.get(_canonical(key_field)), new_fields.get(_canonical(key_field))
                if old_key and new_key:
                    old_ratio = old_sheet["row_count"] / max(old_key["distinct_values"], 1)
                    new_ratio = new_sheet["row_count"] / max(new_key["distinct_values"], 1)
                    if abs(old_ratio - new_ratio) > 0.01:
                        changes.append(
                            {"kind": "grain_changed", "sheet": name, "key_field": key_field, "rows_per_key": {"from": old_ratio, "to": new_ratio}}
                        )

        blocked = bool(reviews or any(change["kind"] in {"sheet_added", "sheet_removed", "field_added", "field_removed", "grain_changed"} for change in changes))
        comparisons.append(
            {
                "candidate": candidate["path"],
                "changes": changes,
                "mapping_suggestions": reviews,
                "impact": {
                    "requires_review": blocked,
                    "append": "blocked" if blocked else "compatible",
                    "join": "blocked" if blocked else "compatible",
                    "delivery_grouping": "review" if any(c["kind"] == "grain_changed" for c in changes) else "unchanged",
                },
            }
        )

    return {
        "report_version": 1,
        "baseline": baseline["path"],
        "confirmed_mapping_file": str(mapping_path) if mapping_path else None,
        "comparisons": comparisons,
    }


def write_drift_report(
    paths: Iterable[str | Path],
    output_path: str | Path,
    *,
    mapping_path: str | Path | None = None,
    key_field: str | None = None,
) -> Path:
    output = Path(output_path)
    report = compare_workbooks(paths, mapping_path=mapping_path, key_field=key_field)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return output
