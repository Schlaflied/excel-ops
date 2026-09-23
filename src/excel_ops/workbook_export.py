"""Deterministic XLSX and CSV export helpers."""

from __future__ import annotations

import csv
import shutil
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string


class WorkbookExportError(ValueError):
    """Raised when an export request cannot be completed safely."""


def export_xlsx(source: str | Path, destination: str | Path) -> Path:
    """Copy an existing workbook to a new XLSX destination without changing it."""

    source_path = Path(source)
    destination_path = Path(destination)
    if source_path.suffix.lower() != ".xlsx":
        raise WorkbookExportError("XLSX export requires an XLSX source")
    if destination_path.suffix.lower() != ".xlsx":
        raise WorkbookExportError("XLSX export destination must end in .xlsx")
    if not source_path.is_file():
        raise WorkbookExportError(f"source workbook does not exist: {source_path}")
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, destination_path)
    return destination_path


def export_csv(
    source: str | Path,
    destination: str | Path,
    *,
    sheet: str | None = None,
    sheets: Iterable[str] | None = None,
) -> tuple[Path, ...]:
    """Export one workbook sheet to UTF-8 CSV, or one CSV file per sheet."""

    source_path = Path(source)
    destination_path = Path(destination)
    if source_path.suffix.lower() != ".xlsx":
        raise WorkbookExportError("CSV export requires an XLSX source")
    if destination_path.suffix.lower() != ".csv" and sheet is not None:
        raise WorkbookExportError("single-sheet CSV destination must end in .csv")
    if sheet is not None and sheets is not None:
        raise WorkbookExportError("choose sheet or sheets, not both")
    workbook = load_workbook(source_path, read_only=True, data_only=False)
    values_workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        selected = [sheet] if sheet is not None else list(workbook.sheetnames if sheets is None else sheets)
        if not selected:
            raise WorkbookExportError("CSV export requires at least one sheet")
        missing = [name for name in selected if name not in workbook.sheetnames]
        if missing:
            raise WorkbookExportError(f"unknown worksheet(s): {', '.join(missing)}")
        destination_path.parent.mkdir(parents=True, exist_ok=True)
        outputs: list[Path] = []
        staged: list[tuple[Path, Path]] = []
        used_names: set[str] = set()
        try:
            for name in selected:
                output = destination_path if sheet is not None else destination_path / _portable_sheet_filename(name, used_names)
                if sheet is None and destination_path.suffix:
                    raise WorkbookExportError("multi-sheet CSV destination must be a directory")
                if output.suffix.lower() != ".csv":
                    raise WorkbookExportError("CSV output must end in .csv")
                output.parent.mkdir(parents=True, exist_ok=True)
                handle = tempfile.NamedTemporaryFile(
                    mode="w", encoding="utf-8-sig", newline="", dir=output.parent,
                    prefix=f".{output.stem}-", suffix=".tmp", delete=False,
                )
                temporary = Path(handle.name)
                try:
                    with handle:
                        writer = csv.writer(handle, lineterminator="\n")
                        cached_formula_cells = _cached_formula_cells(source_path, workbook[name])
                        for row in _cached_rows(
                            workbook[name], values_workbook[name], cached_formula_cells
                        ):
                            writer.writerow([_safe_csv_value(value) for value in row])
                    staged.append((temporary, output))
                except Exception:
                    temporary.unlink(missing_ok=True)
                    raise
                outputs.append(output)
            for temporary, output in staged:
                temporary.replace(output)
            return tuple(outputs)
        except Exception:
            for temporary, _ in staged:
                temporary.unlink(missing_ok=True)
            raise
    finally:
        workbook.close()
        values_workbook.close()


def _cached_rows(formula_sheet, values_sheet, cached_formula_cells: set[str]):
    """Yield cached formula results, rejecting formulas without cached values."""

    for formula_row, values_row in zip(
        formula_sheet.iter_rows(values_only=False), values_sheet.iter_rows(values_only=True)
    ):
        values = []
        for formula_cell, cached in zip(formula_row, values_row):
            if (
                formula_cell.data_type == "f"
                and cached is None
                and formula_cell.coordinate not in cached_formula_cells
            ):
                raise WorkbookExportError(
                    f"formula has no cached result: {formula_sheet.title}!{formula_cell.coordinate}"
                )
            values.append(cached)
        yield values


def _portable_sheet_filename(title: str, used: set[str]) -> str:
    """Return a deterministic CSV filename safe on Windows and POSIX."""

    invalid = '<>:"/\\|?*'
    stem = "".join("_" if char in invalid or ord(char) < 32 else char for char in title).strip(" .")
    if not stem:
        stem = "sheet"
    base = stem.split(".", 1)[0].upper()
    reserved = {"CON", "PRN", "AUX", "NUL"}
    reserved.update(f"COM{i}" for i in range(1, 10))
    reserved.update(f"LPT{i}" for i in range(1, 10))
    reserved.update(f"COM{superscript}" for superscript in "¹²³")
    reserved.update(f"LPT{superscript}" for superscript in "¹²³")
    if base in reserved:
        stem = f"_{stem}"
    candidate = f"{stem}.csv"
    index = 2
    while candidate.casefold() in {item.casefold() for item in used}:
        candidate = f"{stem}-{index}.csv"
        index += 1
    used.add(candidate)
    return candidate


def _safe_csv_value(value):
    """Prefix spreadsheet formula triggers in text cells before CSV output."""

    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + value
    return value


def _cached_formula_cells(source: Path, sheet) -> set[str]:
    """Return formula coordinates whose worksheet XML contains a value node."""

    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    cached: set[str] = set()
    sheet_data = None
    current_row = 0
    next_row = 1
    current_column = 0
    with zipfile.ZipFile(source) as archive:
        with archive.open(sheet._worksheet_path) as stream:
            for event, element in ET.iterparse(stream, events=("start", "end")):
                if event == "start" and element.tag == f"{namespace}sheetData":
                    sheet_data = element
                elif event == "start" and element.tag == f"{namespace}row":
                    current_row = int(element.attrib.get("r", next_row))
                    current_column = 0
                elif event == "end" and element.tag == f"{namespace}c":
                    explicit = element.attrib.get("r")
                    if explicit:
                        column_name, row_number = coordinate_from_string(explicit)
                        current_column = column_index_from_string(column_name)
                        coordinate = f"{column_name}{row_number}"
                    else:
                        current_column += 1
                        coordinate = f"{get_column_letter(current_column)}{current_row}"
                    formula = element.find(f"{namespace}f")
                    value = element.find(f"{namespace}v")
                    if (
                        formula is not None
                        and value is not None
                        and (value.text is not None or element.attrib.get("t") == "str")
                    ):
                        cached.add(coordinate)
                    element.clear()
                elif event == "end" and element.tag == f"{namespace}row":
                    next_row = current_row + 1
                    element.clear()
                    if sheet_data is not None:
                        sheet_data.clear()
    return cached
