from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime
from enum import StrEnum
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Iterable

from openpyxl import load_workbook

from .periods import PeriodResult


class DateRole(StrEnum):
    PERIOD = "report_period"
    PERIOD_START = "period_start"
    PERIOD_END = "period_end"
    GENERATED_ON = "generated_on"
    DEADLINE = "deadline"
    HISTORICAL = "historical_event_date"


class SlotTarget(StrEnum):
    CELL = "cell"
    SHEET_TITLE = "sheet_title"


@dataclass(frozen=True)
class DateSlot:
    sheet: str
    role: DateRole
    target: SlotTarget = SlotTarget.CELL
    cell: str | None = None
    template: str = "{value}"
    include_when_empty: bool = True
    data_start_row: int = 2

    def __post_init__(self) -> None:
        if self.target == SlotTarget.CELL and not self.cell:
            raise ValueError("cell target requires a cell address")
        if self.target == SlotTarget.SHEET_TITLE and self.cell:
            raise ValueError("sheet title target cannot include a cell address")


@dataclass(frozen=True)
class RefreshChange:
    sheet: str
    target: str
    role: str
    old_value: Any
    new_value: Any
    status: str


@dataclass(frozen=True)
class RefreshResult:
    input_path: str
    output_path: str | None
    dry_run: bool
    period: dict[str, Any]
    changes: tuple[RefreshChange, ...]

    def to_dict(self) -> dict[str, Any]:
        return {
            "input_path": self.input_path,
            "output_path": self.output_path,
            "dry_run": self.dry_run,
            "period": self.period,
            "changes": [asdict(change) for change in self.changes],
        }


def _sheet_has_records(worksheet: Any, data_start_row: int) -> bool:
    for row in worksheet.iter_rows(min_row=data_start_row):
        if any(cell.value not in (None, "") for cell in row):
            return True
    return False


def _period_context(
    period: PeriodResult,
    *,
    generated_on: date,
    deadline: date | None,
) -> dict[str, Any]:
    return {
        "value": period.display_text,
        "period": period.display_text,
        "period_start": period.period_start,
        "period_end": period.period_end,
        "generated_on": generated_on,
        "deadline": deadline,
    }


def _role_value(role: DateRole, context: dict[str, Any]) -> Any:
    values = {
        DateRole.PERIOD: context["period"],
        DateRole.PERIOD_START: context["period_start"],
        DateRole.PERIOD_END: context["period_end"],
        DateRole.GENERATED_ON: context["generated_on"],
        DateRole.DEADLINE: context["deadline"],
    }
    if role == DateRole.HISTORICAL:
        raise ValueError("historical event dates are protected and cannot be refreshed")
    value = values[role]
    if value is None:
        raise ValueError(f"{role.value} requires an explicit value")
    return value


def render_period_template(
    template: str,
    period: PeriodResult,
    *,
    generated_on: date | None = None,
    deadline: date | None = None,
) -> str:
    generated = generated_on or period.as_of_date
    context = _period_context(period, generated_on=generated, deadline=deadline)
    return template.format(**context)


def _render_slot(slot: DateSlot, context: dict[str, Any]) -> Any:
    value = _role_value(slot.role, context)
    if slot.template == "{value}":
        return value
    render_context = dict(context)
    render_context["value"] = value
    return slot.template.format(**render_context)


def _values_equal(left: Any, right: Any) -> bool:
    if isinstance(left, datetime) and isinstance(right, date) and not isinstance(right, datetime):
        return left.date() == right
    if isinstance(right, datetime) and isinstance(left, date) and not isinstance(left, datetime):
        return right.date() == left
    return left == right


def _apply_slots(workbook: Any, slots: Iterable[DateSlot], context: dict[str, Any]) -> list[RefreshChange]:
    changes: list[RefreshChange] = []
    for slot in slots:
        if slot.sheet not in workbook.sheetnames:
            raise ValueError(f"declared date slot sheet does not exist: {slot.sheet}")
        worksheet = workbook[slot.sheet]
        target = slot.cell or "<sheet title>"
        if slot.role == DateRole.HISTORICAL:
            old_value = worksheet[slot.cell].value if slot.cell else worksheet.title
            changes.append(RefreshChange(slot.sheet, target, slot.role.value, old_value, old_value, "protected"))
            continue
        if not slot.include_when_empty and not _sheet_has_records(worksheet, slot.data_start_row):
            old_value = worksheet[slot.cell].value if slot.cell else worksheet.title
            changes.append(RefreshChange(slot.sheet, target, slot.role.value, old_value, old_value, "empty_skipped"))
            continue
        new_value = _render_slot(slot, context)
        if slot.target == SlotTarget.CELL:
            old_value = worksheet[slot.cell].value
            if not _values_equal(old_value, new_value):
                worksheet[slot.cell] = new_value
        else:
            old_value = worksheet.title
            if old_value != str(new_value):
                worksheet.title = str(new_value)
        changes.append(
            RefreshChange(slot.sheet, target, slot.role.value, old_value, new_value, "unchanged" if _values_equal(old_value, new_value) else "updated")
        )
    return changes


def _verify_slots(path: Path, slots: Iterable[DateSlot], expected: Iterable[RefreshChange]) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        actual_sheets = set(workbook.sheetnames)
        for slot, change in zip(slots, expected, strict=True):
            sheet_name = str(change.new_value) if slot.target == SlotTarget.SHEET_TITLE and change.status not in {"protected", "empty_skipped"} else slot.sheet
            if sheet_name not in actual_sheets:
                raise ValueError(f"date slot verification failed: sheet {sheet_name!r} is missing")
            if slot.target == SlotTarget.CELL:
                actual = workbook[sheet_name][slot.cell].value
                if not _values_equal(actual, change.new_value):
                    raise ValueError(
                        f"date slot verification failed at {sheet_name}!{slot.cell}: expected {change.new_value!r}, got {actual!r}"
                    )
    finally:
        workbook.close()


def refresh_workbook_period(
    input_path: str | Path,
    output_path: str | Path,
    *,
    period: PeriodResult,
    slots: Iterable[DateSlot],
    generated_on: date | None = None,
    deadline: date | None = None,
    dry_run: bool = False,
) -> RefreshResult:
    source = Path(input_path)
    destination = Path(output_path)
    declared_slots = tuple(slots)
    if not declared_slots:
        raise ValueError("at least one declared date slot is required")
    if source.resolve() == destination.resolve():
        raise ValueError("output_path must differ from input_path")
    generated = generated_on or period.as_of_date
    context = _period_context(period, generated_on=generated, deadline=deadline)
    keep_vba = source.suffix.lower() == ".xlsm"
    workbook = load_workbook(source, read_only=False, data_only=False, keep_vba=keep_vba)
    try:
        changes = _apply_slots(workbook, declared_slots, context)
        if not dry_run:
            destination.parent.mkdir(parents=True, exist_ok=True)
            suffix = destination.suffix or source.suffix
            with NamedTemporaryFile(dir=destination.parent, suffix=suffix, delete=False) as handle:
                temporary = Path(handle.name)
            try:
                workbook.save(temporary)
                temporary.replace(destination)
            finally:
                temporary.unlink(missing_ok=True)
    finally:
        workbook.close()
    if not dry_run:
        _verify_slots(destination, declared_slots, changes)
    return RefreshResult(
        input_path=str(source),
        output_path=None if dry_run else str(destination),
        dry_run=dry_run,
        period=period.to_dict(),
        changes=tuple(changes),
    )
