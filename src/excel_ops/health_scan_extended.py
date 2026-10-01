"""Conservative workbook-health detectors that require an explicit repair baseline."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from hashlib import sha256
import json
from typing import Any, Mapping

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.utils.cell import range_to_tuple

from .health_models import HealthFinding

BLANK_ZERO_CELL_BUDGET = 100_000


def _finding(
    code: str,
    sheet: str,
    location: str,
    evidence: str,
    suggestion: str,
    *,
    details: Mapping[str, Any] | None = None,
    fixability: str = "review",
) -> HealthFinding:
    return HealthFinding(
        code=code,
        severity="warning",
        sheet=sheet,
        location=location,
        evidence=evidence,
        suggestion=suggestion,
        fixability=fixability,
        details=dict(details or {}),
    )


def _data_cells(sheet: Any, header_row: int | None, max_row: int | None, max_column: int | None) -> list[Any]:
    if header_row is None or max_row is None or max_column is None:
        return []
    return [
        cell
        for cell in getattr(sheet, "_cells", {}).values()
        if not isinstance(cell, MergedCell)
        and header_row < cell.row <= max_row
        and cell.column <= max_column
        and cell.value not in (None, "")
    ]


def _column_location(column: int, header_row: int, max_row: int) -> str:
    letter = get_column_letter(column)
    return f"{letter}{header_row + 1}:{letter}{max_row}"


def _mixed_date_findings(
    sheet: Any, cells: list[Any], header_row: int | None, max_row: int | None
) -> list[HealthFinding]:
    if header_row is None or max_row is None:
        return []
    formats: dict[int, set[str]] = defaultdict(set)
    for cell in cells:
        if isinstance(cell.value, (date, datetime)):
            formats[cell.column].add(cell.number_format or "General")
    return [
        _finding(
            "mixed_date_formats",
            sheet.title,
            _column_location(column, header_row, max_row),
            f"Date values use multiple declared number formats: {sorted(values)!r}.",
            "Declare the intended date and locale format before changing any cells.",
            details={
                "formats": {
                    number_format: sorted(
                        cell.coordinate
                        for cell in cells
                        if cell.column == column
                        and isinstance(cell.value, (date, datetime))
                        and (cell.number_format or "General") == number_format
                    )
                    for number_format in sorted(values)
                },
                "required_baseline": ["target_number_format", "locale"],
            },
        )
        for column, values in sorted(formats.items())
        if len(values) > 1
    ]


def _color_summary(color: Any) -> dict[str, Any]:
    if color is None:
        return {"type": None}
    return {
        "type": str(color.type) if color.type is not None else None,
        "rgb": str(color.rgb) if color.type == "rgb" else None,
        "indexed": int(color.indexed) if color.type == "indexed" else None,
        "theme": int(color.theme) if color.type == "theme" else None,
        "tint": float(color.tint or 0),
    }


def _style_summary(cell: Any) -> dict[str, Any]:
    font = cell.font
    fill = cell.fill
    return {
        "font": {
            "name": font.name,
            "size": float(font.sz) if font.sz is not None else None,
            "bold": bool(font.bold),
            "italic": bool(font.italic),
            "color": _color_summary(font.color),
        },
        "fill": {"type": fill.fill_type, "foreground": _color_summary(fill.fgColor)},
        "number_format": cell.number_format,
    }


def _signature(summary: Mapping[str, Any]) -> str:
    encoded = json.dumps(summary, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return sha256(encoded.encode("utf-8")).hexdigest()


def _style_findings(
    sheet: Any, cells: list[Any], header_row: int | None, max_row: int | None
) -> list[HealthFinding]:
    if header_row is None or max_row is None:
        return []
    styles: dict[int, dict[str, dict[str, Any]]] = defaultdict(dict)
    for cell in cells:
        summary = _style_summary(cell)
        digest = _signature(summary)
        group = styles[cell.column].setdefault(
            digest, {"signature_digest": digest, "signature": summary, "coordinates": []}
        )
        group["coordinates"].append(cell.coordinate)
    findings: list[HealthFinding] = []
    for column, groups in sorted(styles.items()):
        if len(groups) <= 1:
            continue
        signatures = [
            {**group, "coordinates": sorted(group["coordinates"])}
            for _, group in sorted(groups.items())
        ]
        samples = [group["coordinates"][0] for group in signatures]
        findings.append(_finding(
            "inconsistent_cell_style",
            sheet.title,
            _column_location(column, header_row, max_row),
            f"The column contains {len(groups)} font/fill/number-format signatures; samples: {samples!r}.",
            "Declare a style baseline for this range; no majority style is inferred.",
            details={
                "signatures": signatures,
                "required_baseline": "target_signature_digest",
            },
        ))
    return findings


def _range_bounds(ref: str) -> tuple[int, int, int, int]:
    if "!" in ref:
        _, boundaries = range_to_tuple(ref)
        return boundaries
    return range_boundaries(ref)


def _ref_covers(
    ref: str, min_row: int, max_row: int, min_column: int, max_column: int
) -> bool:
    try:
        left, top, right, bottom = _range_bounds(ref)
    except ValueError:
        return False
    return left <= min_column and top <= min_row and right >= max_column and bottom >= max_row


def _split_print_refs(print_area: Any) -> list[str]:
    if not print_area:
        return []
    if hasattr(print_area, "ranges"):
        return [str(item) for item in print_area.ranges]
    refs: list[str] = []
    current: list[str] = []
    quoted = False
    for char in str(print_area):
        if char == "'":
            quoted = not quoted
        if char == "," and not quoted:
            refs.append("".join(current))
            current = []
        else:
            current.append(char)
    if current:
        refs.append("".join(current))
    return [ref for ref in refs if ref]


def _layout_findings(
    sheet: Any, header_row: int | None, max_row: int | None, max_column: int | None
) -> list[HealthFinding]:
    if header_row is None or max_row is None or max_column is None:
        return []
    findings: list[HealthFinding] = []
    populated = _data_cells(sheet, header_row, max_row, max_column)
    min_column = min((cell.column for cell in populated), default=1)
    filter_ref = sheet.auto_filter.ref
    data_ref = f"{get_column_letter(min_column)}{header_row}:{get_column_letter(max_column)}{max_row}"
    if filter_ref and not _ref_covers(
        str(filter_ref), header_row, max_row, min_column, max_column
    ):
        findings.append(_finding(
            "filter_range_mismatch", sheet.title, str(filter_ref),
            f"The filter range does not cover the detected table {data_ref}.",
            "Confirm the intended table before changing the filter range.",
            details={"current_ref": str(filter_ref), "candidate_ref": data_ref},
        ))
    freeze = sheet.freeze_panes
    freeze_coordinate = str(getattr(freeze, "coordinate", freeze)) if freeze is not None else None
    freeze_row = coordinate_to_tuple(freeze_coordinate)[0] if freeze_coordinate else None
    if freeze_coordinate is not None and freeze_row != header_row + 1:
        findings.append(_finding(
            "freeze_panes_mismatch", sheet.title, freeze_coordinate,
            f"Freeze panes do not begin immediately below detected header row {header_row}.",
            "Confirm the intended navigation boundary before changing freeze panes.",
            details={"current_coordinate": freeze_coordinate, "candidate_coordinate": f"A{header_row + 1}"},
        ))
    print_refs = _split_print_refs(sheet.print_area)
    if print_refs and not any(
        _ref_covers(ref, header_row, max_row, min_column, max_column) for ref in print_refs
    ):
        rendered = ",".join(print_refs)
        findings.append(_finding(
            "print_area_mismatch", sheet.title, rendered,
            f"The print area does not cover the detected table {data_ref}.",
            "Confirm the intended printed output before changing the print area.",
            details={"current_refs": print_refs, "candidate_refs": [data_ref]},
        ))
    return findings


def _blank_zero_findings(
    sheet: Any, cells: list[Any], header_row: int | None, max_row: int | None, max_column: int | None
) -> list[HealthFinding]:
    if header_row is None or max_row is None or max_column is None:
        return []
    declared_area = max(0, sheet.max_row - header_row) * max(1, sheet.max_column)
    if declared_area > BLANK_ZERO_CELL_BUDGET:
        location = f"A{header_row + 1}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        return [_finding(
            "blank_zero_scan_budget_exceeded",
            sheet.title,
            location,
            f"Blank/zero semantic scan requires {declared_area} cells, above the {BLANK_ZERO_CELL_BUDGET} budget.",
            "Narrow the declared data region or provide an explicit blank/zero policy.",
            details={"declared_cells": declared_area, "cell_budget": BLANK_ZERO_CELL_BUDGET},
            fixability="unsupported",
        )]
    values = {(cell.row, cell.column): cell.value for cell in cells}
    findings: list[HealthFinding] = []
    for column in range(1, max_column + 1):
        column_values = [values.get((row, column)) for row in range(header_row + 1, max_row + 1)]
        zero_rows = [
            row
            for row, value in zip(range(header_row + 1, max_row + 1), column_values)
            if isinstance(value, (int, float)) and not isinstance(value, bool) and value == 0
        ]
        blank_rows = [row for row, value in zip(range(header_row + 1, max_row + 1), column_values) if value in (None, "")]
        if zero_rows and blank_rows:
            zero_locations = [f"{get_column_letter(column)}{row}" for row in zero_rows]
            blank_locations = [f"{get_column_letter(column)}{row}" for row in blank_rows]
            findings.append(_finding(
                "blank_zero_semantics",
                sheet.title,
                _column_location(column, header_row, max_row),
                f"The range mixes {len(zero_rows)} numeric zero(s) with {len(blank_rows)} blank value(s).",
                "Declare whether blank and zero have distinct business meanings before conversion.",
                details={
                    "zero_locations": zero_locations,
                    "blank_locations": blank_locations,
                    "required_baseline": ["blank_action", "zero_action"],
                },
            ))
    return findings


def extended_sheet_findings(
    sheet: Any,
    header_row: int | None,
    max_row: int | None,
    max_column: int | None,
) -> list[HealthFinding]:
    """Return deterministic findings without inferring formatting or value semantics."""

    cells = _data_cells(sheet, header_row, max_row, max_column)
    findings = _mixed_date_findings(sheet, cells, header_row, max_row)
    findings.extend(_style_findings(sheet, cells, header_row, max_row))
    findings.extend(_layout_findings(sheet, header_row, max_row, max_column))
    findings.extend(_blank_zero_findings(sheet, cells, header_row, max_row, max_column))
    return findings
