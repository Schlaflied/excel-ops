"""Orchestrate multi-format artifacts from an already verified XLSX delivery."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook

from .excel_pdf_export import ExcelPdfExportError, export_pdf_with_excel
from .idempotency import file_content_digest
from .pdf_export import PdfExportError, export_pdf
from .workbook_export import WorkbookExportError, export_csv


class DeliveryExportError(ValueError):
    """Raised when a requested delivery artifact cannot be produced honestly."""


@dataclass(frozen=True)
class ExportArtifact:
    """Evidence for one concrete exported file."""

    requested_format: str
    actual_format: str
    source_workbook: str
    output: str
    sheets: tuple[str, ...]
    warnings: tuple[str, ...]
    digest: str
    verification: str = "passed"
    summary: str = ""

    def to_dict(self) -> dict[str, object]:
        return {
            "requested_format": self.requested_format,
            "actual_format": self.actual_format,
            "source_workbook": self.source_workbook,
            "output": self.output,
            "sheets": list(self.sheets),
            "warnings": list(self.warnings),
            "digest": self.digest,
            "verification": self.verification,
            "summary": self.summary,
        }


def export_delivery_artifacts(
    source: str | Path,
    *,
    formats: Sequence[str],
    selection: Mapping[str, Any],
) -> tuple[ExportArtifact, ...]:
    """Generate requested siblings from one verified workbook and return evidence."""

    source_path = Path(source).resolve()
    if not source_path.is_file() or source_path.suffix.lower() != ".xlsx":
        raise DeliveryExportError("multi-format export requires a verified XLSX source")
    sheets = _sheet_names(source_path)
    artifacts: list[ExportArtifact] = []
    for format_name in formats:
        if format_name == "xlsx":
            artifacts.append(_artifact("xlsx", source_path, source_path, sheets, ()))
        elif format_name == "csv":
            artifacts.extend(_export_csv_artifacts(source_path, selection, sheets))
        elif format_name == "pdf":
            artifacts.append(_export_pdf_artifact(source_path, selection, sheets))
        else:
            raise DeliveryExportError(f"unsupported delivery format: {format_name}")
    return tuple(artifacts)


def _export_csv_artifacts(
    source: Path, selection: Mapping[str, Any], sheets: tuple[str, ...]
) -> tuple[ExportArtifact, ...]:
    """Export the declared CSV scope without selecting an implicit active sheet."""

    csv_options = selection.get("csv") or {}
    mode = csv_options.get("mode")
    if mode == "single-sheet":
        chosen = str(csv_options.get("sheet") or "")
        output = source.with_suffix(".csv")
        try:
            paths = export_csv(source, output, sheet=chosen)
        except WorkbookExportError as error:
            raise DeliveryExportError(str(error)) from error
        scopes = ((chosen,),)
    elif mode == "one-file-per-sheet":
        output = source.with_name(f"{source.stem}-csv")
        try:
            paths = export_csv(source, output, sheets=sheets)
        except WorkbookExportError as error:
            raise DeliveryExportError(str(error)) from error
        scopes = tuple((sheet,) for sheet in sheets)
    else:
        raise DeliveryExportError("CSV export requires an explicit mode")
    return tuple(
        _artifact("csv", source, path, scope, ("csv_drops_workbook_features",))
        for path, scope in zip(paths, scopes)
    )


def _export_pdf_artifact(
    source: Path, selection: Mapping[str, Any], sheets: tuple[str, ...]
) -> ExportArtifact:
    """Prefer native Excel on Windows and fall back to LibreOffice elsewhere."""

    pdf_options = selection.get("pdf") or {}
    requested = pdf_options.get("sheets")
    chosen = sheets if requested is None else tuple(str(item) for item in requested)
    output = source.with_suffix(".pdf")
    try:
        if _is_windows():
            if requested is None:
                chosen = _visible_sheet_names(source)
            export_pdf_with_excel(source, output, sheets=chosen)
            summary = "rendered by local Microsoft Excel"
        else:
            export_pdf(source, output, sheets=chosen)
            summary = "rendered by local LibreOffice"
    except (ExcelPdfExportError, PdfExportError) as error:
        raise DeliveryExportError(str(error)) from error
    return _artifact(
        "pdf",
        source,
        output,
        chosen,
        ("pdf_is_not_editable", "pdf_layout_not_verified"),
        summary,
        verification="unverified",
    )


def _artifact(
    format_name: str,
    source: Path,
    output: Path,
    sheets: Sequence[str],
    warnings: Sequence[str],
    summary: str = "verified output",
    verification: str = "passed",
) -> ExportArtifact:
    """Build evidence only after checking suffix and persisted content."""

    expected = f".{format_name}"
    if output.suffix.lower() != expected or not output.is_file():
        raise DeliveryExportError(f"{format_name} export did not create a matching file")
    return ExportArtifact(
        format_name,
        format_name,
        str(source),
        str(output.resolve()),
        tuple(sheets),
        tuple(warnings),
        file_content_digest(output),
        verification=verification,
        summary=summary,
    )


def _sheet_names(source: Path) -> tuple[str, ...]:
    """Read the workbook sheet scope once for evidence and explicit adapters."""

    workbook = load_workbook(source, read_only=True)
    try:
        return tuple(workbook.sheetnames)
    finally:
        workbook.close()


def _visible_sheet_names(source: Path) -> tuple[str, ...]:
    """Return the Windows PDF `all` scope accepted by native Excel."""

    workbook = load_workbook(source, read_only=True)
    try:
        return tuple(
            name
            for name in workbook.sheetnames
            if getattr(workbook[name], "sheet_state", "visible") == "visible"
        )
    finally:
        workbook.close()


def _is_windows() -> bool:
    """Keep platform selection patchable without mutating Python's global os module."""

    import os

    return os.name == "nt"
