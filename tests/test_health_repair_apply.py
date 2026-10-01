from __future__ import annotations

from hashlib import sha256
from dataclasses import replace
import os
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
import pytest

from excel_ops.health_models import HealthRepairAction, HealthRepairApplyError
from excel_ops.health_repair_apply import apply_workbook_health_repairs
from excel_ops.health_repair_plan import plan_workbook_health_repairs


def _digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _source(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["Name", "Amount", "Formula"])
    sheet.append(["  Alpha  ", 2, "=B2*2"])
    validation = DataValidation(type="whole", operator="between", formula1="1", formula2="10")
    sheet.add_data_validation(validation)
    validation.add(sheet["B2"])
    hidden = workbook.create_sheet("辅助")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "preserve"
    workbook.save(path)
    workbook.close()


def _make_valid_macro_package(xlsx: Path, xlsm: Path, marker: bytes) -> None:
    with ZipFile(xlsx) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    content_types = members["[Content_Types].xml"].decode("utf-8")
    content_types = content_types.replace(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    ).replace(
        "</Types>",
        '<Override PartName="/xl/vbaProject.bin" '
        'ContentType="application/vnd.ms-office.vbaProject"/></Types>',
    )
    relationships = members["xl/_rels/workbook.xml.rels"].decode("utf-8").replace(
        "</Relationships>",
        '<Relationship Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
        'Target="vbaProject.bin" Id="rIdVba"/></Relationships>',
    )
    members["[Content_Types].xml"] = content_types.encode("utf-8")
    members["xl/_rels/workbook.xml.rels"] = relationships.encode("utf-8")
    members["xl/vbaProject.bin"] = marker
    with ZipFile(xlsm, "w", ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)


def _selected_whitespace_plan(source: Path):
    initial = plan_workbook_health_repairs(source)
    action = next(item for item in initial.actions if item.finding_code == "extra_whitespace")
    plan = plan_workbook_health_repairs(
        source, dispositions={action.action_id: "selected"}
    )
    return plan, action.action_id


def test_applies_only_allowlisted_action_to_new_copy_and_writes_ledger(tmp_path: Path) -> None:
    source = tmp_path / "中文源.xlsx"
    output = tmp_path / "中文修复.xlsx"
    _source(source)
    source_before = source.read_bytes()
    plan, action_id = _selected_whitespace_plan(source)

    result = apply_workbook_health_repairs(
        source, output, plan, selected_action_ids=[action_id]
    )

    assert source.read_bytes() == source_before
    assert output.exists()
    assert result.output_sha256 == _digest(output)
    assert [change.action_id for change in result.changes] == [action_id]
    assert result.changes[0].before == "  Alpha  "
    assert result.changes[0].after == "Alpha"
    reopened = load_workbook(output, data_only=False)
    try:
        assert reopened["数据"]["A2"].value == "Alpha"
        assert reopened["数据"]["C2"].value == "=B2*2"
        assert len(reopened["数据"].data_validations.dataValidation) == 1
        assert reopened["辅助"].sheet_state == "hidden"
        assert reopened["辅助"]["A1"].value == "preserve"
    finally:
        reopened.close()


def test_rejects_unselected_unknown_empty_and_existing_output(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)

    with pytest.raises(HealthRepairApplyError, match="must not be empty"):
        apply_workbook_health_repairs(source, tmp_path / "empty.xlsx", plan, selected_action_ids=[])
    with pytest.raises(HealthRepairApplyError, match="unknown action"):
        apply_workbook_health_repairs(
            source, tmp_path / "unknown.xlsx", plan, selected_action_ids=["health-unknown"]
        )
    initial = plan_workbook_health_repairs(source)
    with pytest.raises(HealthRepairApplyError, match="not selected and ready"):
        apply_workbook_health_repairs(
            source, tmp_path / "proposed.xlsx", initial, selected_action_ids=[action_id]
        )
    existing = tmp_path / "existing.xlsx"
    existing.write_bytes(b"do-not-replace")
    with pytest.raises(HealthRepairApplyError, match="already exists"):
        apply_workbook_health_repairs(source, existing, plan, selected_action_ids=[action_id])
    assert existing.read_bytes() == b"do-not-replace"


def test_stale_source_is_rejected_without_publishing_output(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)
    workbook = load_workbook(source)
    workbook["数据"]["B2"] = 3
    workbook.save(source)
    workbook.close()

    with pytest.raises(HealthRepairApplyError, match="changed after repair planning"):
        apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])

    assert not output.exists()


def test_save_failure_never_publishes_partial_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)

    def fail_save(self, filename):
        filename.write(b"partial")
        filename.flush()
        raise OSError("synthetic save failure")

    monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", fail_save)
    with pytest.raises(OSError, match="synthetic save failure"):
        apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])

    assert not output.exists()
    assert not list(tmp_path.glob(".output-*"))


def test_xlsm_vba_member_is_preserved(tmp_path: Path) -> None:
    xlsx = tmp_path / "base.xlsx"
    source = tmp_path / "宏源.xlsm"
    output = tmp_path / "宏修复.xlsm"
    _source(xlsx)
    marker = b"synthetic-vba-marker-not-executable"
    _make_valid_macro_package(xlsx, source, marker)
    plan, action_id = _selected_whitespace_plan(source)

    apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])

    with ZipFile(output) as archive:
        assert archive.testzip() is None
        assert archive.read("xl/vbaProject.bin") == marker
        assert b"macroEnabled.main+xml" in archive.read("[Content_Types].xml")
        assert b"relationships/vbaProject" in archive.read("xl/_rels/workbook.xml.rels")
    reopened = load_workbook(output, keep_vba=True, data_only=False)
    try:
        assert reopened["数据"]["C2"].value == "=B2*2"
        assert reopened.vba_archive is not None
    finally:
        reopened.close()
        reopened.vba_archive.close()


@pytest.mark.parametrize("suffix", [".xlsx", ".xlsm"])
def test_unrecognized_package_members_must_not_be_silently_dropped(
    tmp_path: Path, suffix: str
) -> None:
    xlsx = tmp_path / "base.xlsx"
    source = tmp_path / f"source{suffix}"
    output = tmp_path / f"output{suffix}"
    _source(xlsx)
    if suffix == ".xlsm":
        _make_valid_macro_package(xlsx, source, b"synthetic-vba-marker-not-executable")
    else:
        source.write_bytes(xlsx.read_bytes())
    with ZipFile(source, "a") as archive:
        archive.writestr("customXml/item1.xml", b"<custom-preserve />")
    plan, action_id = _selected_whitespace_plan(source)

    with pytest.raises(HealthRepairApplyError, match="did not preserve package member"):
        apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])

    assert not output.exists()


def test_original_workbook_closes_when_repaired_reload_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)
    from excel_ops import health_repair_apply as module

    real_load = module.load_workbook
    real_close = module._close_workbook
    loaded = []
    closed = []

    def tracked_load(*args, **kwargs):
        if len(loaded) == 2:
            raise OSError("synthetic repaired reload failure")
        workbook = real_load(*args, **kwargs)
        loaded.append(workbook)
        return workbook

    def tracked_close(workbook):
        closed.append(workbook)
        real_close(workbook)

    monkeypatch.setattr(module, "load_workbook", tracked_load)
    monkeypatch.setattr(module, "_close_workbook", tracked_close)

    with pytest.raises(OSError, match="synthetic repaired reload failure"):
        apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])

    assert loaded[1] in closed
    assert not output.exists()


def test_forged_or_mutated_plan_is_rejected(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)
    original = next(action for action in plan.actions if action.action_id == action_id)
    forged = replace(
        original,
        proposed_state={"status": "ready", "target": {"candidate_value": "forged"}},
    )
    forged_plan = replace(
        plan,
        actions=tuple(forged if item.action_id == action_id else item for item in plan.actions),
    )

    with pytest.raises(HealthRepairApplyError, match="forged, mutated, or is stale"):
        apply_workbook_health_repairs(
            source, output, forged_plan, selected_action_ids=[action_id]
        )

    assert not output.exists()


def test_structural_repairs_fail_closed_even_when_selected(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet"
    data = workbook.create_sheet("Data")
    data.append(["Name", "Value"])
    data.append(["Alpha", 1])
    workbook.save(source)
    workbook.close()
    initial = plan_workbook_health_repairs(source)
    action = next(item for item in initial.actions if item.finding_code == "default_empty_sheet")
    selected = plan_workbook_health_repairs(
        source, dispositions={action.action_id: "selected"}
    )

    with pytest.raises(HealthRepairApplyError, match="cannot be proven to preserve"):
        apply_workbook_health_repairs(
            source, output, selected, selected_action_ids=[action.action_id]
        )

    assert not output.exists()


def test_persisted_unauthorized_change_is_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)
    from excel_ops import health_repair_apply as module

    original_apply = module._apply_action

    def tampering_apply(workbook, action):
        result = original_apply(workbook, action)
        workbook["辅助"]["A1"] = "tampered"
        return result

    monkeypatch.setattr(module, "_apply_action", tampering_apply)
    with pytest.raises(HealthRepairApplyError, match="unauthorized"):
        apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])

    assert not output.exists()


def test_publish_race_does_not_delete_foreign_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)
    from excel_ops import health_repair_apply as module

    def race(guard, staged, destination):
        destination.write_bytes(b"foreign")
        raise HealthRepairApplyError("synthetic publish race")

    monkeypatch.setattr(module, "_publish_without_overwrite", race)
    with pytest.raises(HealthRepairApplyError, match="synthetic publish race"):
        apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])

    assert output.read_bytes() == b"foreign"


def test_duplicate_header_formula_cannot_bypass_apply_protection(tmp_path: Path) -> None:
    source = tmp_path / "formula-header.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["=A2", "=A2"])
    sheet.append([1, 2])
    workbook.save(source)
    workbook.close()
    initial = plan_workbook_health_repairs(source)
    action = next(item for item in initial.actions if item.finding_code == "duplicate_header")
    selected = plan_workbook_health_repairs(
        source,
        dispositions={action.action_id: "selected"},
        baselines={action.action_id: {"target_header": "Renamed"}},
    )

    with pytest.raises(HealthRepairApplyError, match="protected formula"):
        apply_workbook_health_repairs(
            source, output, selected, selected_action_ids=[action.action_id]
        )

    assert not output.exists()


def test_apply_layer_rejects_merged_follower_independently() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.merge_cells("A1:B1")
    action = HealthRepairAction(
        action_id="health-test",
        finding_id="finding-test",
        finding_code="duplicate_header",
        classification="review",
        decision="selected",
        operation="rename_duplicate_header",
        sheet="Data",
        location="B1",
        impact_scope="Data!B1",
        current_state={},
        proposed_state={"status": "ready", "target": {"target_header": "New"}},
    )
    from excel_ops import health_repair_apply as module

    with pytest.raises(HealthRepairApplyError, match="merged-range follower"):
        module._apply_cell_action(sheet, action, {"target_header": "New"})


def test_persisted_target_mismatch_is_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)
    from excel_ops import health_repair_apply as module

    monkeypatch.setattr(module, "_apply_action", lambda workbook, action: None)
    with pytest.raises(HealthRepairApplyError, match="persisted target mismatch"):
        apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])

    assert not output.exists()


def test_print_area_target_and_ledger_use_persisted_ranges(tmp_path: Path) -> None:
    source = tmp_path / "print.xlsx"
    output = tmp_path / "print-fixed.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Amount", "Note"])
    sheet.append(["Alpha", 1, "x"])
    sheet.print_area = "A1:B2"
    workbook.save(source)
    workbook.close()
    initial = plan_workbook_health_repairs(source)
    action = next(item for item in initial.actions if item.finding_code == "print_area_mismatch")
    selected = plan_workbook_health_repairs(
        source, dispositions={action.action_id: "selected"}
    )

    result = apply_workbook_health_repairs(
        source, output, selected, selected_action_ids=[action.action_id]
    )

    change = result.changes[0]
    assert change.before and change.after
    assert change.before != change.after
    from excel_ops import health_repair_apply as module
    assert {module._range_bounds(item) for item in change.after} == {(1, 1, 3, 2)}


def test_source_path_change_after_snapshot_blocks_publication(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)
    from excel_ops import health_repair_apply as module

    original_canonical = module._canonical_plan

    def mutate_after_snapshot(snapshot, supplied):
        canonical = original_canonical(snapshot, supplied)
        source.write_bytes(source.read_bytes() + b"changed")
        return canonical

    monkeypatch.setattr(module, "_canonical_plan", mutate_after_snapshot)
    with pytest.raises(HealthRepairApplyError, match="changed before repaired copy publication"):
        apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])

    assert not output.exists()


def test_published_identity_must_match_staged_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)
    from excel_ops import health_repair_apply as module

    def copy_instead_of_link(guard, staged, destination):
        destination.write_bytes(staged.read_bytes())

    monkeypatch.setattr(module, "_publish_without_overwrite", copy_instead_of_link)
    with pytest.raises(HealthRepairApplyError, match="identity does not match"):
        apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])

    assert output.exists(), "identity-mismatched external output must not be deleted"


def test_blank_zero_preserve_is_verified_from_source_not_observed_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "blank-zero.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Amount"])
    sheet.append(["Alpha", 0])
    sheet.append(["Beta", None])
    workbook.save(source)
    workbook.close()
    initial = plan_workbook_health_repairs(source)
    action = next(item for item in initial.actions if item.finding_code == "blank_zero_semantics")
    selected = plan_workbook_health_repairs(
        source,
        dispositions={action.action_id: "selected"},
        baselines={
            action.action_id: {"blank_action": "preserve", "zero_action": "preserve"}
        },
    )
    from excel_ops import health_repair_apply as module

    original_apply = module._apply_action

    def tamper_authorized_cell(workbook, selected_action):
        result = original_apply(workbook, selected_action)
        workbook["Data"]["B3"] = 99
        return result

    monkeypatch.setattr(module, "_apply_action", tamper_authorized_cell)
    with pytest.raises(HealthRepairApplyError, match="persisted target mismatch"):
        apply_workbook_health_repairs(
            source, output, selected, selected_action_ids=[action.action_id]
        )

    assert not output.exists()


def test_staged_leaf_swap_cannot_redirect_workbook_save(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    victim = tmp_path / "victim.xlsx"
    victim.write_bytes(b"victim-must-not-change")
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)
    from excel_ops import health_repair_apply as module
    from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

    original_save = OpenpyxlWorkbook.save
    attack = {"swapped": False, "received_stream": False}

    def swap_leaf_then_save(self, destination):
        attack["received_stream"] = hasattr(destination, "fileno")
        staged = next(tmp_path.glob(".output-*.xlsx"))
        try:
            staged.unlink()
            staged.symlink_to(victim)
            attack["swapped"] = True
        except OSError:
            # Windows holds the staged leaf open, so replacement itself fails closed.
            pass
        return original_save(self, destination)

    monkeypatch.setattr(OpenpyxlWorkbook, "save", swap_leaf_then_save)
    if os.name == "nt":
        apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])
        assert not attack["swapped"]
    else:
        with pytest.raises(HealthRepairApplyError, match="staged output changed"):
            apply_workbook_health_repairs(source, output, plan, selected_action_ids=[action_id])
        assert attack["swapped"]
    assert attack["received_stream"]
    assert victim.read_bytes() == b"victim-must-not-change"


def test_symlinked_output_parent_is_rejected_when_supported(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _source(source)
    plan, action_id = _selected_whitespace_plan(source)
    outside = tmp_path / "outside"
    outside.mkdir()
    link = tmp_path / "link"
    try:
        link.symlink_to(outside, target_is_directory=True)
    except OSError:
        pytest.skip("directory symlinks are not available")

    with pytest.raises(HealthRepairApplyError, match="symlink"):
        apply_workbook_health_repairs(
            source, link / "escaped.xlsx", plan, selected_action_ids=[action_id]
        )

    assert not (outside / "escaped.xlsx").exists()
