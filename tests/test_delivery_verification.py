from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel_ops.delivery_verification import (
    DeliveryContract,
    DeliveryVerificationError,
    PeriodExpectation,
    StageCounts,
    verify_and_deliver,
)
from excel_ops.template_writer import TemplateMapping, write_template


def _template(path: Path, *, headers=("record_id", "employee", "hours")) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payroll"
    sheet["A1"] = "Period"
    sheet["B1"] = "2026-09-01 to 2026-09-07"
    for column, header in enumerate(headers, start=1):
        sheet.cell(3, column).value = header
    workbook.save(path)


def _mapping() -> TemplateMapping:
    return TemplateMapping(
        sheet="Payroll",
        header_row=3,
        data_start_row=4,
        field_columns={"record_id": "A", "employee": "B", "hours": "C"},
    )


def _contract(*, written=2, accepted=2, review=1, expected_ids=("src-1", "src-2")) -> DeliveryContract:
    return DeliveryContract(
        counts=StageCounts(input=accepted + review, accepted=accepted, review=review, written=written),
        required_fields=("record_id", "employee", "hours"),
        record_id_field="record_id",
        expected_record_ids=expected_ids,
        period_expectations=(
            PeriodExpectation("Payroll", "B1", "2026-09-01 to 2026-09-07"),
        ),
    )


def test_reopens_staged_and_delivery_files_before_publishing(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    staged = tmp_path / "staging" / "payroll.xlsx"
    delivery = tmp_path / "delivery" / "payroll.xlsx"
    _template(template)
    records = [
        {"record_id": "src-1", "employee": "Ada", "hours": 8},
        {"record_id": "src-2", "employee": "Lin", "hours": 4},
    ]
    write_result = write_template(template, records, _mapping(), output_path=staged)

    result = verify_and_deliver(write_result, delivery, _contract())

    assert result.passed is True
    assert result.delivery_path == delivery.resolve()
    assert delivery.is_file()
    assert json.loads(result.report_path.read_text(encoding="utf-8"))["passed"] is True


@pytest.mark.parametrize(
    ("mutate", "code", "field"),
    [
        (lambda ws: setattr(ws["C3"], "value", "employee"), "column_shift", "hours"),
        (lambda ws: setattr(ws["B4"], "value", None), "missing_required_field", "employee"),
        (lambda ws: setattr(ws["B1"], "value", "2025-01-01"), "period_mismatch", "report_period"),
    ],
)
def test_structured_failures_block_delivery(tmp_path: Path, mutate, code: str, field: str):
    template = tmp_path / "template.xlsx"
    staged = tmp_path / "staging.xlsx"
    delivery = tmp_path / "delivery" / "final.xlsx"
    _template(template)
    write_result = write_template(
        template,
        [{"record_id": "src-1", "employee": "Ada", "hours": 8}],
        _mapping(),
        output_path=staged,
    )
    workbook = load_workbook(staged)
    mutate(workbook["Payroll"])
    workbook.save(staged)
    workbook.close()

    with pytest.raises(DeliveryVerificationError) as caught:
        verify_and_deliver(write_result, delivery, _contract(written=1, accepted=1, expected_ids=("src-1",)))

    assert not delivery.exists()
    finding = next(item for item in caught.value.result.findings if item.code == code)
    assert finding.sheet == "Payroll"
    assert finding.field == field
    assert finding.suggestion
    report = json.loads(caught.value.result.report_path.read_text(encoding="utf-8"))
    assert report["passed"] is False


def test_stage_counts_are_reconciled_before_delivery(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    staged = tmp_path / "staging.xlsx"
    delivery = tmp_path / "delivery.xlsx"
    _template(template)
    write_result = write_template(
        template,
        [{"record_id": "src-1", "employee": "Ada", "hours": 8}],
        _mapping(),
        output_path=staged,
    )
    contract = DeliveryContract(
        counts=StageCounts(input=3, accepted=2, review=0, written=1),
        required_fields=("record_id", "employee", "hours"),
    )

    with pytest.raises(DeliveryVerificationError) as caught:
        verify_and_deliver(write_result, delivery, contract)

    assert {finding.code for finding in caught.value.result.findings} >= {
        "stage_count_mismatch",
        "written_count_mismatch",
    }
    assert not delivery.exists()


def test_duplicate_source_id_blocks_repeated_write(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    staged = tmp_path / "staging.xlsx"
    _template(template)
    records = [
        {"record_id": "src-1", "employee": "Ada", "hours": 8},
        {"record_id": "src-1", "employee": "Ada", "hours": 8},
    ]
    write_result = write_template(template, records, _mapping(), output_path=staged)

    with pytest.raises(DeliveryVerificationError) as caught:
        verify_and_deliver(write_result, tmp_path / "delivery.xlsx", _contract(expected_ids=("src-1", "src-1")))

    duplicate = next(item for item in caught.value.result.findings if item.code == "duplicate_record")
    assert duplicate.row == 5
    assert duplicate.cell == "A5"


def test_same_records_can_be_written_twice_without_appending(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    records = [
        {"record_id": "src-1", "employee": "Ada", "hours": 8},
        {"record_id": "src-2", "employee": "Lin", "hours": 4},
    ]
    _template(template)
    write_template(template, records, _mapping(), output_path=first)
    second_result = write_template(first, records, _mapping(), output_path=second)

    result = verify_and_deliver(second_result, tmp_path / "delivery" / "final.xlsx", _contract())

    assert result.passed
    workbook = load_workbook(result.delivery_path)
    assert workbook["Payroll"].max_row == 5
    workbook.close()


def test_zero_record_target_requires_marker_and_current_period(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    staged = tmp_path / "zero.xlsx"
    _template(template, headers=("status", "employee", "hours"))
    mapping = TemplateMapping(
        sheet="Payroll",
        header_row=3,
        data_start_row=4,
        field_columns={"status": "A", "employee": "B", "hours": "C"},
    )
    write_result = write_template(template, [{"status": "No records"}], mapping, output_path=staged)
    contract = DeliveryContract(
        counts=StageCounts(input=0, accepted=0, review=0, written=0),
        required_fields=(),
        empty_marker_field="status",
        empty_marker_value="No records",
        period_expectations=(PeriodExpectation("Payroll", "B1", "2026-09-01 to 2026-09-07"),),
    )

    result = verify_and_deliver(write_result, tmp_path / "delivery" / "zero.xlsx", contract)

    assert result.passed


def test_extension_verifier_uses_the_shared_finding_contract(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    staged = tmp_path / "staging.xlsx"
    _template(template)
    write_result = write_template(
        template,
        [{"record_id": "src-1", "employee": "Ada", "hours": 8}],
        _mapping(),
        output_path=staged,
    )

    def reject_formula_placeholder(workbook, _write_result, _contract):
        from excel_ops.delivery_verification import VerificationFinding

        assert workbook["Payroll"]["C4"].value == 8
        return [VerificationFinding("custom_check", "Custom verifier failed.", "Fix the custom rule.", "Payroll", 4, "hours", "C4")]

    contract = DeliveryContract(
        counts=StageCounts(input=1, accepted=1, review=0, written=1),
        required_fields=("record_id", "employee", "hours"),
        verifiers=(reject_formula_placeholder,),
    )

    with pytest.raises(DeliveryVerificationError) as caught:
        verify_and_deliver(write_result, tmp_path / "delivery.xlsx", contract)

    assert any(item.code == "custom_check" for item in caught.value.result.findings)
