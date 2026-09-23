import csv
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from excel_ops.workbook_export import WorkbookExportError, export_csv, export_xlsx


def _workbook(path):
    workbook = Workbook()
    first = workbook.active
    first.title = "Payroll"
    first.append(["员工", "金额"])
    first.append(["林", 12.5])
    second = workbook.create_sheet("Summary")
    second.append(["总数"])
    second.append([1])
    workbook.save(path)


def test_xlsx_export_copies_a_real_workbook(tmp_path):
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "out.xlsx"
    _workbook(source)
    assert export_xlsx(source, destination) == destination
    assert load_workbook(destination, read_only=True).sheetnames == ["Payroll", "Summary"]


def test_csv_single_sheet_is_utf8_and_preserves_values(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "payroll.csv"
    _workbook(source)
    outputs = export_csv(source, output, sheet="Payroll")
    assert outputs == (output,)
    with output.open(encoding="utf-8-sig", newline="") as stream:
        assert list(csv.reader(stream)) == [["员工", "金额"], ["林", "12.5"]]


def test_csv_can_export_each_sheet_to_a_directory(tmp_path):
    source = tmp_path / "source.xlsx"
    output_dir = tmp_path / "csv"
    _workbook(source)
    outputs = export_csv(source, output_dir, sheets=["Payroll", "Summary"])
    assert [path.name for path in outputs] == ["Payroll.csv", "Summary.csv"]


def test_csv_rejects_ambiguous_or_missing_sheets(tmp_path):
    source = tmp_path / "source.xlsx"
    _workbook(source)
    with pytest.raises(WorkbookExportError, match="choose sheet"):
        export_csv(source, tmp_path / "out.csv", sheet="Payroll", sheets=["Summary"])
    with pytest.raises(WorkbookExportError, match="unknown worksheet"):
        export_csv(source, tmp_path / "out.csv", sheet="Missing")


def test_csv_preserves_explicit_empty_selection_as_an_error(tmp_path):
    source = tmp_path / "source.xlsx"
    _workbook(source)
    with pytest.raises(WorkbookExportError, match="at least one sheet"):
        export_csv(source, tmp_path / "csv", sheets=[])


def test_csv_sanitizes_windows_sheet_names(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "CON"
    workbook.create_sheet("Summary")
    workbook.save(source)
    outputs = export_csv(source, tmp_path / "csv", sheets=["CON", "Summary"])
    assert [path.name for path in outputs] == ["_CON.csv", "Summary.csv"]


def test_csv_sanitizes_reserved_names_with_extensions_and_superscripts(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "CON.report"
    workbook.create_sheet("COM²")
    workbook.create_sheet("LPT³")
    workbook.save(source)
    outputs = export_csv(source, tmp_path / "csv")
    assert [path.name for path in outputs] == ["_CON.report.csv", "_COM².csv", "_LPT³.csv"]


def test_csv_rejects_formula_without_cached_result(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "=1+1"
    workbook.save(source)
    with pytest.raises(WorkbookExportError, match="no cached result"):
        export_csv(source, tmp_path / "out.csv", sheet="Sheet")


def test_formula_cache_parser_infers_missing_cell_coordinate(tmp_path):
    source = tmp_path / "source.xlsx"
    rewritten = tmp_path / "rewritten.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "=1+1"
    workbook.save(source)
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(rewritten, "w") as target:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<c r="A1"', b"<c", 1)
            target.writestr(info, data)
    with pytest.raises(WorkbookExportError, match="no cached result"):
        export_csv(rewritten, tmp_path / "out.csv", sheet="Sheet")


def test_csv_neutralizes_formula_like_text(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    cell = workbook.active["A1"]
    cell.value = "=not-a-formula"
    cell.data_type = "s"
    workbook.save(source)
    outputs = export_csv(source, tmp_path / "out.csv", sheet="Sheet")
    assert outputs[0].read_text(encoding="utf-8-sig").strip() == "'=not-a-formula"


@pytest.mark.parametrize("value", ["\t=SUM(A1:A2)", "\r+1", "\n@cmd"])
def test_csv_neutralizes_control_prefixes(tmp_path, value):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = value
    workbook.save(source)
    export_csv(source, tmp_path / "out.csv", sheet="Sheet")
    with (tmp_path / "out.csv").open(encoding="utf-8-sig", newline="") as stream:
        assert next(csv.reader(stream))[0].startswith("'")


def test_xlsm_is_not_silently_renamed_to_xlsx(tmp_path):
    source = tmp_path / "source.xlsm"
    source.write_bytes(b"not a workbook")
    with pytest.raises(WorkbookExportError, match="XLSX source"):
        export_xlsx(source, tmp_path / "out.xlsx")
