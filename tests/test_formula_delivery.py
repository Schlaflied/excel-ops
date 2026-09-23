from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel_ops.formula_delivery import (
    FormulaDeliveryError,
    FormulaDeliveryRule,
    apply_formula_delivery,
    formula_rule_from_mapping,
)
from excel_ops.formula_planning import FormulaPlan
from excel_ops.formula_recalculation import FormulaValueExpectation


def _formula_plan() -> FormulaPlan:
    return FormulaPlan(
        business_rule="Add three to the source value.",
        operation="addition",
        output_mode="formula",
        target_excel_version="365",
        value="=C2+3",
        function="addition",
        compatibility_strategy="basic arithmetic works in all target versions",
        requires_independent_recalculation=True,
    )


def test_formula_delivery_applies_and_independently_recalculates(tmp_path: Path):
    workbook_path = tmp_path / "delivery.xlsx"
    workbook = Workbook()
    workbook.active.title = "Report"
    workbook["Report"]["C2"] = 2
    workbook.save(workbook_path)
    workbook.close()

    def recalculate(source: Path, destination: Path) -> str:
        calculated = load_workbook(source)
        calculated["Report"]["D2"] = 5
        calculated.save(destination)
        calculated.close()
        return "test-engine"

    evidence = apply_formula_delivery(
        workbook_path,
        (
            FormulaDeliveryRule(
                _formula_plan(),
                "Report",
                "D2",
                (FormulaValueExpectation("Report", "D2", 5),),
            ),
        ),
        recalculator=recalculate,
    )

    assert evidence["status"] == "verified"
    assert evidence["engine"] == "test-engine"
    assert evidence["rules"][0]["written_cells"] == 1
    assert "expected" not in str(evidence)


def test_formula_rule_requires_independent_expectations():
    with pytest.raises(FormulaDeliveryError, match="requires independently derived"):
        formula_rule_from_mapping(
            {
                "plan": _formula_plan().to_dict(),
                "sheet": "Report",
                "target_range": "D2",
            }
        )


def test_expected_values_are_part_of_the_idempotency_fingerprint():
    first = FormulaDeliveryRule(
        _formula_plan(),
        "Report",
        "D2",
        (FormulaValueExpectation("Report", "D2", 5),),
    )
    second = FormulaDeliveryRule(
        _formula_plan(),
        "Report",
        "D2",
        (FormulaValueExpectation("Report", "D2", 6),),
    )

    assert first.fingerprint_payload() != second.fingerprint_payload()
    assert first.manifest_payload() == second.manifest_payload()


def test_static_formula_rule_rejects_unverified_expectations():
    plan = _formula_plan().to_dict()
    plan.update(
        output_mode="static",
        value="approved",
        requires_independent_recalculation=False,
    )

    with pytest.raises(FormulaDeliveryError, match="static mode expectations"):
        formula_rule_from_mapping(
            {
                "plan": plan,
                "sheet": "Report",
                "target_range": "D2",
                "expectations": [
                    {"sheet": "Report", "cell": "D2", "expected": "approved"}
                ],
            }
        )


@pytest.mark.parametrize(
    ("field", "value", "message"),
    [
        (
            "requires_independent_recalculation",
            "false",
            "requires_independent_recalculation must be a boolean",
        ),
        ("overwrite_formulas", "false", "overwrite_formulas must be a boolean"),
    ],
)
def test_formula_rule_requires_strict_booleans(field, value, message):
    payload = {
        "plan": _formula_plan().to_dict(),
        "sheet": "Report",
        "target_range": "D2",
        "expectations": [{"sheet": "Report", "cell": "D2", "expected": 5}],
    }
    if field == "requires_independent_recalculation":
        payload["plan"][field] = value
    else:
        payload[field] = value

    with pytest.raises(FormulaDeliveryError, match=message):
        formula_rule_from_mapping(payload)
