from openpyxl import load_workbook

from excel_ops.pipeline import run_pipeline


def test_pipeline_routes_uncertain_rows(tmp_path):
    output = tmp_path / "delivery.xlsx"
    result = run_pipeline("examples/extracted-records.json", output)
    assert result["total"] == 2
    assert result["accepted"] == 1
    assert result["review"] == 1

    wb = load_workbook(output, data_only=True)
    assert wb["Accepted"].max_row == 2
    assert wb["Review"].max_row == 2
    assert "low_confidence" in wb["Review"]["G2"].value
    wb.close()
