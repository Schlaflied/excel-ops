from __future__ import annotations

import json
from hashlib import sha256
import os
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
import pytest

from excel_ops.cli import main
from excel_ops.health_repair_plan import plan_workbook_health_repairs
from excel_ops.health_repair_verification import repair_and_verify_workbook_health
from excel_ops.health_models import HealthFinding, HealthRepairAction, HealthRepairApplyError
import excel_ops.health_repair_verification as verification


def _source(path: Path, values=(" Alpha ",)) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Amount"])
    for index, value in enumerate(values, start=1):
        sheet.append([value, index])
    workbook.save(path)
    workbook.close()


def _whitespace_actions(source: Path):
    return [
        action
        for action in plan_workbook_health_repairs(source).actions
        if action.finding_code == "extra_whitespace"
    ]


def test_verified_repair_reopens_rescans_reconciles_and_publishes(tmp_path: Path) -> None:
    source = tmp_path / "中文源.xlsx"
    output = tmp_path / "中文修复.xlsx"
    _source(source)
    action = _whitespace_actions(source)[0]

    result = repair_and_verify_workbook_health(
        source, output, selected_action_ids=[action.action_id]
    )

    assert result.verified is True
    assert result.output == str(output.absolute())
    assert result.failures == ()
    assert [item["action_id"] for item in result.authorized_actions] == [action.action_id]
    assert [item.action_id for item in result.actual_changes] == [action.action_id]
    assert action.finding_id not in {item["finding_id"] for item in result.remaining_findings}
    reopened = load_workbook(output)
    try:
        assert reopened["Data"]["A2"].value == "Alpha"
    finally:
        reopened.close()
    assert not list(tmp_path.glob(".*-health-verify-*.xlsx"))


def test_unauthorized_finding_remains_byte_for_byte_in_reconciliation(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source, (" Alpha ", " Beta "))
    actions = _whitespace_actions(source)

    result = repair_and_verify_workbook_health(
        source, output, selected_action_ids=[actions[0].action_id]
    )

    assert result.verified is True
    remaining_ids = {item["finding_id"] for item in result.remaining_findings}
    assert actions[1].finding_id in remaining_ids


def test_health_cli_scan_dry_run_and_confirmed_write(tmp_path: Path, capsys) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    request = tmp_path / "request.json"
    _source(source)

    main(["health-scan", str(source)])
    scan_payload = json.loads(capsys.readouterr().out)
    action = next(item for item in scan_payload["actions"] if item["finding_code"] == "extra_whitespace")
    request.write_text(
        json.dumps({"selected_action_ids": [action["action_id"]], "baselines": {}}),
        encoding="utf-8",
    )
    main(["health-repair", str(source), str(output), "--request", str(request), "--dry-run"])
    planned = json.loads(capsys.readouterr().out)
    common = {"schema", "mode", "source", "source_sha256", "output", "output_sha256",
              "verified", "before", "actions", "changes", "remaining", "new", "failures"}
    assert common <= scan_payload.keys()
    assert common <= planned.keys()
    assert planned["mode"] == "dry_run"
    assert not output.exists()

    main(["health-repair", str(source), str(output), "--request", str(request), "--confirm"])
    repaired = json.loads(capsys.readouterr().out)
    assert repaired["verified"] is True
    assert repaired["mode"] == "verified_repair"
    assert {"before", "actions", "changes", "remaining", "new", "failures"} <= repaired.keys()
    assert output.exists()


def test_authorized_finding_that_remains_fails_without_publishing(tmp_path: Path) -> None:
    source = tmp_path / "numeric.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Identifier"])
    sheet.append(["Alpha", "9007199254740993"])
    workbook.save(source)
    workbook.close()
    initial = plan_workbook_health_repairs(source)
    action = next(item for item in initial.actions if item.finding_code == "numeric_text")

    result = repair_and_verify_workbook_health(
        source,
        output,
        selected_action_ids=[action.action_id],
        baselines={action.action_id: {"target_representation": "preserve_text"}},
    )

    assert result.verified is False
    assert result.output is None
    assert result.failures[0]["code"] == "authorized_finding_remains"
    assert not output.exists()


def test_reconciliation_rejects_changed_unauthorized_and_new_findings() -> None:
    authorized = HealthFinding(
        code="extra_whitespace", severity="warning", sheet="Data", location="A2",
        evidence="synthetic authorized finding", suggestion="normalize", fixability="safe",
    )
    unauthorized = HealthFinding(
        code="duplicate_header", severity="warning", sheet="Data", location="A1:B1",
        evidence="synthetic unauthorized finding", suggestion="review", fixability="review",
    )
    introduced = HealthFinding(
        code="invisible_character", severity="warning", sheet="Data", location="C2",
        evidence="synthetic new finding", suggestion="remove", fixability="safe",
    )
    action = HealthRepairAction(
        action_id="health-authorized", finding_id=authorized.finding_id,
        finding_code=authorized.code, classification="safe", decision="selected",
        operation="normalize_whitespace", sheet="Data", location="A2",
        impact_scope="Data!A2", current_state={}, proposed_state={"status": "ready"},
    )

    _, new, failures = verification._reconcile(
        SimpleNamespace(findings=(authorized, unauthorized)),
        SimpleNamespace(findings=(introduced,)),
        (action,),
    )

    assert [item["finding_id"] for item in new] == [introduced.finding_id]
    assert {item["code"] for item in failures} == {
        "unauthorized_finding_changed", "new_finding_introduced",
    }


def test_source_change_after_candidate_rescan_blocks_final_publication(
    tmp_path: Path, monkeypatch,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    action = _whitespace_actions(source)[0]
    original_reconcile = verification._reconcile

    def mutate_source_before_publish(before, after, actions):
        result = original_reconcile(before, after, actions)
        source.write_bytes(source.read_bytes() + b"changed-after-rescan")
        return result

    monkeypatch.setattr(verification, "_reconcile", mutate_source_before_publish)
    with pytest.raises(HealthRepairApplyError, match="source changed before verified"):
        repair_and_verify_workbook_health(
            source, output, selected_action_ids=[action.action_id]
        )

    assert not output.exists()
    assert not list(tmp_path.glob(".*-health-verify-*.xlsx"))


def test_candidate_same_inode_rewrite_during_link_is_detected_and_unpublished(
    tmp_path: Path, monkeypatch,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    action = _whitespace_actions(source)[0]
    original_link = verification._DirectoryGuard.link

    def rewrite_then_link(guard, candidate, destination):
        if destination == output.absolute():
            with candidate.open("r+b") as stream:
                stream.seek(0)
                stream.write(b"corrupt-after-digest")
                stream.truncate()
        return original_link(guard, candidate, destination)

    monkeypatch.setattr(verification._DirectoryGuard, "link", rewrite_then_link)
    with pytest.raises(HealthRepairApplyError, match="changed during publication"):
        repair_and_verify_workbook_health(
            source, output, selected_action_ids=[action.action_id]
        )

    assert not output.exists()
    assert not list(tmp_path.glob(".*-health-verify-*.xlsx"))


def test_source_change_during_final_link_removes_published_output(
    tmp_path: Path, monkeypatch,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    action = _whitespace_actions(source)[0]
    original_link = verification._DirectoryGuard.link

    def mutate_source_then_link(guard, candidate, destination):
        if destination == output.absolute():
            source.write_bytes(source.read_bytes() + b"changed-during-link")
        return original_link(guard, candidate, destination)

    monkeypatch.setattr(verification._DirectoryGuard, "link", mutate_source_then_link)
    with pytest.raises(HealthRepairApplyError, match="source changed during verified"):
        repair_and_verify_workbook_health(
            source, output, selected_action_ids=[action.action_id]
        )

    assert not output.exists()
    assert not list(tmp_path.glob(".*-health-verify-*.xlsx"))


@pytest.mark.skipif(os.name == "nt", reason="POSIX permits unlinking an open hardlink")
def test_replaced_final_identity_is_removed_without_touching_open_candidate(
    tmp_path: Path, monkeypatch,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    action = _whitespace_actions(source)[0]
    original_link = verification._DirectoryGuard.link

    def replace_published_link(guard, candidate, destination):
        original_link(guard, candidate, destination)
        if destination == output.absolute():
            destination.unlink()
            destination.write_bytes(b"foreign-corrupt-final")

    monkeypatch.setattr(verification._DirectoryGuard, "link", replace_published_link)
    with pytest.raises(HealthRepairApplyError, match="identity does not match"):
        repair_and_verify_workbook_health(
            source, output, selected_action_ids=[action.action_id]
        )

    assert not output.exists()
    assert not list(tmp_path.glob(".*-health-verify-*.xlsx"))


def test_scan_and_plan_hash_mismatch_fails_closed() -> None:
    scan = SimpleNamespace(
        source="book.xlsx", source_sha256=sha256(b"scan").hexdigest(), findings=(),
    )
    plan = SimpleNamespace(source_sha256=sha256(b"plan").hexdigest(), actions=())

    with pytest.raises(HealthRepairApplyError, match="between health scan and repair planning"):
        verification.health_workflow_contract(scan, plan, mode="dry_run")
