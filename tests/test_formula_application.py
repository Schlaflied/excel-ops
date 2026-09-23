import datetime as dt
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from excel_ops.formula_application import (
    FormulaApplicationError,
    apply_formula_plan,
    plan_formula_application,
)
from excel_ops.formula_planning import DateAddFormulaSpec, plan_formula


def _source(path: Path, *, protected_formula: bool = False, merged: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["C2"] = "2026-01-01"
    sheet["C3"] = "2026-02-01"
    sheet["C4"] = "2026-03-01"
    if protected_formula:
        sheet["D3"] = "=C3"
    if merged:
        sheet.merge_cells("D3:D4")
    workbook.save(path)
    workbook.close()


def _formula_plan():
    return plan_formula(
        DateAddFormulaSpec("Move each renewal date forward one month.", "C2", 1),
        target_excel_version="365",
    )


def test_dry_run_translates_formulas_without_writing(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)

    result = plan_formula_application(
        source, output, _formula_plan(), sheet="Report", target_range="D2:D4"
    )

    assert result.dry_run is True
    assert result.verified is False
    assert [item.cell for item in result.changes] == ["D2", "D3", "D4"]
    assert not output.exists()


def test_confirmed_application_writes_copy_and_verifies_relative_references(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)

    result = apply_formula_plan(
        source,
        output,
        _formula_plan(),
        sheet="Report",
        target_range="d2:d4",
        confirmed=True,
    )

    assert result.verified is True
    written = load_workbook(output, data_only=False)
    try:
        assert [written["Report"][cell].value for cell in ("D2", "D3", "D4")] == [
            "=EDATE(C2,1)",
            "=EDATE(C3,1)",
            "=EDATE(C4,1)",
        ]
    finally:
        written.close()
    original = load_workbook(source, data_only=False)
    try:
        assert original["Report"]["D2"].value is None
    finally:
        original.close()


def test_existing_formula_is_protected_by_default(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source, protected_formula=True)

    result = apply_formula_plan(
        source,
        output,
        _formula_plan(),
        sheet="Report",
        target_range="D2:D4",
        confirmed=True,
    )

    assert [(item.cell, item.reason) for item in result.skipped] == [
        ("D3", "protected_formula")
    ]
    written = load_workbook(output, data_only=False)
    try:
        assert written["Report"]["D3"].value == "=C3"
    finally:
        written.close()


def test_explicit_overwrite_replaces_existing_formula(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source, protected_formula=True)

    result = apply_formula_plan(
        source,
        output,
        _formula_plan(),
        sheet="Report",
        target_range="D2:D4",
        confirmed=True,
        overwrite_formulas=True,
    )

    assert result.skipped == ()


def test_array_formula_is_protected_by_default(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    workbook = load_workbook(source)
    workbook["Report"]["D3"] = ArrayFormula(ref="D3:D3", text="=C3")
    workbook.save(source)
    workbook.close()

    result = apply_formula_plan(
        source,
        output,
        _formula_plan(),
        sheet="Report",
        target_range="D2:D4",
        confirmed=True,
    )

    assert ("D3", "protected_formula") in [
        (item.cell, item.reason) for item in result.skipped
    ]


def test_merged_non_anchor_is_skipped(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source, merged=True)

    result = apply_formula_plan(
        source,
        output,
        _formula_plan(),
        sheet="Report",
        target_range="D2:D4",
        confirmed=True,
    )

    assert ("D4", "merged_non_anchor") in [
        (item.cell, item.reason) for item in result.skipped
    ]


def test_static_mode_requires_one_cell(tmp_path):
    source = tmp_path / "source.xlsx"
    _source(source)
    plan = plan_formula(
        DateAddFormulaSpec("Use the independently approved date.", "C2", 1),
        target_excel_version="365",
        output_mode="static",
        computed_value="2026-02-01",
    )

    with pytest.raises(FormulaApplicationError, match="single target cell"):
        plan_formula_application(
            source,
            tmp_path / "output.xlsx",
            plan,
            sheet="Report",
            target_range="D2:D3",
        )


def test_static_mode_rejects_formula_like_text(tmp_path):
    source = tmp_path / "source.xlsx"
    _source(source)
    plan = plan_formula(
        DateAddFormulaSpec("Use an independently computed value.", "C2", 1),
        target_excel_version="365",
        output_mode="static",
        computed_value="=1+1",
    )

    with pytest.raises(FormulaApplicationError, match="must not be formula text"):
        apply_formula_plan(
            source,
            tmp_path / "output.xlsx",
            plan,
            sheet="Report",
            target_range="D2",
            confirmed=True,
        )


def test_static_date_is_normalized_during_read_back_verification(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)
    plan = plan_formula(
        DateAddFormulaSpec("Use the independently approved date.", "C2", 1),
        target_excel_version="365",
        output_mode="static",
        computed_value=dt.date(2026, 2, 1),
    )

    result = apply_formula_plan(
        source,
        output,
        plan,
        sheet="Report",
        target_range="D2",
        confirmed=True,
    )

    assert result.verified is True
    written = load_workbook(output)
    try:
        assert written["Report"]["D2"].value == dt.datetime(2026, 2, 1)
    finally:
        written.close()


def test_write_requires_confirmation_and_new_output(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _source(source)

    with pytest.raises(FormulaApplicationError, match="confirmed=True"):
        apply_formula_plan(
            source, output, _formula_plan(), sheet="Report", target_range="D2"
        )
    output.write_bytes(b"existing")
    with pytest.raises(FormulaApplicationError, match="already exists"):
        apply_formula_plan(
            source,
            output,
            _formula_plan(),
            sheet="Report",
            target_range="D2",
            confirmed=True,
        )
