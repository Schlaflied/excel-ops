from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel_ops.delivery import (
    DeliveryPlanError,
    DeliveryTarget,
    _mapping_payload,
    load_delivery_targets,
    plan_delivery,
    run_delivery,
)
from excel_ops.matching import Destination
from excel_ops.number_formats import (
    FieldFormatPolicy,
    FormatPolicy,
    NumberFormatPolicyError,
    build_number_format,
    resolve_format_policy,
)
from excel_ops.template_writer import TemplateMapping, TemplateWriteError, write_template


def _template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    headers = ("CAD pay", "USD pay", "CNY pay", "FX", "Hours", "Headcount", "Rate")
    for index, header in enumerate(headers, start=1):
        sheet.cell(1, index).value = header
        sheet.cell(2, index).number_format = "General"
    workbook.save(path)


def _policy() -> FormatPolicy:
    return FormatPolicy.from_mapping(
        {
            "locale": "en-CA",
            "currency": "CAD",
            "fields": {
                "cad_pay": {"kind": "amount", "style": "accounting", "decimals": 2},
                "usd_pay": {"kind": "amount", "currency": "USD", "negative_style": "red"},
                "cny_pay": {"kind": "amount", "currency": "CNY"},
                "fx": {"kind": "exchange_rate", "decimals": 6, "calculation_decimals": 8},
                "hours": {"kind": "hours", "decimals": 2},
                "headcount": {"kind": "headcount", "decimals": 0},
                "rate": {"kind": "percentage", "decimals": 2},
            },
        }
    )


def test_cad_usd_cny_percentage_and_accounting_formats_preserve_values(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    _template(template)
    record = {
        "cad_pay": -1234.567,
        "usd_pay": -90.5,
        "cny_pay": 700.125,
        "fx": 1.37654321,
        "hours": 37.555,
        "headcount": 12,
        "rate": 0.1375,
    }
    mapping = TemplateMapping(
        sheet="Report",
        field_columns={name: index for index, name in enumerate(record, start=1)},
        format_policy=_policy(),
    )

    result = write_template(template, [record], mapping, output_path=output)

    workbook = load_workbook(output, data_only=False)
    sheet = workbook["Report"]
    assert [sheet.cell(2, index).value for index in range(1, 8)] == list(record.values())
    assert "CA$" in sheet["A2"].number_format
    assert "* " in sheet["A2"].number_format
    assert sheet["A2"].number_format != build_number_format(
        FieldFormatPolicy("amount", currency="CAD", negative_style="parentheses"),
        currency="CAD",
    )
    assert "US$" in sheet["B2"].number_format
    assert "[Red]" in sheet["B2"].number_format
    assert "¥" in sheet["C2"].number_format
    assert "000000" in sheet["D2"].number_format
    assert sheet["F2"].number_format.startswith("#,##0;")
    assert sheet["G2"].number_format.startswith("0.00%")
    assert sheet["F2"].number_format.split(";")[2] == "#,##0"
    assert sheet["G2"].number_format.split(";")[2] == "0.00%"
    workbook.close()
    assert len(result.format_changes) == 7
    assert result.format_policy["fields"]["fx"]["calculation_decimals"] == 8
    assert json.loads(result.change_log_path.read_text(encoding="utf-8"))["formatted_cells"] == 7


def test_mixed_or_ambiguous_currency_never_silently_uses_one_format(tmp_path: Path):
    mixed = FormatPolicy(
        currency_field="currency",
        fields={"amount": FieldFormatPolicy("amount")},
    )
    with pytest.raises(NumberFormatPolicyError, match="mixed currencies") as caught:
        resolve_format_policy(
            mixed,
            [{"amount": 10, "currency": "CAD"}, {"amount": 20, "currency": "USD"}],
        )
    assert caught.value.code == "mixed_currency_confirmation_required"

    with pytest.raises(NumberFormatPolicyError, match="no currency") as missing:
        resolve_format_policy(
            mixed,
            [{"amount": 10, "currency": ""}, {"amount": 20, "currency": "CAD"}],
        )
    assert missing.value.code == "currency_confirmation_required"

    explicit = FormatPolicy(
        currency="USD",
        currency_field="currency",
        fields={"amount": FieldFormatPolicy("amount", currency="CAD")},
    )
    assert resolve_format_policy(explicit, [{"amount": 10, "currency": ""}])["amount"].currency == "CAD"

    with pytest.raises(NumberFormatPolicyError, match="ambiguous"):
        FormatPolicy(currency="$", fields={"amount": FieldFormatPolicy("amount")})

    template = tmp_path / "template.xlsx"
    _template(template)
    mapping = TemplateMapping(
        sheet="Report",
        field_columns={"amount": "A", "currency": "B"},
        format_policy=mixed,
    )
    with pytest.raises(TemplateWriteError, match="mixed_currency_confirmation_required"):
        write_template(
            template,
            [{"amount": 10, "currency": "CAD"}, {"amount": 20, "currency": "USD"}],
            mapping,
            output_path=tmp_path / "blocked.xlsx",
        )
    assert not (tmp_path / "blocked.xlsx").exists()


def test_workbook_currency_is_default_and_field_currency_overrides_it():
    resolved = resolve_format_policy(_policy(), [{"cad_pay": 1, "usd_pay": 2, "cny_pay": 3}])
    assert resolved["cad_pay"].currency == "CAD"
    assert resolved["usd_pay"].currency == "USD"
    assert resolved["cny_pay"].currency == "CNY"


def test_storage_calculation_and_display_precision_remain_distinct():
    rule = FieldFormatPolicy(
        "exchange_rate",
        display_decimals=6,
        storage_decimals=12,
        calculation_decimals=10,
        rounding="half_even",
    )
    resolved = resolve_format_policy(FormatPolicy(fields={"fx": rule}), [{"fx": 1.234567891234}])["fx"]
    assert resolved.display_decimals == 6
    assert resolved.storage_decimals == 12
    assert resolved.calculation_decimals == 10
    assert resolved.rounding == "half_even"
    assert build_number_format(rule).startswith("#,##0.000000")


@pytest.mark.parametrize("value", [1.9, "2", True, -1, 16])
def test_recipe_precision_rejects_fractional_string_boolean_and_out_of_range_values(value):
    with pytest.raises(NumberFormatPolicyError) as caught:
        FormatPolicy.from_mapping(
            {"fields": {"hours": {"kind": "hours", "display_decimals": value}}}
        )
    assert caught.value.code == "invalid_precision"


def test_format_policy_is_part_of_the_idempotency_mapping_component():
    def target(decimals: int) -> DeliveryTarget:
        return DeliveryTarget(
            destination=Destination("North", ()),
            template_path="template.xlsx",
            mapping=TemplateMapping(
                sheet="Report",
                field_columns={"record_id": "A", "confidence": "B"},
                format_policy=FormatPolicy(
                    fields={"confidence": FieldFormatPolicy("percentage", display_decimals=decimals)}
                ),
            ),
        )

    assert _mapping_payload([target(1)]) != _mapping_payload([target(2)])


def test_json_recipe_loads_workbook_defaults_and_field_overrides(tmp_path: Path):
    payload = {
        "inputs": ["input.csv"],
        "staging_dir": "staging",
        "delivery_dir": "delivery",
        "format_policy": {
            "locale": "en-CA",
            "currency": "CAD",
            "defaults": {"negative_style": "red", "rounding": "half_even"},
            "fields": {"confidence": {"kind": "percentage", "decimals": 1}},
        },
        "targets": [
            {
                "key": "North",
                "template": "template.xlsx",
                "sheet": "Report",
                "field_columns": {"record_id": "A", "confidence": "B"},
            }
        ],
    }

    targets, _ = load_delivery_targets(payload, base_dir=tmp_path)

    policy = targets[0].mapping.format_policy
    assert policy is not None
    assert policy.currency == "CAD"
    assert policy.fields["confidence"].kind == "percentage"
    assert policy.fields["confidence"].decimals == 1
    assert policy.fields["confidence"].negative_style == "red"
    assert policy.fields["confidence"].rounding == "half_even"

    payload["format_policy"] = False
    with pytest.raises(DeliveryPlanError, match="target North format_policy must be an object"):
        load_delivery_targets(payload, base_dir=tmp_path)


def test_delivery_manifest_records_the_applied_policy_and_persisted_format(tmp_path: Path):
    source = tmp_path / "source.csv"
    source.write_text(
        "Location,Date,Asset ID,Category,Amount,Currency\n"
        "North,2026-09-21,A-1,routine,1250.75,CAD\n",
        encoding="utf-8",
    )
    template = tmp_path / "delivery-template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append(["record_id", "amount", "currency"])
    workbook.save(template)
    policy = FormatPolicy(
        currency_field="currency",
        fields={"amount": FieldFormatPolicy("amount", display_decimals=2)},
    )
    target = DeliveryTarget(
        destination=Destination("North", ("North",)),
        template_path=template,
        mapping=TemplateMapping(
            sheet="Report",
            field_columns={"record_id": "A", "amount": "B", "currency": "C"},
            format_policy=policy,
        ),
        required_fields=("record_id", "amount", "currency"),
    )

    run = run_delivery(
        [source],
        [target],
        staging_dir=tmp_path / "staging",
        delivery_dir=tmp_path / "delivery",
    )

    assert run.delivered is True
    manifest = run.manifests[0]
    assert manifest.format_policy["fields"]["amount"]["currency"] == "CAD"
    persisted = load_workbook(manifest.output, read_only=True)
    assert persisted["Report"]["B2"].value == 1250.75
    assert "CA$" in persisted["Report"]["B2"].number_format
    assert persisted["Report"]["C2"].value == "CAD"
    persisted.close()
    reloaded = json.loads(manifest.manifest_path().read_text(encoding="utf-8"))
    assert reloaded["format_policy"] == manifest.format_policy


def test_delivery_plan_routes_missing_currency_to_the_ambiguity_queue(tmp_path: Path):
    source = tmp_path / "missing-currency.csv"
    source.write_text(
        "Location,Date,Asset ID,Category,Amount,Currency\n"
        "North,2026-09-21,A-1,routine,1250.75,\n",
        encoding="utf-8",
    )
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    workbook.active.title = "Report"
    workbook.active.append(["record_id", "amount", "currency"])
    workbook.save(template)
    target = DeliveryTarget(
        destination=Destination("North", ("North",)),
        template_path=template,
        mapping=TemplateMapping(
            sheet="Report",
            field_columns={"record_id": "A", "amount": "B", "currency": "C"},
            format_policy=FormatPolicy(
                currency_field="currency",
                fields={"amount": FieldFormatPolicy("amount")},
            ),
        ),
    )

    run = plan_delivery(
        [source],
        [target],
        staging_dir=tmp_path / "staging",
        delivery_dir=tmp_path / "delivery",
    )

    blocker = "format_policy:North:currency_confirmation_required:amount"
    assert run.plan.writable is False
    assert blocker in run.plan.blocking_items
    assert blocker in run.plan.unresolved_ambiguities


def test_excel_unsafe_large_integer_is_not_reported_as_preserved(tmp_path: Path):
    template = tmp_path / "large-int-template.xlsx"
    workbook = Workbook()
    workbook.active.title = "Report"
    workbook.active.append(["headcount"])
    workbook.save(template)
    mapping = TemplateMapping(
        sheet="Report",
        field_columns={"headcount": "A"},
        format_policy=FormatPolicy(fields={"headcount": FieldFormatPolicy("headcount")}),
    )
    output = tmp_path / "large-int-output.xlsx"

    with pytest.raises(TemplateWriteError, match="numeric_value_verification_failed"):
        write_template(
            template,
            [{"headcount": 9_007_199_254_740_993}],
            mapping,
            output_path=output,
        )
    assert not output.exists()
