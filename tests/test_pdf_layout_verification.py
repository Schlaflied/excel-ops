from pathlib import Path

from openpyxl import Workbook

from excel_ops.pdf_layout_verification import verify_pdf_layout


class _Box:
    width = 612
    height = 792


class _Page:
    mediabox = _Box()

    def __init__(self, text: str):
        self._text = text

    def extract_text(self) -> str:
        return self._text


class _BrokenPage:
    @property
    def mediabox(self):
        raise ValueError("broken media box")

    def extract_text(self) -> str:
        raise ValueError("broken content stream")


def _source(path: Path, *, rows: int = 2, repeat_header: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["A1"] = "Header"
    for row in range(2, rows + 1):
        sheet.cell(row, 1, f"row-{row}")
    sheet["B2"] = "Boundary"
    if repeat_header:
        sheet.print_title_rows = "1:1"
    workbook.save(path)
    workbook.close()


def test_layout_passes_when_pages_are_readable_and_boundaries_are_present(tmp_path, monkeypatch):
    source = tmp_path / "report.xlsx"
    _source(source)
    monkeypatch.setattr(
        "excel_ops.pdf_layout_verification.PdfReader",
        lambda path: type("Reader", (), {"pages": [_Page("Header Boundary")]})(),
    )

    result = verify_pdf_layout(source, tmp_path / "report.pdf", sheets=["Report"])

    assert result.passed is True
    assert result.pages == 1
    assert result.findings == ()


def test_layout_reports_clipping_unreadable_pages_and_missing_repeat_headers(
    tmp_path, monkeypatch
):
    source = tmp_path / "long-report.xlsx"
    _source(source, rows=60)
    monkeypatch.setattr(
        "excel_ops.pdf_layout_verification.PdfReader",
        lambda path: type("Reader", (), {"pages": [_Page("")]})(),
    )

    result = verify_pdf_layout(source, tmp_path / "report.pdf", sheets=["Report"])

    assert result.passed is False
    assert {item.code for item in result.findings} == {
        "unreadable_page",
        "repeated_header_not_configured",
        "boundary_content_missing",
    }


def test_layout_fails_closed_when_the_pdf_cannot_be_parsed(tmp_path, monkeypatch):
    source = tmp_path / "report.xlsx"
    _source(source)

    def fail(path):
        raise ValueError("broken xref")

    monkeypatch.setattr("excel_ops.pdf_layout_verification.PdfReader", fail)

    result = verify_pdf_layout(source, tmp_path / "report.pdf", sheets=["Report"])

    assert result.passed is False
    assert result.pages == 0
    assert [item.code for item in result.findings] == ["pdf_unreadable"]


def test_configured_repeat_header_must_appear_on_every_page(tmp_path, monkeypatch):
    source = tmp_path / "multipage.xlsx"
    _source(source, rows=60, repeat_header=True)
    monkeypatch.setattr(
        "excel_ops.pdf_layout_verification.PdfReader",
        lambda path: type(
            "Reader",
            (),
            {"pages": [_Page("Header Boundary row-60"), _Page("Boundary row-60")]},
        )(),
    )

    result = verify_pdf_layout(source, tmp_path / "report.pdf", sheets=["Report"])

    assert result.passed is False
    assert "repeated_header_missing" in {item.code for item in result.findings}


def test_page_read_errors_are_reported_without_stopping_later_pages(tmp_path, monkeypatch):
    source = tmp_path / "report.xlsx"
    _source(source)
    monkeypatch.setattr(
        "excel_ops.pdf_layout_verification.PdfReader",
        lambda path: type(
            "Reader", (), {"pages": [_BrokenPage(), _Page("Header Boundary")]}
        )(),
    )

    result = verify_pdf_layout(source, tmp_path / "report.pdf", sheets=["Report"])

    assert result.pages == 2
    assert {item.code for item in result.findings} == {
        "invalid_page_bounds",
        "unreadable_page",
    }


def test_boundary_tokens_ignore_formula_and_number_cells(tmp_path, monkeypatch):
    source = tmp_path / "report.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["A1"] = "Header"
    sheet["A2"] = "Boundary"
    sheet["A3"] = "=SUM(B3:B4)"
    sheet["A4"] = 42
    workbook.save(source)
    workbook.close()
    monkeypatch.setattr(
        "excel_ops.pdf_layout_verification.PdfReader",
        lambda path: type("Reader", (), {"pages": [_Page("Header Boundary")]})(),
    )

    result = verify_pdf_layout(source, tmp_path / "report.pdf", sheets=["Report"])

    assert result.passed is True
