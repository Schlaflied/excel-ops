from __future__ import annotations

from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import pytest

from excel_ops.health_models import HealthRepairPlanError
from excel_ops.health_repair_plan import plan_workbook_health_repairs
from excel_ops.health_scan import scan_workbook_health


def _digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _save_extended_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Date", "Amount", "Notes"])
    sheet.append([datetime(2026, 1, 2), 0, "Alpha"])
    sheet.append([datetime(2026, 2, 3), None, "Beta"])
    sheet["A2"].number_format = "yyyy-mm-dd"
    sheet["A3"].number_format = "mm/dd/yyyy"
    sheet["C3"].font = Font(name="Arial", color="FF0000")
    sheet.auto_filter.ref = "A1:B3"
    sheet.freeze_panes = "A3"
    sheet.print_area = "A1:B3"
    workbook.save(path)
    workbook.close()


def test_extended_scan_detects_declared_health_risks_without_writing(tmp_path: Path) -> None:
    source = tmp_path / "extended.xlsx"
    _save_extended_fixture(source)
    before = source.read_bytes()

    result = scan_workbook_health(source)

    assert source.read_bytes() == before
    codes = {finding.code for finding in result.findings}
    assert {
        "blank_zero_semantics",
        "filter_range_mismatch",
        "freeze_panes_mismatch",
        "inconsistent_cell_style",
        "mixed_date_formats",
        "print_area_mismatch",
    } <= codes


def test_declared_layout_that_covers_the_table_is_not_flagged(tmp_path: Path) -> None:
    source = tmp_path / "layout.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Amount"])
    sheet.append(["Alpha", 1])
    sheet.append(["Beta", 2])
    sheet.auto_filter.ref = "A1:B3"
    sheet.freeze_panes = "A2"
    sheet.print_area = "A1:B3"
    workbook.save(source)
    workbook.close()

    codes = {finding.code for finding in scan_workbook_health(source).findings}

    assert not codes & {
        "filter_range_mismatch",
        "freeze_panes_mismatch",
        "print_area_mismatch",
    }


def test_plan_is_json_serializable_stable_and_traces_every_finding(tmp_path: Path) -> None:
    source = tmp_path / "plan.xlsx"
    _save_extended_fixture(source)
    before_digest = _digest(source)

    first = plan_workbook_health_repairs(source)
    second = plan_workbook_health_repairs(source)

    assert _digest(source) == before_digest
    assert first.to_dict() == second.to_dict()
    assert json.loads(json.dumps(first.to_dict(), ensure_ascii=False)) == first.to_dict()
    scan_ids = {finding.finding_id for finding in scan_workbook_health(source).findings}
    assert {action.finding_id for action in first.actions} == scan_ids
    assert len({action.action_id for action in first.actions}) == len(first.actions)
    assert first.to_dict()["counts"].keys() == {"safe", "review", "unsupported"}
    assert all(action.current_state["digest"] for action in first.actions)


def test_format_and_semantic_actions_fail_closed(tmp_path: Path) -> None:
    source = tmp_path / "closed.xlsx"
    _save_extended_fixture(source)

    plan = plan_workbook_health_repairs(source)
    by_code = {action.finding_code: action for action in plan.actions}

    for code in ("blank_zero_semantics", "inconsistent_cell_style", "mixed_date_formats"):
        assert by_code[code].classification == "review"
        assert by_code[code].decision == "proposed"
        assert by_code[code].proposed_state["status"] == "requires_declared_baseline"
    style_summary = by_code["inconsistent_cell_style"].current_state["summary"]
    assert len(style_summary["signatures"]) == 2
    assert all(item["signature_digest"] for item in style_summary["signatures"])
    blank_summary = by_code["blank_zero_semantics"].current_state["summary"]
    assert blank_summary["zero_locations"] == ["B2"]
    assert blank_summary["blank_locations"] == ["B3"]
    assert by_code["filter_range_mismatch"].proposed_state == {
        "status": "ready",
        "target": {"candidate_ref": "A1:C3"},
    }
    assert by_code["freeze_panes_mismatch"].proposed_state["target"] == {
        "candidate_coordinate": "A2"
    }
    assert by_code["print_area_mismatch"].proposed_state["target"] == {
        "candidate_refs": ["A1:C3"]
    }


def test_risky_content_changes_are_never_safe(tmp_path: Path) -> None:
    source = tmp_path / "risky.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Code", "Amount", "Comment"])
    sheet.append([" 42 ", 1, "merged"])
    sheet.append([None, None, None])
    sheet.append(["Alpha", 2, None])
    sheet.merge_cells("C2:D2")
    workbook.save(source)
    workbook.close()

    plan = plan_workbook_health_repairs(source)
    risky = {
        action.finding_code: action.classification
        for action in plan.actions
        if action.finding_code in {"numeric_text", "blank_data_rows", "abnormal_merged_cells"}
    }

    assert risky["numeric_text"] == "review"
    assert risky["blank_data_rows"] == "review"
    assert risky["abnormal_merged_cells"] == "unsupported"
    by_code = {action.finding_code: action for action in plan.actions}
    assert by_code["numeric_text"].proposed_state["target"] == {"candidate_value": 42}
    assert by_code["extra_whitespace"].proposed_state["target"] == {"candidate_value": "42"}


def test_callers_can_select_or_reject_individual_actions(tmp_path: Path) -> None:
    source = tmp_path / "decisions.xlsx"
    _save_extended_fixture(source)
    initial = plan_workbook_health_repairs(source)
    ready_ids = [
        action.action_id
        for action in initial.actions
        if action.classification == "review" and action.proposed_state["status"] == "ready"
    ]

    decided = plan_workbook_health_repairs(
        source,
        dispositions={ready_ids[0]: "selected", ready_ids[1]: "rejected"},
    )
    decisions = {action.action_id: action.decision for action in decided.actions}

    assert decisions[ready_ids[0]] == "selected"
    assert decisions[ready_ids[1]] == "rejected"
    assert all(action.action_id in decisions for action in initial.actions)


def test_unknown_and_unsupported_action_selection_fail_closed(tmp_path: Path) -> None:
    source = tmp_path / "unsupported.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Value"])
    sheet.append(["Alpha", 1])
    sheet.merge_cells("A2:B2")
    workbook.save(source)
    workbook.close()
    initial = plan_workbook_health_repairs(source)
    unsupported = next(action for action in initial.actions if action.classification == "unsupported")

    with pytest.raises(HealthRepairPlanError, match="unsupported action"):
        plan_workbook_health_repairs(source, dispositions={unsupported.action_id: "selected"})
    with pytest.raises(HealthRepairPlanError, match="unknown or stale action"):
        plan_workbook_health_repairs(source, dispositions={"health-missing": "rejected"})


def test_baseline_actions_cannot_be_selected_without_validated_targets(tmp_path: Path) -> None:
    source = tmp_path / "baselines.xlsx"
    _save_extended_fixture(source)
    initial = plan_workbook_health_repairs(source)
    by_code = {action.finding_code: action for action in initial.actions}
    style = by_code["inconsistent_cell_style"]

    with pytest.raises(HealthRepairPlanError, match="requires a validated baseline"):
        plan_workbook_health_repairs(source, dispositions={style.action_id: "selected"})
    with pytest.raises(HealthRepairPlanError, match="current target_signature_digest"):
        plan_workbook_health_repairs(
            source,
            baselines={style.action_id: {"target_signature_digest": "stale"}},
        )
    ready_action = next(
        action for action in initial.actions if action.proposed_state["status"] == "ready"
    )
    with pytest.raises(HealthRepairPlanError, match="does not accept a baseline"):
        plan_workbook_health_repairs(
            source, baselines={ready_action.action_id: {"unexpected": "value"}}
        )

    style_digest = style.current_state["summary"]["signatures"][0]["signature_digest"]
    baselines = {
        style.action_id: {"target_signature_digest": style_digest},
        by_code["mixed_date_formats"].action_id: {
            "target_number_format": "yyyy-mm-dd",
            "locale": "en-CA",
        },
        by_code["blank_zero_semantics"].action_id: {
            "blank_action": "preserve",
            "zero_action": "preserve",
        },
    }
    dispositions = {action_id: "selected" for action_id in baselines}

    decided = plan_workbook_health_repairs(
        source, dispositions=dispositions, baselines=baselines
    )

    selected = {action.action_id: action for action in decided.actions if action.decision == "selected"}
    assert set(selected) == set(baselines)
    assert all(action.proposed_state["status"] == "ready" for action in selected.values())


def test_action_ids_cannot_cross_workbooks_or_source_versions(tmp_path: Path) -> None:
    first_source = tmp_path / "first.xlsx"
    second_source = tmp_path / "second.xlsx"
    _save_extended_fixture(first_source)
    _save_extended_fixture(second_source)
    workbook = load_workbook(second_source)
    workbook["Data"]["C2"] = "different"
    workbook.save(second_source)
    workbook.close()
    first = plan_workbook_health_repairs(first_source)
    old_id = first.actions[0].action_id

    with pytest.raises(HealthRepairPlanError, match="unknown or stale action"):
        plan_workbook_health_repairs(second_source, dispositions={old_id: "rejected"})

    workbook = load_workbook(first_source)
    workbook["Data"]["C2"] = "changed after planning"
    workbook.save(first_source)
    workbook.close()
    with pytest.raises(HealthRepairPlanError, match="unknown or stale action"):
        plan_workbook_health_repairs(first_source, dispositions={old_id: "rejected"})


def test_style_and_blank_zero_distribution_changes_invalidate_state(tmp_path: Path) -> None:
    source = tmp_path / "distribution.xlsx"
    _save_extended_fixture(source)
    before = {
        action.finding_code: action
        for action in plan_workbook_health_repairs(source).actions
    }
    workbook = load_workbook(source)
    sheet = workbook["Data"]
    sheet["B2"] = None
    sheet["B3"] = 0
    sheet["C2"].font = Font(name="Arial", color="FF0000")
    sheet["C3"].font = Font(name="Calibri", color="0000FF")
    workbook.save(source)
    workbook.close()
    after = {
        action.finding_code: action
        for action in plan_workbook_health_repairs(source).actions
    }

    for code in ("blank_zero_semantics", "inconsistent_cell_style"):
        assert before[code].action_id != after[code].action_id
        assert before[code].current_state["digest"] != after[code].current_state["digest"]


def test_large_sparse_sheet_reports_budget_instead_of_scanning_rectangle(tmp_path: Path) -> None:
    source = tmp_path / "sparse.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sparse"
    sheet.append(["Name", "Amount"])
    sheet.append(["Alpha", 0])
    sheet["XFD1048576"] = "stray"
    workbook.save(source)
    workbook.close()

    result = scan_workbook_health(source)
    budget = next(
        finding for finding in result.findings
        if finding.code == "blank_zero_scan_budget_exceeded"
    )

    assert budget.fixability == "unsupported"
    assert budget.details["declared_cells"] > budget.details["cell_budget"]


def test_special_sheet_name_and_multiple_print_ranges_parse_safely(tmp_path: Path) -> None:
    source = tmp_path / "special-sheet.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "A ! 中文"
    sheet.append(["Name", "Amount", "Note"])
    sheet.append(["Alpha", 1, "x"])
    sheet.append(["Beta", 2, "y"])
    sheet.print_area = "A1:C3,E1:E3"
    workbook.save(source)
    workbook.close()

    result = scan_workbook_health(source)

    assert not [finding for finding in result.findings if finding.code == "print_area_mismatch"]


def test_numeric_candidates_preserve_precision_without_float_rounding(tmp_path: Path) -> None:
    source = tmp_path / "numeric-precision.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Label", "Value"])
    sheet.append(["unsafe integer", "9007199254740993"])
    sheet.append(["very long", "9" * 400])
    sheet.append(["decimal", "123.4500"])
    sheet.append(["percentage", "12.5%"])
    workbook.save(source)
    workbook.close()

    plan = plan_workbook_health_repairs(source)
    by_location = {
        action.location: action
        for action in plan.actions
        if action.finding_code == "numeric_text"
    }

    for location in ("B2", "B3"):
        assert by_location[location].proposed_state["status"] == "requires_declared_baseline"
        assert by_location[location].proposed_state["target"] is None
    assert by_location["B4"].proposed_state["target"] == {
        "candidate_value": {"kind": "decimal", "canonical": "123.45"}
    }
    assert by_location["B5"].proposed_state["target"] == {
        "candidate_value": {"kind": "decimal", "canonical": "0.125"}
    }

    unsafe = by_location["B2"]
    with pytest.raises(HealthRepairPlanError, match="preserve_text"):
        plan_workbook_health_repairs(
            source,
            dispositions={unsafe.action_id: "selected"},
            baselines={unsafe.action_id: {"target_representation": "number"}},
        )
    preserved = plan_workbook_health_repairs(
        source,
        dispositions={unsafe.action_id: "selected"},
        baselines={unsafe.action_id: {"target_representation": "preserve_text"}},
    )
    selected = next(action for action in preserved.actions if action.action_id == unsafe.action_id)
    assert selected.proposed_state["target"] == {
        "target_representation": "preserve_text",
        "value": "9007199254740993",
    }


def test_duplicate_header_baseline_checks_the_complete_header_row(tmp_path: Path) -> None:
    source = tmp_path / "duplicate-headers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", " amount ", "Amount", "Notes"])
    sheet.append(["Alpha", 1, 2, "ok"])
    workbook.save(source)
    workbook.close()
    initial = plan_workbook_health_repairs(source)
    action = next(
        item for item in initial.actions if item.finding_code == "duplicate_header"
    )
    summary = action.current_state["summary"]

    assert summary["current_normalized"] == "amount"
    assert [(item["coordinate"], item["normalized"]) for item in summary["headers"]] == [
        ("A1", "name"),
        ("B1", "amount"),
        ("C1", "amount"),
        ("D1", "notes"),
    ]
    with pytest.raises(HealthRepairPlanError, match="must differ"):
        plan_workbook_health_repairs(
            source, baselines={action.action_id: {"target_header": " Amount "}}
        )
    with pytest.raises(HealthRepairPlanError, match="conflicts with another header"):
        plan_workbook_health_repairs(
            source, baselines={action.action_id: {"target_header": "  NAME  "}}
        )

    valid = plan_workbook_health_repairs(
        source,
        dispositions={action.action_id: "selected"},
        baselines={action.action_id: {"target_header": "Gross   Amount"}},
    )
    renamed = next(item for item in valid.actions if item.action_id == action.action_id)
    assert renamed.proposed_state["target"] == {
        "target_header": "Gross Amount",
        "normalized": "gross amount",
    }


def test_selected_duplicate_header_targets_must_be_unique_within_header_row(
    tmp_path: Path,
) -> None:
    source = tmp_path / "multiple-duplicate-headers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Amount", " amount ", "AMOUNT"])
    sheet.append(["Alpha", 1, 2, 3])
    workbook.save(source)
    workbook.close()
    initial = plan_workbook_health_repairs(source)
    actions = [
        item for item in initial.actions if item.finding_code == "duplicate_header"
    ]
    assert len(actions) == 2
    baselines = {
        action.action_id: {"target_header": "Revenue"} for action in actions
    }

    proposed = plan_workbook_health_repairs(source, baselines=baselines)
    assert all(action.decision == "proposed" for action in proposed.actions)
    plan_workbook_health_repairs(
        source,
        dispositions={
            actions[0].action_id: "selected",
            actions[1].action_id: "rejected",
        },
        baselines=baselines,
    )
    with pytest.raises(HealthRepairPlanError, match="targets conflict"):
        plan_workbook_health_repairs(
            source,
            dispositions={action.action_id: "selected" for action in actions},
            baselines=baselines,
        )
