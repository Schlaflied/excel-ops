from pathlib import Path
from subprocess import CompletedProcess

import pytest
from openpyxl import Workbook, load_workbook

from excel_ops.pdf_export import PdfExportError, export_pdf


def _source(path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "Report"
    workbook.active["A1"] = "synthetic"
    workbook.create_sheet("Archive")
    workbook.save(path)


def _renderer(command, **kwargs):
    outdir = Path(command[command.index("--outdir") + 1])
    prepared = Path(command[-1])
    workbook = load_workbook(prepared, read_only=True)
    try:
        assert workbook.sheetnames == ["Report"]
    finally:
        workbook.close()
    (outdir / "prepared.pdf").write_bytes(b"%PDF-1.7\nsynthetic")
    return CompletedProcess(command, 0, "", "")


def test_pdf_export_uses_selected_sheets_and_verifies_signature(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "report.pdf"
    _source(source)
    assert export_pdf(source, output, sheets=["Report"], soffice="soffice", runner=_renderer) == output
    assert output.read_bytes().startswith(b"%PDF-")


def test_pdf_export_requires_renderer_and_explicit_valid_scope(tmp_path, monkeypatch):
    source = tmp_path / "source.xlsx"
    _source(source)
    monkeypatch.setattr("excel_ops.pdf_export.shutil.which", lambda _: None)
    with pytest.raises(PdfExportError, match="renderer unavailable"):
        export_pdf(source, tmp_path / "out.pdf", soffice=None)
    with pytest.raises(PdfExportError, match="at least one sheet"):
        export_pdf(source, tmp_path / "out.pdf", sheets=[], soffice="soffice")
    with pytest.raises(PdfExportError, match="unknown worksheet"):
        export_pdf(source, tmp_path / "out.pdf", sheets=["Missing"], soffice="soffice")
    with pytest.raises(PdfExportError, match="contains duplicates"):
        export_pdf(source, tmp_path / "out.pdf", sheets=["Report", "Report"], soffice="soffice")


def test_pdf_export_rejects_non_pdf_renderer_output(tmp_path):
    source = tmp_path / "source.xlsx"
    _source(source)

    def invalid_renderer(command, **kwargs):
        outdir = Path(command[command.index("--outdir") + 1])
        (outdir / "prepared.pdf").write_bytes(b"not-a-pdf")
        return CompletedProcess(command, 0, "", "")

    with pytest.raises(PdfExportError, match="not a PDF"):
        export_pdf(source, tmp_path / "out.pdf", soffice="soffice", runner=invalid_renderer)


def test_pdf_export_reports_renderer_start_failure(tmp_path):
    source = tmp_path / "source.xlsx"
    _source(source)

    def missing_renderer(command, **kwargs):
        raise FileNotFoundError("soffice")

    with pytest.raises(PdfExportError, match="renderer could not run"):
        export_pdf(source, tmp_path / "out.pdf", soffice="soffice", runner=missing_renderer)
