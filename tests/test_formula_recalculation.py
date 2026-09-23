from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_ops.formula_recalculation import (
    FormulaEngineUnavailable,
    FormulaValueExpectation,
    verify_formula_recalculation,
)
from excel_ops.formula_verification import FormulaRegion


def _source(path: Path, formula: str = "=SUM(A2:B2)") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["A2"] = 2
    sheet["B2"] = 3
    sheet["C2"] = formula
    workbook.save(path)
    workbook.close()


def _calculated(value):
    def recalculate(source: Path, destination: Path) -> str:
        workbook = load_workbook(source, data_only=False)
        workbook["Report"]["C2"] = value
        workbook.save(destination)
        workbook.close()
        return "test-engine"

    return recalculate


def test_recalculation_verifies_an_independent_expected_value(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "recalculated.xlsx"
    _source(source)

    result = verify_formula_recalculation(
        source,
        output,
        expectations=(FormulaValueExpectation("Report", "C2", 5),),
        formula_regions=(FormulaRegion("Report", "C2:C2"),),
        recalculator=_calculated(5),
    )

    assert result.status == "verified"
    assert result.passed is True
    assert result.engine == "test-engine"
    assert result.recalculated_path == output.resolve()
    assert output.exists()


def test_recalculation_uses_numeric_tolerance(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "recalculated.xlsx"
    _source(source)

    result = verify_formula_recalculation(
        source,
        output,
        expectations=(
            FormulaValueExpectation("Report", "C2", Decimal("5.00"), Decimal("0.01")),
        ),
        recalculator=_calculated(5.005),
    )

    assert result.status == "verified"


def test_mismatch_fails_without_publishing_output(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "recalculated.xlsx"
    _source(source)

    result = verify_formula_recalculation(
        source,
        output,
        expectations=(FormulaValueExpectation("Report", "C2", 6),),
        recalculator=_calculated(5),
    )

    assert result.status == "failed"
    assert {item.code for item in result.findings} == {"recalculation_mismatch"}
    assert not output.exists()


def test_recalculated_formula_error_is_blocking(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "recalculated.xlsx"
    _source(source)

    result = verify_formula_recalculation(
        source,
        output,
        expectations=(FormulaValueExpectation("Report", "C2", 5),),
        recalculator=_calculated("#N/A"),
    )

    assert result.status == "failed"
    assert {item.code for item in result.findings} == {
        "recalculated_formula_error"
    }


def test_recalculation_scans_errors_outside_expected_cells(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "recalculated.xlsx"
    _source(source)

    def recalculate(source: Path, destination: Path) -> str:
        workbook = load_workbook(source)
        workbook["Report"]["C2"] = 5
        workbook["Report"]["D8"] = "#REF!"
        workbook.save(destination)
        workbook.close()
        return "test-engine"

    result = verify_formula_recalculation(
        source,
        output,
        expectations=(FormulaValueExpectation("Report", "C2", 5),),
        recalculator=recalculate,
    )

    assert result.status == "failed"
    finding = next(
        item for item in result.findings if item.code == "recalculated_formula_error"
    )
    assert (finding.sheet, finding.cell) == ("Report", "D8")


def test_static_formula_failure_blocks_engine_execution(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "recalculated.xlsx"
    _source(source, "=#REF!")
    called = False

    def should_not_run(source: Path, destination: Path) -> str:
        nonlocal called
        called = True
        return "test-engine"

    result = verify_formula_recalculation(
        source,
        output,
        expectations=(FormulaValueExpectation("Report", "C2", 5),),
        recalculator=should_not_run,
    )

    assert result.status == "failed"
    assert "broken_reference" in {item.code for item in result.findings}
    assert called is False


def test_missing_expectations_are_unverified(tmp_path):
    source = tmp_path / "source.xlsx"
    _source(source)

    result = verify_formula_recalculation(
        source,
        tmp_path / "recalculated.xlsx",
        expectations=(),
        recalculator=_calculated(5),
    )

    assert result.status == "unverified"
    assert [item.code for item in result.findings] == [
        "recalculation_expectations_missing"
    ]


def test_missing_engine_is_reported_as_unverified(tmp_path):
    source = tmp_path / "source.xlsx"
    _source(source)

    def unavailable(source: Path, destination: Path) -> str:
        raise FormulaEngineUnavailable("No calculation engine is installed")

    result = verify_formula_recalculation(
        source,
        tmp_path / "recalculated.xlsx",
        expectations=(FormulaValueExpectation("Report", "C2", 5),),
        recalculator=unavailable,
    )

    assert result.status == "unverified"
    assert [item.code for item in result.findings] == [
        "recalculation_not_verified"
    ]


def test_circular_reference_is_detected_before_recalculation(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["A1"] = "=B1"
    sheet["B1"] = "=A1"
    workbook.save(source)
    workbook.close()

    result = verify_formula_recalculation(
        source,
        tmp_path / "recalculated.xlsx",
        expectations=(FormulaValueExpectation("Report", "A1", 1),),
        recalculator=_calculated(1),
    )

    assert result.status == "failed"
    assert "circular_reference" in {item.code for item in result.findings}
