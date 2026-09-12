from __future__ import annotations

from datetime import date

from openpyxl import Workbook, load_workbook

from excel_ops.period_refresh import DateRole, DateSlot, SlotTarget, refresh_workbook_period
from excel_ops.periods import resolve_period


def _template(path):
    workbook = Workbook()
    report = workbook.active
    report.title = "Report"
    report["A1"] = "Period"
    report["B1"] = "old"
    report["A2"] = "Generated"
    report["B2"] = date(2020, 1, 1)
    report["A3"] = "Deadline"
    report["B3"] = date(2020, 1, 2)
    report["A5"] = "Event date"
    report["A6"] = date(2024, 4, 3)
    empty = workbook.create_sheet("Empty")
    empty["A1"] = "Period"
    empty["B1"] = "old"
    workbook.save(path)


def _slots():
    return [
        DateSlot("Report", DateRole.PERIOD, cell="B1", template="{period_start:%Y-%m-%d} to {period_end:%Y-%m-%d}"),
        DateSlot("Report", DateRole.GENERATED_ON, cell="B2"),
        DateSlot("Report", DateRole.DEADLINE, cell="B3"),
        DateSlot("Report", DateRole.HISTORICAL, cell="A6"),
        DateSlot("Empty", DateRole.PERIOD, cell="B1", include_when_empty=True),
    ]


def test_refreshes_declared_slots_and_protects_history(tmp_path):
    source = tmp_path / "template.xlsx"
    output = tmp_path / "delivery.xlsx"
    _template(source)
    period = resolve_period("上一个完整周", as_of=date(2026, 1, 2))

    result = refresh_workbook_period(
        source,
        output,
        period=period,
        slots=_slots(),
        generated_on=date(2026, 1, 2),
        deadline=date(2026, 1, 5),
    )

    workbook = load_workbook(output, data_only=False)
    assert workbook["Report"]["B1"].value == "2025-12-22 to 2025-12-28"
    assert workbook["Report"]["B2"].value.date() == date(2026, 1, 2)
    assert workbook["Report"]["B3"].value.date() == date(2026, 1, 5)
    assert workbook["Report"]["A6"].value.date() == date(2024, 4, 3)
    assert workbook["Empty"]["B1"].value == period.display_text
    workbook.close()
    assert [change.status for change in result.changes] == ["updated", "updated", "updated", "protected", "updated"]


def test_dry_run_lists_changes_without_writing(tmp_path):
    source = tmp_path / "template.xlsx"
    output = tmp_path / "delivery.xlsx"
    _template(source)
    period = resolve_period("本月", as_of=date(2026, 2, 10))

    result = refresh_workbook_period(source, output, period=period, slots=_slots(), deadline=date(2026, 3, 3), dry_run=True)

    assert not output.exists()
    assert result.output_path is None
    assert result.changes[0].old_value == "old"
    assert result.changes[0].new_value == "2026-02-01 to 2026-02-28"


def test_sheet_title_and_empty_skip_are_scoped(tmp_path):
    source = tmp_path / "template.xlsx"
    output = tmp_path / "delivery.xlsx"
    _template(source)
    period = resolve_period("本季度", as_of=date(2026, 5, 20))
    slots = [
        DateSlot("Report", DateRole.PERIOD, target=SlotTarget.SHEET_TITLE, template="Q{period_end:%m}"),
        DateSlot("Empty", DateRole.PERIOD, cell="B1", include_when_empty=False),
    ]

    result = refresh_workbook_period(source, output, period=period, slots=slots)

    workbook = load_workbook(output)
    assert "Q06" in workbook.sheetnames
    assert workbook["Empty"]["B1"].value == "old"
    workbook.close()
    assert result.changes[1].status == "empty_skipped"


def test_repeated_run_is_idempotent(tmp_path):
    source = tmp_path / "template.xlsx"
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    _template(source)
    period = resolve_period("本月", as_of=date(2024, 2, 29))

    refresh_workbook_period(source, first, period=period, slots=_slots(), deadline=date(2024, 3, 5))
    result = refresh_workbook_period(first, second, period=period, slots=_slots(), deadline=date(2024, 3, 5))

    assert all(change.status in {"unchanged", "protected"} for change in result.changes)


def test_missing_declared_sheet_stops_without_partial_output(tmp_path):
    source = tmp_path / "template.xlsx"
    output = tmp_path / "delivery.xlsx"
    _template(source)
    period = resolve_period("本月", as_of=date(2026, 5, 20))

    try:
        refresh_workbook_period(source, output, period=period, slots=[DateSlot("Missing", DateRole.PERIOD, cell="A1")])
    except ValueError as error:
        assert "does not exist" in str(error)
    else:
        raise AssertionError("expected missing sheet to fail")
    assert not output.exists()
