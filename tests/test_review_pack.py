from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
import pytest

from excel_ops.review_pack import (
    ReviewPackRow,
    import_review_pack,
    review_rows_from_match_results,
    write_review_pack,
)


def row(record_id: str = "rec_demo_1") -> ReviewPackRow:
    return ReviewPackRow(
        record_id=record_id,
        source_value="Tornto Office",
        source_location="synthetic.xlsx:Sheet1!A2",
        proposed_value="Toronto",
        review_reason="fuzzy_requires_review",
        confidence=0.91,
        candidates=("Toronto", "London"),
    )


def pack_path() -> Path:
    return Path.cwd() / f".test-review-pack-{uuid4().hex}.xlsx"


def test_review_pack_round_trip_and_history():
    path = pack_path()
    try:
        write_review_pack([row(), row("rec_demo_2"), row("rec_demo_3")], path)
        wb = load_workbook(path)
        assert wb["Review"].protection.sheet
        assert wb["Review"]["H2"].protection.locked is False
        assert wb["Review"]["A2"].protection.locked is True
        assert wb["_Manifest"].sheet_state == "veryHidden"
        wb["Review"]["H2"] = "accept"
        wb["Review"]["H3"] = "correct"
        wb["Review"]["I3"] = "Toronto"
        wb["Review"]["J3"] = "Confirmed against source"
        wb["Review"]["H4"] = "cannot determine"
        wb.save(path)
        wb.close()

        result = import_review_pack(path)

        assert result.accepted_values == {"rec_demo_1": "Toronto", "rec_demo_2": "Toronto"}
        assert result.decisions["rec_demo_3"].decision == "cannot determine"
        assert "rec_demo_3" not in result.accepted_values
        assert len(result.history) == 3
        repeated = import_review_pack(path)
        assert len(repeated.history) == 3
        wb = load_workbook(path, read_only=True)
        assert wb["Review History"].max_row == 4
        wb.close()
    finally:
        path.unlink(missing_ok=True)


def test_blank_decision_remains_pending():
    path = pack_path()
    try:
        write_review_pack([row()], path)
        result = import_review_pack(path)
        assert result.pending_record_ids == ("rec_demo_1",)
        assert result.decisions == {}
    finally:
        path.unlink(missing_ok=True)


@pytest.mark.parametrize("decision", ["yes", "maybe", "ACCEPTED"])
def test_import_rejects_unknown_decision(decision):
    path = pack_path()
    try:
        write_review_pack([row()], path)
        wb = load_workbook(path)
        wb["Review"]["H2"] = decision
        wb.save(path)
        wb.close()
        with pytest.raises(ValueError, match="Invalid decision"):
            import_review_pack(path)
    finally:
        path.unlink(missing_ok=True)


def test_correct_requires_value():
    path = pack_path()
    try:
        write_review_pack([row()], path)
        wb = load_workbook(path)
        wb["Review"]["H2"] = "correct"
        wb.save(path)
        wb.close()
        with pytest.raises(ValueError, match="requires a corrected value"):
            import_review_pack(path)
    finally:
        path.unlink(missing_ok=True)


@pytest.mark.parametrize("damage", ["header", "missing_id", "duplicate_id", "candidate"])
def test_tampering_blocks_import(damage):
    path = pack_path()
    try:
        write_review_pack([row(), row("rec_demo_2")], path)
        wb = load_workbook(path)
        ws = wb["Review"]
        if damage == "header":
            ws["A1"] = "ID"
        elif damage == "missing_id":
            ws["A2"] = ""
        elif damage == "duplicate_id":
            ws["A3"] = "rec_demo_1"
        else:
            ws["G2"] = '["Montreal"]'
        wb.save(path)
        wb.close()
        with pytest.raises(ValueError):
            import_review_pack(path)
    finally:
        path.unlink(missing_ok=True)


@dataclass
class Candidate:
    destination_key: str


@dataclass
class Record:
    location: str
    source: str


@dataclass
class MatchResult:
    record_id: str
    record: Record
    status: str
    destination_key: str | None
    confidence: float
    candidates: tuple[Candidate, ...]
    reasons: tuple[str, ...]


def test_issue_2_match_results_have_a_compatibility_adapter():
    matched = MatchResult("rec_1", Record("Toronto", "a.xlsx:A2"), "matched", "Toronto", 1, (), ())
    review = MatchResult(
        "rec_2",
        Record("Tornto", "a.xlsx:A3"),
        "review",
        None,
        0.9,
        (Candidate("Toronto"),),
        ("fuzzy_requires_review",),
    )

    rows = review_rows_from_match_results([matched, review])

    assert len(rows) == 1
    assert rows[0].record_id == "rec_2"
    assert rows[0].candidates == ("Toronto",)
