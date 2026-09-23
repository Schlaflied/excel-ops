from datetime import datetime
from math import inf, nan
from pathlib import Path

import pytest
from openpyxl import Workbook

from excel_ops.ingestion import LayoutDetectionError, load_input_records
from excel_ops.models import ExtractedRecord


def test_two_xlsx_layouts_produce_same_contract(tmp_path: Path):
    period_path = tmp_path / "period-layout.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    ws.append(["Please use the Data tab"])
    data = wb.create_sheet("Arbitrary weekly upload")
    data.append(["Synthetic inspection export"])
    data.append([])
    data.append(["Location", "Period", "Identifier", "Category"])
    data.append(["100 Example Avenue", datetime(2026, 9, 8), "DEMO 123", "routine"])
    wb.save(period_path)

    datetime_path = tmp_path / "datetime-layout.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Not a fixed data name"
    ws.append(["Site", "Inspection Date", "Asset ID", "Type"])
    ws.append(["100 Example Avenue", "2026/09/08 08:30", "DEMO 123", "routine"])
    wb.save(datetime_path)

    period = load_input_records(period_path)
    timestamp = load_input_records(datetime_path)
    assert [(r.location, r.event_date, r.identifier, r.category) for r in period] == [
        (r.location, r.event_date, r.identifier, r.category) for r in timestamp
    ]
    assert period[0].source_file == "period-layout.xlsx"
    assert period[0].source_sheet == "Arbitrary weekly upload"
    assert period[0].source_row == 4


def test_excel_numeric_date_and_preamble_are_supported(tmp_path: Path):
    path = tmp_path / "numeric-date.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Synthetic preface"])
    ws.append([])
    ws.append(["Address", "Date", "Code", "Event Type"])
    ws.append(["200 Sample Street", 45543, "CHECK 9", "follow-up"])
    wb.save(path)

    records = load_input_records(path)
    assert records[0].event_date == "2024-09-08"


def test_csv_layout_is_normalized(tmp_path: Path):
    path = tmp_path / "inspection.csv"
    path.write_text(
        "report generated for a synthetic fixture\n\nsite,date time,id,type\n"
        "100 Example Avenue,2026-09-08 11:30,DEMO 123,routine\n",
        encoding="utf-8",
    )
    records = load_input_records(path)
    assert records[0].event_date == "2026-09-08"
    assert records[0].source_row == 4


def test_optional_amount_and_currency_flow_into_the_normalized_contract(tmp_path: Path):
    path = tmp_path / "amounts.csv"
    path.write_text(
        "Location,Date,Asset ID,Category,Gross Pay,Currency Code\n"
        "100 Example Avenue,2026-09-08,DEMO 123,routine,1234.50,CAD\n",
        encoding="utf-8",
    )

    record = load_input_records(path)[0]

    assert record.amount == 1234.5
    assert record.currency == "CAD"


@pytest.mark.parametrize("value", [nan, inf, -inf, "NaN", "Infinity"])
def test_non_finite_amounts_are_not_accepted_as_numeric_values(value):
    record = ExtractedRecord.from_dict({"amount": value}, "synthetic.json")
    assert record.amount is None


def test_embedded_currency_is_preserved_and_checked_against_explicit_currency():
    matching = ExtractedRecord.from_dict(
        {"amount": "US$1,234.50", "currency": "USD"}, "synthetic.json"
    )
    inferred = ExtractedRecord.from_dict({"amount": "CA$100"}, "synthetic.json")
    disambiguated = ExtractedRecord.from_dict(
        {"amount": "$100", "currency": "CAD"}, "synthetic.json"
    )

    assert matching.amount == 1234.5
    assert matching.currency == "USD"
    assert inferred.amount == 100
    assert inferred.currency == "CAD"
    assert disambiguated.currency == "CAD"
    with pytest.raises(ValueError, match="USD.*conflicts.*CAD"):
        ExtractedRecord.from_dict(
            {"amount": "US$100", "currency": "CAD"}, "synthetic.json"
        )
    with pytest.raises(ValueError, match=r"\$.*conflicts.*EUR"):
        ExtractedRecord.from_dict(
            {"amount": "$100", "currency": "EUR"}, "synthetic.json"
        )


def test_unrecognized_layout_stops_without_partial_output(tmp_path: Path):
    path = tmp_path / "unknown.xlsx"
    wb = Workbook()
    wb.active.append(["maybe", "some", "data"])
    wb.active.append(["100 Example Avenue", "DEMO 123", "routine"])
    wb.save(path)

    with pytest.raises(LayoutDetectionError, match="No worksheet"):
        load_input_records(path)


def test_image_json_retains_region_provenance(tmp_path: Path):
    path = tmp_path / "image-records.json"
    path.write_text(
        '{"source":"synthetic.png","records":[{"location":"100 Example Avenue",'
        '"event_date":"2026-09-08","identifier":"DEMO 123","category":"routine",'
        '"confidence":0.96,"source_region":"page=1;x=20;y=40;w=300;h=60"}]}',
        encoding="utf-8",
    )
    record = load_input_records(path)[0]
    assert record.source == "synthetic.png@page=1;x=20;y=40;w=300;h=60"
