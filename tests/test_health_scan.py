from __future__ import annotations

from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
import pytest

from excel_ops.health_models import HealthScanError
from excel_ops.health_scan import scan_workbook_health


def _digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _save_healthy(path: Path, *, sheet_name: str = "Data") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(["Name", "Amount"])
    sheet.append(["Alpha", 12.5])
    sheet.append(["Beta", 7])
    workbook.save(path)
    workbook.close()


def test_normal_workbook_returns_no_known_issues_and_json(tmp_path: Path) -> None:
    source = tmp_path / "healthy.xlsx"
    _save_healthy(source)

    result = scan_workbook_health(source)

    assert result.assessment == "no_known_issues"
    assert result.findings == ()
    assert result.source_sha256 == _digest(source)
    payload = result.to_dict()
    assert json.loads(json.dumps(payload, ensure_ascii=False)) == payload
    assert payload["coverage"] == sorted(payload["coverage"])


def test_multiple_findings_have_stable_codes_and_precise_locations(tmp_path: Path) -> None:
    source = tmp_path / "many-issues.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Amount", " amount ", "Notes", None])
    sheet.append([" 123 ", "Alpha\u200b", None, "merged", None])
    sheet.append([None, None, None, None, None])
    sheet.append(["Beta", 5, None, None, None])
    sheet.merge_cells("D2:E2")
    sheet["Z200"] = "stray"
    workbook.save(source)
    workbook.close()
    before = source.read_bytes()

    first = scan_workbook_health(source)
    second = scan_workbook_health(source)

    assert source.read_bytes() == before
    assert first.to_dict() == second.to_dict()
    locations = {(item.code, item.location) for item in first.findings}
    assert ("duplicate_header", "C1") in locations
    assert ("extra_whitespace", "A2") in locations
    assert ("numeric_text", "A2") in locations
    assert ("invisible_characters", "B2") in locations
    assert ("blank_data_rows", "A3:D3") in locations
    assert ("abnormal_merged_cells", "D2:E2") in locations
    assert ("content_beyond_working_area", "Z200") in locations
    assert all(item.sheet == "Data" for item in first.findings)


def test_chinese_path_and_sheet_name_are_preserved(tmp_path: Path) -> None:
    source = tmp_path / "月度报表.xlsx"
    _save_healthy(source, sheet_name="销售数据")

    result = scan_workbook_health(source)

    assert result.source == str(source)
    assert result.sheets == ("销售数据",)
    assert result.assessment == "no_known_issues"


def test_empty_workbook_is_not_reported_as_healthy(tmp_path: Path) -> None:
    source = tmp_path / "empty.xlsx"
    workbook = Workbook()
    workbook.save(source)
    workbook.close()

    result = scan_workbook_health(source)

    assert result.assessment == "findings_detected"
    assert [(item.code, item.sheet, item.location) for item in result.findings] == [
        ("default_empty_sheet", "Sheet", "A1"),
        ("empty_workbook", "Sheet", "A1"),
    ]


def test_xlsm_bytes_and_embedded_vba_member_remain_unchanged(tmp_path: Path) -> None:
    xlsx = tmp_path / "macro-source.xlsx"
    source = tmp_path / "宏工作簿.xlsm"
    _save_healthy(xlsx)
    source.write_bytes(xlsx.read_bytes())
    marker = b"synthetic-vba-marker-not-executable"
    with ZipFile(source, "a") as archive:
        archive.writestr("xl/vbaProject.bin", marker)
    before = source.read_bytes()

    result = scan_workbook_health(source)

    assert result.source_sha256 == sha256(before).hexdigest()
    assert source.read_bytes() == before
    with ZipFile(source) as archive:
        assert archive.read("xl/vbaProject.bin") == marker


@pytest.mark.parametrize("name", ["book.csv", "book.xls"])
def test_unsupported_formats_fail_closed(tmp_path: Path, name: str) -> None:
    source = tmp_path / name
    source.write_bytes(b"not a supported workbook")

    with pytest.raises(HealthScanError, match="supports only"):
        scan_workbook_health(source)


def test_unreadable_workbook_fails_closed(tmp_path: Path) -> None:
    source = tmp_path / "broken.xlsx"
    source.write_bytes(b"not a zip archive")

    with pytest.raises(HealthScanError, match="could not be scanned"):
        scan_workbook_health(source)


def test_leading_zero_identifier_is_not_classified_as_numeric_text(tmp_path: Path) -> None:
    source = tmp_path / "identifiers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Identifier", "Value"])
    sheet.append(["00123", "42"])
    workbook.save(source)
    workbook.close()

    result = scan_workbook_health(source)

    numeric_locations = [item.location for item in result.findings if item.code == "numeric_text"]
    assert numeric_locations == ["B2"]
