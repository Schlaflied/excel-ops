from __future__ import annotations

import json
from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from excel_ops.template_writer import TemplateMapping, TemplateWriteError, write_template


@dataclass(frozen=True)
class FixedPeriod:
    period_start: date
    period_end: date
    as_of_date: date


def _department_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payroll"
    sheet["A1"] = "Company banner"
    sheet.merge_cells("A1:D1")
    sheet["A3"], sheet["B3"], sheet["C3"] = "Employee", "Hours", "Pay"
    sheet["A4"] = "do-not-replace"
    sheet["B4"] = 0
    sheet["C4"] = "=B4*25"
    sheet["A4"].fill = PatternFill("solid", fgColor="FFFF00")
    sheet["D8"] = "Approval: Finance"
    workbook.save(path)


def _client_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Upload"
    sheet["A1"] = "Client template"
    sheet["B5"], sheet["D5"], sheet["F5"] = "Pay", "Employee", "Hours"
    sheet["H2"] = "UNAUTHORISED"
    workbook.save(path)


def test_same_records_write_to_two_layouts_and_preserve_templates(tmp_path: Path):
    records = [{"employee": "Ada", "hours": 8, "pay": 200}, {"employee": "Lin", "hours": 4, "pay": 100}]
    first = tmp_path / "payroll-2026-08-31.xlsx"
    second = tmp_path / "client.xlsx"
    _department_template(first)
    _client_template(second)
    first_bytes, second_bytes = first.read_bytes(), second.read_bytes()
    period = FixedPeriod(date(2026, 9, 1), date(2026, 9, 7), date(2026, 9, 12))

    department = write_template(
        first,
        records,
        TemplateMapping(
            sheet="Payroll",
            header_row=3,
            data_start_row=4,
            field_columns={"employee": "A", "hours": "B", "pay": "C"},
            style_source_row=4,
        ),
        output_dir=tmp_path / "out",
        period=period,
    )
    client = write_template(
        second,
        records,
        TemplateMapping(
            sheet="Upload",
            header_row=5,
            data_start_row=6,
            field_columns={"employee": "D", "hours": "F", "pay": "B"},
            template_type="client_upload",
        ),
        output_path=tmp_path / "out" / "client-delivery.xlsx",
    )

    assert first.read_bytes() == first_bytes
    assert second.read_bytes() == second_bytes
    assert department.output_path.name == "payroll-2026-09-12.xlsx"
    assert department.output_path.is_file()
    assert client.output_path.is_file()

    written = load_workbook(department.output_path, data_only=False)
    sheet = written["Payroll"]
    assert sheet["A4"].value == "Ada"
    assert sheet["B5"].value == 4
    assert sheet["C4"].value == "=B4*25"
    assert sheet["D8"].value == "Approval: Finance"
    assert sheet["A5"].fill.fgColor.rgb == sheet["A4"].fill.fgColor.rgb
    written.close()
    assert any(item.reason == "protected_formula" for item in department.skipped)

    written = load_workbook(client.output_path)
    assert written["Upload"]["D6"].value == "Ada"
    assert written["Upload"]["F7"].value == 4
    assert written["Upload"]["H2"].value == "UNAUTHORISED"
    written.close()

    log = json.loads(department.change_log_path.read_text(encoding="utf-8"))
    assert log["output_path"] == str(department.output_path)
    assert log["written_cells"] == 5
    assert log["skipped_items"] == 1
    assert {item["reason"] for item in log["skipped"]} == {"protected_formula"}


def test_rejects_source_overwrite_missing_sheet_and_row_overflow(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    _client_template(template)
    mapping = TemplateMapping(sheet="Upload", field_columns={"employee": "D"}, max_rows=1)

    with pytest.raises(TemplateWriteError, match="must not overwrite"):
        write_template(template, [{"employee": "Ada"}], mapping, output_path=template)
    existing = tmp_path / "existing.xlsx"
    existing.write_bytes(b"preserve this delivery")
    with pytest.raises(TemplateWriteError, match="already exists"):
        write_template(template, [{"employee": "Ada"}], mapping, output_path=existing)
    assert existing.read_bytes() == b"preserve this delivery"
    with pytest.raises(TemplateWriteError, match="exceeds"):
        write_template(
            template,
            [{"employee": "Ada"}, {"employee": "Lin"}],
            mapping,
            output_path=tmp_path / "overflow.xlsx",
        )
    with pytest.raises(TemplateWriteError, match="mapped sheet"):
        write_template(
            template,
            [{"employee": "Ada"}],
            TemplateMapping(sheet="Missing", field_columns={"employee": "A"}),
            output_path=tmp_path / "missing.xlsx",
        )


def test_merged_non_anchor_and_missing_fields_are_logged(tmp_path: Path):
    template = tmp_path / "merged.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.merge_cells("A2:B2")
    workbook.save(template)

    result = write_template(
        template,
        [{"name": "Ada"}],
        TemplateMapping(sheet="Data", field_columns={"name": "B"}),
        output_path=tmp_path / "written.xlsx",
    )

    assert result.changes == ()
    assert result.skipped[0].reason == "merged_non_anchor"
